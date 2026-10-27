"""不连罐头的测试:找日期 / 员工 / 视频、认封面、裁剪框算多远、记录、挪文件夹、界面能打开、绝不点发布。
    PYTHONIOENCODING=utf-8 python tests/test_basic.py
"""
import os
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QE_DATA_DIR", tempfile.mkdtemp(prefix="qe_data_"))
os.environ.setdefault("QE_NO_WINFIX", "1")                 # 测试时别去恢复真罐头的窗口
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

fails = []

from qe import gamecheck as _gc                             # noqa: E402
AI_GAME_CALLS = []                                           # 这台电脑上有真的百炼 Key:测试里绝不真问 AI(花钱)


def _fake_ask_ai(own, titles, key, **k):
    AI_GAME_CALLS.append((own, list(titles)))
    return {t: ("PUBG" if "苟分" in t else "") for t in titles}


_gc.ask_ai = _fake_ask_ai
import gui.api_tab as _api_mod                               # noqa: E402
_api_mod.ApiTab.AUTO_UNSURE = False                          # 测试里不自动去企鹅号查结果不明(会拿假钥匙调真接口)
_REAL_ASK = _api_mod.ask                                     # 有的测试会换掉 ask():留一份真的


def check(c, m):
    print(("  OK   " if c else "  FAIL ") + m)
    if not c:
        fails.append(m)


def _img(path, w, h):
    from PIL import Image
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (w, h), (200, 200, 200)).save(path)


def test_folders(tmp: Path):
    print("[素材文件夹]")
    from qe import folders as F
    from qe.config import Settings
    st = Settings(tmp / "s.json")
    root = tmp / "小学生教育"
    st["material_root"] = str(root)
    st["archive_root"] = str(tmp / "内容查看" / "{month}月发布视频" / "企鹅号" / "教育")
    for p in ("10.15/研习空间/31条", "10.15/研习空间/51", "10.15/小天地/51", "10.14/研习社/卓13", "10.15/不认识/甲", "9.30/小天地/乙"):
        (root / p).mkdir(parents=True)
    e = root / "10.15/研习空间/31条"
    for n in ("b.第二条", "a.第一条"):
        for ext in (".mp4", ".png", ".jpg"):
            (e / f"{n}{ext}").write_bytes(b"x")
    (e / "只有视频.mp4").write_bytes(b"x")
    (e / "多余的图.png").write_bytes(b"x")
    check(F.list_dates(str(root)) == ["9.30", "10.14", "10.15"], f"日期按日子排 -> {F.list_dates(str(root))}")
    ad = F.account_dirs(str(root), "10.15", st)
    check([os.path.basename(x) for x in ad["家长研习空间"]] == ["研习空间"] and [os.path.basename(x) for x in ad["家长小天地"]] == ["小天地"], "缩写认号")
    check([os.path.basename(x) for x in F.account_dirs(str(root), "10.14", st)["家长研习空间"]] == ["研习社"], "10.14 叫研习社也认")
    check(F.unknown_dirs(str(root), "10.15", st) == ["不认识"], "不认识的缩写挑出来提醒")
    vs = F.videos_in(str(e), "家长研习空间", "10.15")
    check([v.title for v in vs] == ["a.第一条", "b.第二条", "只有视频"], f"视频按文件名排,图片不算视频 -> {[v.title for v in vs]}")
    check(vs[0].png.endswith("a.第一条.png") and vs[0].jpg.endswith("a.第一条.jpg") and not vs[2].png, "同名的横 / 竖封面配上")
    check(vs[0].employee == "31条" and vs[0].date == "10.15", "员工 = 员工文件夹名")
    dst = F.archive_target(st, "10.15")
    check(dst.endswith(os.path.join("10月发布视频", "企鹅号", "教育", "10.15")), f"挪到 10月发布视频 -> {dst}")
    F.move_date(st, "10.15")
    check(not (root / "10.15").exists() and (Path(dst) / "研习空间" / "31条" / "a.第一条.mp4").is_file(), "整天剪切过去")
    (root / "10.15/研习空间/31条").mkdir(parents=True)
    (root / "10.15/研习空间/31条/a.第一条.mp4").write_bytes(b"y")
    try:
        F.move_date(st, "10.15")
        check(False, "那边已经有了要先问")
    except FileExistsError:
        check(True, "那边已经有同名日期:先问人(抛 FileExistsError)")
    F.move_date(st, "10.15", merge=True)
    check((Path(dst) / "研习空间/31条/a.第一条 (2).mp4").is_file() and not (root / "10.15").exists(), "合并:同名不覆盖,改名 (2)")


def test_cover_dir(tmp: Path):
    print("[员工「梁」的放法:封面\标题\横屏16x9.png]")
    from qe import folders as F
    e = tmp / "梁"
    (e / "封面" / "湖南省科技馆_C_馆上新").mkdir(parents=True)
    (e / "封面" / "两个都有").mkdir(parents=True)
    (e / "封面" / "只有修改").mkdir(parents=True)
    (e / "封面" / "心理学视角：区分情绪，减少吼叫").mkdir(parents=True)
    for n in ("湖南省科技馆 C 馆上新", "两个都有", "只有修改", "心理学视角：区分情绪,减少吼叫"):
        (e / f"{n}.mp4").write_bytes(b"x")
        (e / f"{n}.jpg").write_bytes(b"x")
    _img(e / "封面" / "湖南省科技馆_C_馆上新" / "横屏16x9.png", 320, 180)
    _img(e / "封面" / "两个都有" / "横屏16x9.png", 320, 180)
    _img(e / "封面" / "两个都有" / "竖屏9x16.png", 180, 320)
    _img(e / "封面" / "两个都有" / "原图.png", 320, 180)
    _img(e / "封面" / "只有修改" / "原图.png", 320, 180)
    _img(e / "封面" / "只有修改" / "修改 (1).png", 320, 180)
    _img(e / "封面" / "心理学视角：区分情绪，减少吼叫" / "横屏16x9.png", 320, 180)
    vs = {v.title: v for v in F.videos_in(str(e), "号", "10.15")}
    check(vs["湖南省科技馆 C 馆上新"].png.endswith("横屏16x9.png"), "名字里空格 / 下划线对不齐也认")
    check(vs["心理学视角：区分情绪,减少吼叫"].png.endswith("横屏16x9.png"), "全角半角逗号对不齐也认")
    check(vs["两个都有"].png.endswith("横屏16x9.png") and vs["两个都有"].vpng.endswith("竖屏9x16.png"), "横屏、竖屏都认出来")
    check(vs["只有修改"].png.endswith("修改 (1).png"), "没有横屏:用「修改 (1).png」,不用原图")
    check(all("原图" not in v.png and not v.png.endswith(".jpg") for v in vs.values()), "原图、带字的 jpg 一律不用")


def test_cover(tmp: Path):
    print("[封面]")
    from qe import cover as C
    from qe.folders import Video
    # 实测:画布 (393,167,677,342),图 2848×1600,3:4 框默认在 603~860
    cv = {"left": 393, "top": 167, "width": 677, "height": 342}
    x0, y0, x1, y1 = C.box_drag(cv, 2848, 1600, 0.75, {})
    check(abs(x0 - 731.6) < 2 and x0 == x1, f"不用拖:起点 = 框中间 -> {x0:.1f}")
    x0, y0, x1, y1 = C.box_drag(cv, 2848, 1600, 0.75, {"cx": 0.684, "cy": 0.4})
    want = 427.1 + 0.684 * 608.8 - 128.25 - 603.3
    check(abs((x1 - x0) - want) < 2 and abs(y1 - y0) < 1, f"3:4 往右拖到人中间 -> {x1 - x0:.1f}(应 {want:.1f})")
    x0, y0, x1, y1 = C.box_drag(cv, 2848, 1600, 9 / 16, {"cx": 0.98, "cy": 0.4})
    check(abs(x1 - (427.1 + 608.8 - 96.2)) < 2, f"人在最边上:框只拖到图边为止 -> {x1:.1f}")
    x0, y0, x1, y1 = C.box_drag(cv, 2848, 1600, 2.35, {"cx": 0.5, "cy": 0.2})
    check(abs(x1 - x0) < 1 and y1 < y0, "2.35:1 框和图一样宽:只上下挪")
    check(C.ratio_of("3:4") == 0.75 and C.ratio_of("2.35:1") == 2.35 and C.ratio_of("") == 1.0, "比例字样认得")
    import glob
    found = glob.glob(r"X:\知乎视频\小学生教育\10.15\研习空间\31条*")          # 用户发完会改名成「31条（已发）」
    real = found[0] if found else ""
    if real and os.path.isdir(real):
        def mk(stem):
            return Video(path=os.path.join(real, stem + ".mp4"), title=stem, account="家长研习空间", employee="31条", date="10.15",
                         png=os.path.join(real, stem + ".png"), jpg=os.path.join(real, stem + ".jpg"))
        v = mk("七条同学交往准则，帮助小学生建立良好校园人际关系")
        C.decide(v)
        check(v.cover_kind == "横" and 0.6 < v.plan.get("cx", 0) < 0.75 and not v.warn, f"一个人在右边:横封面,框对准他 -> {v.cover_why} {v.plan.get('cx')}")
        v = mk("三款每日十分钟亲子小游戏，帮助提升儿童注意力水平")
        C.decide(v)
        check(v.cover.endswith(".png") and v.warn and "脸最大" in v.cover_why, f"三个小孩分得开:还是用 png(jpg 有字),框对准脸最大的,标黄 -> {v.cover_why}")
        v = mk("十种添加剂含量较高零食，建议家长减少给孩子食用")
        C.decide(v)
        check(v.cover_kind == "横" and not v.warn and 0.45 < v.plan.get("cx", 0) < 0.7, f"没人(零食):横封面,框对准零食堆 -> {v.cover_why} {v.plan}")
        g = mk("五款辅助写作业工具，帮助改善孩子写作业专注力问题")
        C.decide(g)
        check(g.plan.get("cx", 0) > 0.58, f"背影写作业的小孩(认不出脸):框对准小孩那边 -> {g.plan.get('cx')}")
        liang = r"X:\知乎视频\小学生教育\10.15\研习空间\梁"
        if os.path.isdir(liang):
            from qe import folders as F
            lv = next((x for x in F.videos_in(liang, "家长研习空间", "10.15") if x.title.startswith("孩子遭遇同伴孤立")), None)
            if lv:
                C.decide(lv)
                check(lv.vcover.endswith("竖屏9x16.png") and lv.cover.endswith("横屏16x9.png") and not lv.warn,
                      f"两个人分得开、有无字竖屏:竖的比例用竖屏图 -> {lv.cover_why}")
                check(C.pick_for(lv, 0.75)[0] == lv.vcover and C.pick_for(lv, 2.35)[0] == lv.cover and C.pick_for(lv, 1.0)[0] == lv.cover,
                      "3:4 用竖图,2.35:1 / 1:1 用横图")
    else:
        print("  (网络盘上没有真封面,跳过)")
    v = Video(path="x.mp4", title="x", account="a", employee="e", date="d")
    C.decide(v)
    check(v.warn and not v.cover, "一张封面都没有:标黄")
    v = Video(path="x.mp4", title="x", account="a", employee="e", date="d", jpg="x.jpg")
    C.decide(v)
    check(v.warn and not v.cover and "jpg 有字" in v.cover_why, f"只有 jpg:不用(有字),标黄 -> {v.cover_why}")


def test_records(tmp: Path):
    print("[记录]")
    import openpyxl
    from qe.folders import Video
    from qe.records import Records
    r = Records(tmp / "r.sqlite")
    a = Video(path="p1", title="一", account="家长研习空间", employee="31条", date="10.15", cover_kind="横")
    b = Video(path="p2", title="二", account="家长研习空间", employee="31条", date="10.15", cover_kind="竖")
    c = Video(path="p3", title="三", account="家长研习空间", employee="51", date="10.15")
    r.prepared(a, "B1")
    r.prepared(b, "B1")
    r.failed(c, "B1", "封面没设上")
    check(r.status_of("p1") == "摆好" and r.status_of("p3") == "", "摆好 / 没摆成(没摆成的下次还摆)")
    check(r.open_batches()[0]["n"] == 2, "有一批等确认")
    check(r.mark_batch_published("B1") == 2 and r.status_of("p2") == "已发", "这批发好了:记已发")
    check(all(x["published_at"].startswith("20") for x in r.rows("10.15")), "记了发布时间")
    s = r.summary("10.15")
    check(len(s) == 1 and s[0]["employee"] == "31条" and s[0]["sent"] == 2, f"按员工数 -> {s}")
    r.prepared(c, "B2")
    check(r.drop_batch("B2") == 1 and r.status_of("p3") == "", "这批不要了:作废,下次还摆")
    out = tmp / "记录.xlsx"
    r.export_xlsx(str(out))
    wb = openpyxl.load_workbook(out)
    check(wb.sheetnames == ["明细", "按员工汇总"] and wb["明细"].max_row == 3, "导出 Excel")


def test_failed_steps(tmp: Path):
    print("[封面没成也接着做声明 / 手动发的标成已发]")
    from qe import prepare as P
    from qe.folders import Video
    from qe.records import Records
    done = []

    class FakeRunner(P.Runner):
        def __init__(self):
            self.log, self.stopped, self.shot_dir, self.page = (lambda m: None), (lambda: False), "", None

        def open_account(self, name):
            pass

        def pending_videos(self):
            return []

        def go_create(self):
            pass

        def upload(self, vs):
            pass

        def ensure_visible(self):
            pass

        def select_tab(self, t):
            pass

        def wait_uploaded(self, t):
            pass

        def set_cover(self, v):
            done.append(("封面", v.title))
            if v.title == "坏封面":
                raise P.PrepError("企鹅号说封面图尺寸太小")

        def set_declaration(self):
            done.append(("自主声明",))

        def set_ai_declaration(self):
            done.append(("AI",))

        def shot(self, name):
            pass

        def _close_dialog(self):
            pass

    vs = [Video(path="a", title="坏封面", account="号", employee="员", date="1.1"),
          Video(path="b", title="好的", account="号", employee="员", date="1.1")]
    res = FakeRunner().prepare_batch("号", vs)
    check(done[:3] == [("封面", "坏封面"), ("自主声明",), ("AI",)], f"封面没成:自主声明、AI 补充照样做 -> {done[:3]}")
    check(not res[0][1] and "封面" in res[0][2] and "手动补" in res[0][2] and res[1][1], f"这条算没摆成,说清楚是封面 -> {res[0][2]}")
    r = Records(tmp / "rf.sqlite")
    r.prepared(vs[1], "B")
    r.failed(vs[0], "B", "封面")
    check([x["title"] for x in r.failed_in_batch("B")] == ["坏封面"], "这批里没摆成的认得出")
    check(r.mark_failed_published("B") == 1 and r.status_of("a") == "已发", "手动补好一起发了:也记成已发")
    c = Video(path="c", title="手动发的", account="号", employee="员", date="1.1")
    r.manual_published(c)
    check(r.status_of("c") == "已发" and r.rows("1.1")[-1]["published_at"], "「选中的标成已发」:记已发、有时间")


def test_api(tmp: Path):
    print("[接口发:找素材 / 话题留一个 / 平台名 / 转竖屏 / 账号加密 / 发布流程(假接口)]")
    from qe import apiacc, apisrc, checks, omapi, topic, vertical
    root = tmp / "api"
    g = root / "游戏视频"
    for d in ("CSGO--CSGO战神君/10.15(20)", "CSGO--CSGO战神君/不能用", "CSGO--CSGO战神君/10.14", "1疑似脏话/10.15",
              "地下城与勇士--DNF连招档案馆/10.15（20）"):
        (g / d).mkdir(parents=True)
    (root / "动漫视频" / "国漫趣谈" / "10.15（星期四）").mkdir(parents=True)
    for f in ("CSGO--CSGO战神君/10.15(20)/a #CSGO #反恐精英.mp4", "CSGO--CSGO战神君/10.15(20)/a #CSGO #反恐精英.jpg",
              "CSGO--CSGO战神君/不能用/坏.mp4", "1疑似脏话/10.15/脏.mp4", "地下城与勇士--DNF连招档案馆/10.15（20）/b.mp4"):
        (g / f).write_bytes(b"x")
    (root / "动漫视频" / "国漫趣谈" / "10.15（星期四）" / "抖音同款 #牧神记.mp4").write_bytes(b"x")
    roots = [str(g), str(root / "动漫视频")]
    check(apisrc.date_key("10.14（星期三）") == "10.14" and apisrc.date_key("10.15(20)") == "10.15" and apisrc.date_key("不能用") == "",
          "日期文件夹只认开头的「月.日」")
    known = {"CSGO战神君", "DNF连招档案馆", "国漫趣谈"}
    check(apisrc.dates(roots, known) == ["10.14", "10.15"], "有哪些日子")
    vs = apisrc.videos(roots, "10.15", known)
    check(sorted(v.account for v in vs) == ["CSGO战神君", "DNF连招档案馆", "国漫趣谈"], "「不能用」「1疑似脏话」不算")
    check(apisrc.archive_dir({}, "10.15") == r"X:\知乎视频\内容查看\10月发布视频\企鹅号\企鹅号软件发布\10.15",
          "发完挪到「企鹅号软件发布\日期」(日期 = 发出去那天)")
    check(apisrc.archive_dir({}).endswith(apisrc.pub_day()) and f"{apisrc.pub_day().split('.')[0]}月发布视频" in apisrc.archive_dir({}),
          "不给日期就按今天(北京时间)算,{month} 也是")

    p = topic.split_topics("标题 #a #B #b")
    check(p["body"] == "标题" and p["sep"] == " " and p["tags"] == ["a", "B"], "话题拆开、不分大小写去重")
    check(topic.split_topics("标题#a 中间 #b")["tags"] == ["b"], "只认结尾那一串")
    k = "狼队危险了# 2026kpl年度总决赛 # kpl十周年 # 2026kpl年度总决赛 # kpl十周年 "
    check(topic.split_topics(k)["tags"] == ["2026kpl年度总决赛", "kpl十周年"], "认「# 话题」(# 后面带空格)和重复的")
    check(topic.new_name(k, lambda b, t: t[0]) == "狼队危险了#2026kpl年度总决赛", "留下的写成「#话题」")
    check(topic.new_name("标题 # 牧神记") == "标题 #牧神记" and topic.split_topics("标题＃甲＃乙")["tags"] == ["甲", "乙"],
          "只有一个「# 话题」也改写法;全角 ＃ 也认")
    check(topic.split_topics("取胜n# 2026kpl # kpl十周年")["body"] == "取胜" and topic.split_topics("神作 ##怪化猫")["tags"] == ["怪化猫"]
          and topic.split_topics("Join#话题")["body"] == "Join", "认「##话题」「取胜n#话题」(n 是丢了的换行),英文单词的 n 不动")
    check(topic.with_topic("团战#英雄联盟手游 #lol", "#英雄 联盟") == "团战#英雄联盟" and topic.with_topic("没话题 ", "英雄联盟") == "没话题#英雄联盟"
          and topic.with_topic("a#英雄联盟", "英雄联盟") is None, "批量改话题:全换成填的那个;没话题就接上;一样的不改")
    check(topic.bad_format("标题 # 剑来") and topic.bad_format("神作 ##怪化猫") and topic.bad_format("《剑来》 #剑来 ")
          and topic.bad_format("打法。n#三角洲行动") and topic.bad_format("没话题 ") and not topic.bad_format("a #x #y")
          and not topic.bad_format("正常#三角洲行动"), "「话题写法不对」判严格(结尾空格、n# 也算)")
    check(topic.tidy("《剑来》 #剑来 ") == "《剑来》 #剑来" and topic.tidy("打法。n#三角洲行动") == "打法。#三角洲行动"
          and topic.tidy("神作 ##怪化猫") == "神作 #怪化猫" and topic.tidy("x # a # a") == "x #a" and topic.tidy("没话题 ") == "没话题",
          "整理写法:一键改好(不用 AI)")
    check("动画名" in topic.TOPIC_SYSTEM and "留下的就必须是作品名" in topic.TOPIC_SYSTEM, "作品名(游戏 / 剧 / 电影 / 动画)优先")
    check(topic.new_name("x #a #a") == "x #a" and topic.new_name("x #a") is None and topic.new_name("x") is None, "0/1 个不问 AI")
    check(topic.new_name("x #活动 #三角洲行动", lambda b, t: t[1]) == "x #三角洲行动", "2 个以上留 AI 挑的")
    check(topic.parse_order('```json\n{"理由":"x","顺序":[2,1]}\n```', 2) == [2, 1] and topic.parse_order("2", 3) == [2, 1, 3]
          and topic.parse_order("不知道", 2) == [], "AI 回话读顺序")
    m = topic.topic_messages("正文", ["a", "b"])
    check("排第一的那个会被留下" in m[0]["content"] and "1. #a\n2. #b" in m[1]["content"], "提示词和 ainame 一样")

    check(checks.platform_hits("抖音同款 #牧神记") == ["抖音"] and checks.platform_hits("everybody books") == []
          and checks.platform_hits("来自 dy 的") == ["dy"], "标题里的别的平台名(dy 只认单独的)")

    v = [x for x in vs if x.account == "CSGO战神君"][0]
    led = topic.Ledger(tmp / "led.json")
    moves = topic.rename_with_sidecars(v.path, topic.new_name(v.title, lambda b, t: t[0]))
    led.add("t1", moves)
    d = os.path.dirname(v.path)
    check(sorted(os.listdir(d)) == ["a #CSGO.jpg", "a #CSGO.mp4"], "改名:同名封面跟着改")
    check(led.undo_last() == (2, 0, 0) and sorted(os.listdir(d)) == ["a #CSGO #反恐精英.jpg", "a #CSGO #反恐精英.mp4"], "能撤销")

    check(apiacc.unseal(apiacc.seal("秘密abc")) == "秘密abc" and "秘密" not in apiacc.seal("秘密abc"), "密钥加密存")
    import sqlite3
    db = tmp / "old.db"
    con = sqlite3.connect(db)
    con.execute("create table accounts(id integer primary key, source_name, source_id, client_id, client_secret, nick, category, enabled, sort_order)")
    con.execute("insert into accounts values(1,'CSGO','1','cid','sec1','CSGO战神君',1604,1,0)")
    con.commit()
    con.close()
    check(apiacc.migrate(str(db)) == 1 and apiacc.find("CSGO战神君")["category"] == 1604, "从旧数据库搬账号")
    check("sec1" not in apiacc.store_path().read_text("utf-8"), "存下来的文件里没有明文密钥")

    class R:
        def __init__(self, j):
            self.j, self.status_code, self.text = j, 200, ""

        def json(self):
            return self.j

    class S:
        def __init__(self):
            self.calls, self.n, self.trunks = [], 0, []

        def post(self, url, **kw):
            self.calls.append(("token", kw["params"]["client_secret"]))
            return R({"code": "0", "data": {"access_token": "T", "expires_in": 7200}})

        def request(self, method, url, **kw):
            path = url.split("api.om.qq.com")[1]
            self.calls.append((path, kw.get("json")))
            assert kw["params"]["access_token"] == "T"
            if path == "/video/clientuploadready":
                self.size = kw["json"]["size"]
                return R({"code": "0", "data": {"transaction_id": "U1"}})
            if path == "/video/clientuploadtrunk":       # 和真接口一样:回下一片的起止;最后一片回 文件大小-1;文件名只能英文
                off = int(kw["data"]["start_offset"])
                name, chunk, _ct = kw["files"]["mediatrunk"]
                assert name.isascii(), "文件名得是英文"
                assert off == sum(len(c) for c in self.trunks), "多传 / 少传了"
                self.trunks.append(chunk)
                nxt = off + len(chunk)
                if nxt >= self.size:              # 最后一片:回 start = end = 文件大小 - 1(真接口)
                    return R({"code": "0", "data": {"start_offset": self.size - 1, "end_offset": self.size - 1}})
                return R({"code": "0", "data": {"start_offset": nxt, "end_offset": min(self.size, nxt + len(chunk)) - 1}})
            if path == "/articlev2/clientpubvideo":
                return R({"code": "0", "data": {"transaction_id": "P1", "article_id": "A9"}})
            if path == "/transaction/infoclient":
                if kw["params"]["transaction_id"] == "U1":
                    return R({"code": "0", "data": {"article_info": {"article_video_info": {}}, "transaction_status": "成功",
                                                    "vid": "V7"}})       # 真接口 vid 在最外层
                self.n += 1
                return R({"code": "0", "data": {"transaction_status": "处理中" if self.n < 2 else "成功"}})
            return R({"code": "-1", "msg": "?"})
    s = S()
    real_sleep = omapi.time.sleep
    omapi.time.sleep = lambda x: None
    try:
        r = omapi.Client("cid", "sec1", session=s).send(v.path, v.title, 1604)
    finally:
        omapi.time.sleep = real_sleep
    pub = [c[1] for c in s.calls if c[0] == "/articlev2/clientpubvideo"][0]
    big = tmp / "big.mp4"
    big.write_bytes(bytes(range(10)))
    s2 = S()
    real_trunk = omapi.TRUNK
    omapi.TRUNK = 3
    try:
        omapi.Client("cid", "sec1", session=s2).upload(str(big))
    finally:
        omapi.TRUNK = real_trunk
    check(b"".join(s2.trunks) == bytes(range(10)) and len(s2.trunks) == 4, "分片上传:多片按平台说的下一片位置接着传,不跳片、传完不多传")
    class Flaky(S):                              # 第二片第一次断网:原地重传这一片,接着传完
        def request(self, method, url, **kw):
            if url.endswith("clientuploadtrunk") and kw["data"]["start_offset"] == "3" and not getattr(self, "boom", False):
                self.boom = True
                import requests as _rq
                raise _rq.ConnectionError("断网")
            return S.request(self, method, url, **kw)
    s3 = Flaky()
    omapi.TRUNK = 3
    real_sleep = omapi.time.sleep
    omapi.time.sleep = lambda x: None
    try:
        omapi.Client("cid", "sec1", session=s3).upload(str(big))
    finally:
        omapi.TRUNK = real_trunk
        omapi.time.sleep = real_sleep
    check(b"".join(s3.trunks) == bytes(range(10)) and s3.boom, "上传一片断网:原地重传这一片,不用整条重来")
    check(r == {"vid": "V7", "article_id": "A9", "transaction_id": "P1", "status": "成功", "cover": ""}, "发布流程:传 → vid → 发 → 等成功")
    check(pub == {"title": v.title, "tags": "", "vid": "V7", "cat": 1604, "desc": v.title, "aigc_userstatement": "2",
                  "ext_info": "{\"self_declare\": 0}"},
          "发的内容:标题 = 简介 = 文件名,分类按号,AI 声明非 AI,自主声明无需标注(self_declare=0)")
    dst = apisrc.move_published(v.path, str(tmp / "已发布"))
    check(os.path.isfile(dst) and not os.path.exists(v.path), "发完挪走")

    ff = vertical.ffmpeg_path()
    if ff:
        src = tmp / "横.mp4"
        import subprocess
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-f", "lavfi",
                        "-i", "sine", "-t", "2", "-c:v", "libx264", "-c:a", "aac", str(src)], check=True)
        check(vertical.is_landscape(vertical.probe(str(src), ff)), "认出横屏")
        new, orig = vertical.to_vertical(str(src), ff)
        info = vertical.probe(new, ff)
        check((info["w"], info["h"]) == (1080, 1920) and os.path.isfile(orig) and "横屏原片" in orig, "转成竖屏,原片挪进「横屏原片」")
    else:
        print("  (没有 ffmpeg,跳过转竖屏)")


def _settle(tab):
    """等后台(换日期 / 刷新后自动检查)做完。"""
    from PySide6.QtWidgets import QApplication
    if tab.job:
        tab.job.wait()
    QApplication.processEvents()


def test_api_tab(tmp: Path):
    print("[接口发页:手选文件夹(旧软件的「选择」)/ 一条一条发、间隔 / 发完挪走、记录]")
    from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton
    QApplication.instance() or QApplication([])
    from qe import apiacc, apisrc, omapi
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "api2"
    (root / "游戏视频" / "CSGO--CSGO战神君" / "10.15").mkdir(parents=True)
    (root / "游戏视频" / "2电竞赛事（腾讯系相关游戏赛事）").mkdir(parents=True)
    for f in ("CSGO--CSGO战神君/10.15/c1.mp4", "CSGO--CSGO战神君/10.15/c2.mp4",
              "2电竞赛事（腾讯系相关游戏赛事）/e1.mp4", "2电竞赛事（腾讯系相关游戏赛事）/e1.jpg"):
        (root / "游戏视频" / f).write_bytes(b"x")
    esp = str(root / "游戏视频" / "2电竞赛事（腾讯系相关游戏赛事）")
    roots = [str(root / "游戏视频")]
    known = {"CSGO战神君", "电竞情报馆"}
    vs = apisrc.videos(roots, "10.15", known, {"电竞情报馆": esp})
    check(sorted((v.account, v.title) for v in vs) == [("CSGO战神君", "c1"), ("CSGO战神君", "c2"), ("电竞情报馆", "e1")],
          "手选的文件夹:电竞情报馆发「2电竞赛事」里的")
    vs = apisrc.videos(roots, "10.15", known, {"CSGO战神君": esp})
    check([v.account for v in vs] == ["CSGO战神君"] and vs[0].title == "e1", "手选了就不再自动找这个号")

    accs = [{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True},
            {"nick": "电竞情报馆", "source_name": "电竞赛事", "client_id": "b", "secret": apiacc.seal("s"), "category": 1605, "enabled": True}]
    apiacc.save(accs)
    st = Settings(tmp / "st_api.json")
    st["api_roots"] = roots
    st["api_archive"] = str(tmp / "发布" / "{date}已发布")
    st["api_interval"] = 1
    check(st.get("api_threads") == 1 and "api_parallel" not in st, "默认一条一条发(旧软件「线程 1」)")
    rec = Records(tmp / "rec_api.sqlite")
    logs = []
    tab = T.ApiTab(st, rec, logs.append)
    tab.c_date.setCurrentText("10.15")
    tab._set_pick("10.15", "电竞情报馆", esp)
    tab.load()
    check(len(tab.videos) == 3, "页面上:自动找的 + 手选的")
    sent = []

    class FakeClient:
        def __init__(self, cid, sec, **k):
            self.cid = cid

        def send(self, path, title, cat, before_publish=None, **k):
            if before_publish:
                before_publish("V")
            sent.append((time.time(), title, cat))
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}

    real_c, real_q = omapi.Client, QMessageBox.question
    omapi.Client = FakeClient
    QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
    try:
        tab.run_check()
        tab.job.wait()
        QApplication.processEvents()
        for v in tab.videos:                       # 假视频 ffmpeg 读不出来:当成竖屏的好视频
            v.checks.update(bad_read=False, landscape=False, w=1080, h=1920)
        tab.run_publish()
        tab.job.wait()
        QApplication.processEvents()
    finally:
        omapi.Client, QMessageBox.question = real_c, real_q
    check(sorted((t, c) for _x, t, c in sent) == [("c1", 1604), ("c2", 1604), ("e1", 1605)], "三条都发了,分类按号")
    gaps = [b[0] - a[0] for a, b in zip(sent, sent[1:])]
    check(all(g >= 0.95 for g in gaps), f"每条之间歇够间隔 -> {[round(g, 2) for g in gaps]}")
    out = tmp / "发布" / f"{apisrc.pub_day()}已发布"          # 按发出去那天(今天)算,不是素材文件夹的 10.15
    check(sorted(os.listdir(out / "CSGO")) == ["c1.mp4", "c2.mp4"] and sorted(os.listdir(out / "电竞赛事")) == ["e1.jpg", "e1.mp4"]
          and not os.path.isfile(os.path.join(esp, "e1.jpg")), "发完挪走,按游戏分子文件夹(名字 = 旧软件「账号」),同名封面跟着挪")
    check(apisrc.archive_dir({"api_archive_by_game": False}, "10.15", accs[0]).endswith("企鹅号软件发布" + os.sep + "10.15")
          and apisrc.archive_dir({}, "10.15", accs[0]).endswith(os.path.join("企鹅号软件发布", "10.15", "CSGO")),
          "日期下面再分游戏;不勾「按游戏分」就只到日期")
    check(len([r for r in rec.rows("10.15") if r["employee"] == "接口" and r["status"] == "已发"]) == 3, "记进记录(员工 = 接口)")
    check(not (root / "游戏视频" / "CSGO--CSGO战神君" / "10.15").exists(), "这个号 10.15 发空了:文件夹自动删了")
    (root / "游戏视频" / "CSGO--CSGO战神君" / "10.15").mkdir()          # 员工又放了两条进来
    for f in ("c3.mp4", "c4.mp4"):
        (root / "游戏视频" / "CSGO--CSGO战神君" / "10.15" / f).write_bytes(b"x")
    tab.load()
    tab.run_check()
    tab.job.wait()
    QApplication.processEvents()
    for v in tab.videos:
        v.checks.update(bad_read=False, landscape=False, w=1080, h=1920)
    sent.clear()
    omapi.Client = FakeClient
    QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
    try:
        row = [i for i, v in enumerate(tab.videos) if v.title == "c4"][0]
        tab.tree.clearSelection()
        tab.items[row].setSelected(True)
        check(tab.selected_rows() == [row], "选中一条")
        tab.run_publish(tab.selected_rows())
        tab.job.wait()
        QApplication.processEvents()
    finally:
        omapi.Client, QMessageBox.question = real_c, real_q
    check([t for _x, t, _c in sent] == ["c4"], "只发选中的:c3 没发")
    for f in ("d1.mp4", "d2.mp4"):                    # 前面发掉的都挪走了,补两条
        (root / "游戏视频" / "CSGO--CSGO战神君" / "10.15" / f).write_bytes(b"x")
    tab.load()
    _settle(tab)
    check({n for n, g in tab.groups.items() if g.childCount()} == {"CSGO战神君"} and
          all(not g.isExpanded() for g in tab.groups.values()), "按号分组,默认收起")
    check("这天已发 1 条" in tab.groups["电竞情报馆"].text(0) and tab.cards.vals["sent"].text() != "0",
          "发完挪走了的号:写「这天已发 N 条」,不是「没找到视频」;卡片「已发」算上挪走的")
    tab.expand_all(True)
    check(all(g.isExpanded() for g in tab.groups.values()), "全部展开")
    tab.expand_all(False)
    check(all(not g.isExpanded() for g in tab.groups.values()), "全部收起")
    tab.groups["CSGO战神君"].setExpanded(True)
    tab.load()
    _settle(tab)
    check(tab.groups["CSGO战神君"].isExpanded(), "刷新后展开的还展开着")
    tab.tree.clearSelection()
    tab.groups["CSGO战神君"].setSelected(True)
    check([tab.videos[i].title for i in tab.selected_rows()] == ["c3", "d1", "d2"], "选中号 = 这个号全部")
    check("3 条" in tab.groups["CSGO战神君"].text(0) and "有问题 3" in tab.groups["CSGO战神君"].text(2)
          and "没检查" not in tab.groups["CSGO战神君"].text(2), "号那一行有统计;刷新后自动检查了")
    from PySide6.QtWidgets import QInputDialog
    real_t, real_i, real_gi = QInputDialog.getText, QMessageBox.information, QInputDialog.getItem
    QInputDialog.getText = staticmethod(lambda *a, **k: ("英雄联盟", True))
    QInputDialog.getItem = staticmethod(lambda *a, **k: ("英雄联盟", True))
    QMessageBox.information = staticmethod(lambda *a, **k: None)
    QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
    try:
        rows = [i for i, v in enumerate(tab.videos) if v.title in ("c3",)]
        tab.set_topic(rows)
        check([v.title for v in tab.videos if v.title.startswith("c3")] == ["c3#英雄联盟"] and
              os.path.isfile(os.path.join(os.path.dirname(tab.videos[rows[0]].path), "c3#英雄联盟.mp4")), "批量改话题改了文件名")
        sent_row = [i for i, v in enumerate(tab.videos) if v.title == "d1"]
        tab.state[sent_row[0]] = "已发"
        tab.set_topic(sent_row)
        check(tab.videos[sent_row[0]].title == "d1", "发过的不改")
        tab._ask_title = lambda *a, **k: ("c3 改过的标题", True)        # 改标题的框(宽的那个)
        tab.edit_title(rows[0])
        check(tab.videos[rows[0]].title == "c3 改过的标题", "右键改标题")
        check("改标题" in tab.b_undo.toolTip(), "撤销按钮上写着上一次是什么")
        tab.undo_topic()                           # 真点「撤销」(审查:原来一点就崩)
        _settle(tab)
        check(any(v.title == "c3#英雄联盟" for v in tab.videos), "撤销:只撤上一次(改标题),批量改话题还在")
        tab.undo_topic()
        _settle(tab)
        check(any(v.title == "c3" for v in tab.videos), "再撤销一次:批量改话题也撤回去")
    finally:
        QInputDialog.getText, QMessageBox.question, QMessageBox.information = real_t, real_q, real_i
        QInputDialog.getItem = real_gi
    tab.e_find.setText("c1")
    vis = [tab.videos[i].title for i, it in tab.items.items() if not it.isHidden()]
    check(vis == [], "搜标题(c1 发掉了,搜不到)")
    tab.e_find.setText("d2")
    vis = [tab.videos[i].title for i, it in tab.items.items() if not it.isHidden()]
    check(vis == ["d2"], "搜标题")
    tab.e_find.setText("")
    with open(os.path.join(esp, "e2.mp4"), "wb") as f:
        f.write(b"x")
    tab.load()
    _settle(tab)
    check("管全部" in tab.lab_scope.text(), "没选:写着「管全部」")
    tab.k_bad.setChecked(True)
    check("过滤中" in tab.lab_scope.text(), "开着「只看有问题的」:写着「过滤中:只管看得见的」,不说全部")
    tab.k_bad.setChecked(False)
    for v in tab.videos:
        v.checks = {}
    tab.tree.clearSelection()
    tab.groups["电竞情报馆"].setSelected(True)
    check("选中" in tab.lab_scope.text() and "电竞情报馆 1 条" in tab.lab_scope.text(), "选了一个号:写着「只管选中的:号 N 条」")
    tab.run_check()
    tab.job.wait()
    QApplication.processEvents()
    check([bool(v.checks) for v in tab.videos] == [v.account == "电竞情报馆" for v in tab.videos], "只检查选中的号")
    tab.tree.clearSelection()
    check("管全部" in tab.lab_scope.text(), "取消选择:回到全部")
    tab.unpick("电竞情报馆")
    check("10.15" not in (st.get("api_picked") or {}), "取消手选")


def test_audit_r1(tmp: Path):
    print("[审查第 1 轮修的:密钥不进日志 / 结果不明不重发 / 撤销不改已发的 / 转竖屏不覆盖 / 重名序号位置 / 间隔]")
    import requests
    from PySide6.QtWidgets import QApplication, QMessageBox
    QApplication.instance() or QApplication([])
    from qe import apiacc, omapi, topic, vertical
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T

    # 1. 网络错:报错里没有密钥 / token
    class Boom:
        def post(self, *a, **k):
            raise requests.ConnectionError("Max retries exceeded with url: /x?client_secret=SEC123&access_token=TOK9")
        request = post
    try:
        omapi.Client("cid", "SEC123", session=Boom()).token()
        msg = ""
    except omapi.ApiError as e:
        msg = str(e)
    check(msg and "SEC123" not in msg and "TOK9" not in msg and "网络" in msg, f"断网报错不带密钥、说人话 -> {msg}")
    check(omapi.redact("a?client_secret=xx&access_token=yy") == "a?client_secret=***&access_token=***", "redact 去掉密钥")

    # 2. 发布请求发出去没回话 → 结果不明,不再发;发成功后查进度出错 → 算已发
    class S:
        def __init__(self, pub_err=None, poll_err=None):
            self.pub_err, self.poll_err, self.pubs = pub_err, poll_err, 0

        def post(self, url, **kw):
            class R:
                status_code = 200

                def json(self_):
                    return {"code": "0", "data": {"access_token": "T", "expires_in": 7200}}
            return R()

        def request(self, method, url, **kw):
            path = url.split("api.om.qq.com")[1]
            d = {}
            if path == "/video/clientuploadready":
                d = {"transaction_id": "U"}
            elif path == "/video/clientuploadtrunk":
                d = {"end_offset": len(kw["files"]["mediatrunk"][1]) - 1 + int(kw["data"]["start_offset"])}
            elif path == "/transaction/infoclient":
                if kw["params"]["transaction_id"] == "U":
                    d = {"vid": "V"}
                elif self.poll_err:
                    raise self.poll_err
                else:
                    d = {"transaction_status": "成功"}
            elif path == "/articlev2/clientpubvideo":
                self.pubs += 1
                if self.pub_err:
                    raise self.pub_err
                d = {"transaction_id": "P", "article_id": "A"}

            class R:
                status_code = 200

                def json(self_):
                    return {"code": "0", "data": d}
            return R()
    f = tmp / "aud.mp4"
    f.write_bytes(b"0123456789")
    try:
        omapi.Client("c", "s", session=S(pub_err=requests.ReadTimeout("t"))).send(str(f), "t", 1)
        kind = "none"
    except omapi.Unknown:
        kind = "unknown"
    check(kind == "unknown", "发布请求超时 = 结果不明(不是没发成)")
    r = omapi.Client("c", "s", session=S(poll_err=requests.ReadTimeout("t"))).send(str(f), "t", 1)
    check(r["status"] == "处理中" and r["article_id"] == "A", "发成功后查进度出错 = 已发(平台处理中)")

    rec = Records(tmp / "aud.sqlite")

    class V:
        date, account, title = "1.1", "号", "t"
        path = str(f)
    rid = rec.api_begin(V)
    check(rec.blocked(V.path) and rec.status_of(V.path) == "结果不明", "发布中(程序被关)= 不能自动再发")
    rec.api_finish(rid, "结果不明", "x")
    check(rec.blocked(V.path) and len(rec.unsure_paths()) == 1, "结果不明:不能自动再发,列出来让人看")
    rec.resolve_unsure(V.path, False)
    check(not rec.blocked(V.path), "人看过了标「没发」:能再发")

    # 3. 撤销不改已发的
    d = tmp / "aud_undo"
    d.mkdir()
    (d / "c1 #a #b.mp4").write_bytes(b"x")
    led = topic.Ledger(tmp / "aud_led.json")
    moves = topic.rename_logged(str(d / "c1 #a #b.mp4"), "c1 #a", led, "t")
    rec.api_result(type("X", (), {"date": "1.1", "account": "号", "title": "c1 #a", "path": moves[0][1]}), "已发")
    ok, bad, skipped = led.undo_last(skip=rec.blocked)
    check(skipped == 1 and (d / "c1 #a.mp4").exists(), "撤销:已经发了的不改回去(防重发)")

    # 7. 重名:序号插在话题前面,话题还在结尾
    (d / "a #x.mp4").write_bytes(b"x")
    (d / "a #x #y.mp4").write_bytes(b"x")
    mv = topic.rename_with_sidecars(str(d / "a #x #y.mp4"), "a #x")
    check(os.path.basename(mv[0][1]) == "a (2) #x.mp4" and topic.split_topics("a (2) #x")["tags"] == ["x"],
          "重名:「a (2) #x」话题还在结尾")

    # 5. 转竖屏:a.mov 转出来不覆盖已有的 a.mp4
    ff = vertical.ffmpeg_path()
    if ff:
        import subprocess
        vd = tmp / "aud_vert"
        vd.mkdir()
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-t", "1",
                        "-c:v", "libx264", str(vd / "a.mov")], check=True)
        (vd / "a.mp4").write_bytes(b"keep")
        out, orig = vertical.to_vertical(str(vd / "a.mov"), ff)
        check((vd / "a.mp4").read_bytes() == b"keep" and os.path.basename(out) == "a (2).mp4", "转竖屏:不覆盖同名的另一条")
        check(not [n for n in os.listdir(vd) if n.startswith("_处理中_")], "转竖屏:不留临时文件")
        (vd / "b.mov").write_bytes(b"x")
        (vd / "横屏原片").rename(vd / "横屏原片_d")
        (vd / "横屏原片").write_bytes(b"block")
        try:
            vertical.to_vertical(str(vd / "b.mov"), ff)
            err = ""
        except RuntimeError as e:
            err = str(e)
        check(err and (vd / "b.mov").exists() and not [n for n in os.listdir(vd) if n.startswith("_处理中_")],
              f"转竖屏失败:原片还在、不留临时文件 -> {err[:40]}")

    # 8. 间隔:上一条发完以后才开始算
    root = tmp / "aud_gap"
    (root / "号A").mkdir(parents=True)
    for n in ("g1.mp4", "g2.mp4"):
        (root / "号A" / n).write_bytes(b"x")
    apiacc.save([{"nick": "号A", "source_name": "A", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "aud_st.json")
    st["api_roots"] = [str(tmp / "nothing")]
    st["api_archive"] = str(tmp / "aud_out")
    st["api_interval"] = 1
    st["api_picked"] = {"1.1": {"号A": str(root / "号A")}}
    tab = T.ApiTab(st, Records(tmp / "aud2.sqlite"), lambda m: None)
    tab.c_date.addItem("1.1")
    tab.c_date.setCurrentText("1.1")
    tab.load()
    _settle(tab)
    for v in tab.videos:
        v.checks = {"w": 1080, "h": 1920, "tags": [], "platform": []}
    marks = []

    class Slow:
        def __init__(self, *a, **k):
            pass

        def send(self, path, title, cat, before_publish=None, **k):
            marks.append(("start", time.time()))
            before_publish("V")
            time.sleep(0.6)
            marks.append(("end", time.time()))
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}
    real_c, real_q = omapi.Client, QMessageBox.question
    omapi.Client = Slow
    QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
    try:
        tab.run_publish()
        check(tab.busy() and not tab.b_pick.isEnabled(), "后台发布时「给号选文件夹」点不了")
        tab.job.wait()
        QApplication.processEvents()
    finally:
        omapi.Client, QMessageBox.question = real_c, real_q
    gap = marks[2][1] - marks[1][1] if len(marks) >= 3 else 0
    check(gap >= 0.95, f"间隔:上一条发完以后歇够了再发 -> {gap:.2f} 秒")


def test_audit_r2(tmp: Path):
    print("[审查第 1 轮(教育 / 数据安全)修的:挪文件不留两份 / 设置不互相冲 / 坏文件留底 / 路径写法 / 停了的批次 / 没 png 不挡路]")
    import json
    from qe import fsops
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records

    # 1. 挪一天:有文件被占着 → 占着的留在原处,其它挪过去,不会两边各一份;关掉再挪接着挪,没有 (2)
    root = tmp / "mv_src"
    d = root / "10.15" / "小天地" / "张三"
    d.mkdir(parents=True)
    for n in ("a", "b", "c"):
        (d / f"{n}.mp4").write_bytes(n.encode() * 10)
    st = {"material_root": str(root), "archive_root": str(tmp / "mv_dst" / "{month}月")}
    fh = open(d / "b.mp4", "rb")
    try:
        try:
            F.move_date(st, "10.15")
            err = ""
        except fsops.MoveError as e:
            err = str(e)
    finally:
        fh.close()
    dst = tmp / "mv_dst" / "10月" / "10.15" / "小天地" / "张三"
    both = [n for n in os.listdir(d) if (dst / n).exists()] if d.exists() else []
    check(err and "b.mp4" in err and not both, f"挪一天:被占着的留在原处、不留两份 -> {err[:30]}")
    F.move_date(st, "10.15", merge=True)
    check(sorted(os.listdir(dst)) == ["a.mp4", "b.mp4", "c.mp4"] and not (root / "10.15").exists(),
          "关掉再挪:接着挪完,没有 (2)")
    try:
        F.move_date({"material_root": str(root), "archive_root": ""}, "10.15")
        err = ""
    except Exception as e:
        err = str(e)
    check("发完挪到" in err, "「发完挪到」空着:说人话")

    # 4. 设置:两份各改各的,不互相冲掉
    sp = tmp / "st_merge.json"
    a = Settings(sp)
    b = Settings(sp)
    b["api_picked"] = {"10.15": {"号": "X"}}
    b.save()
    a["last_date"] = "10.16"
    a.save()
    disk = json.loads(sp.read_text("utf-8"))
    check(disk.get("api_picked") == {"10.15": {"号": "X"}} and disk.get("last_date") == "10.16" and a["api_picked"] == b["api_picked"],
          "设置:各改各的,后保存的不会冲掉先保存的")

    # 5. 设置文件坏了:原文件留底
    bad = tmp / "st_bad.json"
    bad.write_text('{"material_root": "Q:\\\\x", ', "utf-8")
    c = Settings(bad)
    check(c.broken and os.path.exists(c.broken) and c["material_root"] == r"X:\知乎视频\小学生教育", "设置坏了:原文件改名留着,用默认")
    nul = tmp / "st_null.json"
    nul.write_text("null", "utf-8")
    check(Settings(nul).broken, "设置内容是 null:不崩")

    # 7. 路径写法不同也认得出发过
    r = Records(tmp / "a2.sqlite")

    class V:
        date, account, employee, title, cover_kind = "10.15", "号", "员", "t", ""
        path = r"X:\知乎视频\小学生教育\10.15\a.mp4"
    r.manual_published(V)
    check(r.status_of("x:/知乎视频/小学生教育/10.15/a.mp4") == "已发", "素材文件夹换个写法(x:/),发过的照样认得")

    # 6. 补记没摆成:已经记过已发的不再记,汇总不多算
    class W:
        date, account, employee, title, cover_kind = "10.15", "号", "员", "w", ""
        path = r"X:\w.mp4"
    r.failed(W, "B1", "封面没成")
    r.manual_published(W)
    r.mark_failed_published("B1")
    srow = [x for x in r.summary("10.15") if x["employee"] == "员"][0]
    check(srow["sent"] == 2, f"补记没摆成:不重复算 -> 已发 {srow['sent']}")

    # 2. 摆到一半停了:已经摆好的那几条照样等人确认
    from PySide6.QtWidgets import QApplication, QMessageBox
    QApplication.instance() or QApplication([])
    from gui.window import MainWindow
    st2 = Settings(tmp / "aw.json")
    st2["material_root"] = str(tmp / "nothing")
    st2["api_roots"] = [str(tmp / "nothing")]
    st2["ai_key"] = "test"
    rec = Records(tmp / "aw.sqlite")
    w = MainWindow(st2, rec)
    w.batch_id = "B9"
    rec.prepared(V, "B9")
    real_w = QMessageBox.warning
    QMessageBox.warning = staticmethod(lambda *a, **k: None)
    try:
        w._batch_done(None, "停下了")
    finally:
        QMessageBox.warning = real_w
    check(w.batch_id == "B9" and w.b_sent.isEnabled(), "摆到一半停了:摆好的那几条还能点「这批发好了 / 不要了」")
    w.close()


def test_ui_r1(tmp: Path):
    print("[界面审查修的:没发成说人话 / 号密钥坏了剩下的标「没轮到」/ 发完总结 / 结尾空格单独说 / 窗口能缩窄]")
    from PySide6.QtWidgets import QApplication, QPushButton
    QApplication.instance() or QApplication([])
    from qe import apiacc, checks, omapi, topic
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T

    short, tip = T.friendly(omapi.ApiError("上传视频失败:invalid_file_upload(code 40023)", "40023"))
    check("平台不收" in short and "重新导出" in tip, "40023:说人话 + 下一步")
    short, tip = T.friendly(omapi.ApiError("换 token 失败", "40001", fatal=True))
    check("密钥" in short and "登录取钥匙" in tip, "换 token 失败:说密钥不对 + 去登录取钥匙")

    class V:
        title = "A小到底该怎么控 "
        checks = {"w": 1080, "h": 1920, "tags": []}
    check(checks.problems(V) == [] and checks.format_kind(V.title) == "结尾多了空格",
          "只是结尾多空格:不算要处理(发的时候自己去掉,新员工试用 2026-10-21)")
    V.title = "标题 # 剑来"
    V.checks = {"w": 1080, "h": 1920, "tags": ["剑来"]}
    check(checks.problems(V) == ["话题写法不对"], "「# 话题」:说「话题写法不对」")

    root = tmp / "ui_r1"
    for nick in ("号A", "号B"):
        (root / nick).mkdir(parents=True)
        for n in ("1", "2", "3"):
            (root / nick / f"{nick}{n}.mp4").write_bytes(b"x")
    apiacc.save([{"nick": n, "source_name": n, "client_id": n, "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                 for n in ("号A", "号B")])
    st = Settings(tmp / "ui_st.json")
    st["api_roots"] = [str(tmp / "none")]
    st["api_archive"] = str(tmp / "ui_out")
    st["api_interval"] = 0
    st["api_picked"] = {"1.1": {n: str(root / n) for n in ("号A", "号B")}}
    tab = T.ApiTab(st, Records(tmp / "ui.sqlite"), lambda m: None)
    tab.c_date.addItem("1.1")
    tab.c_date.setCurrentText("1.1")
    tab.load()
    _settle(tab)
    for v in tab.videos:
        v.checks = {"w": 1080, "h": 1920, "tags": [], "platform": []}

    class C:
        def __init__(self, cid, sec, **k):
            self.cid = cid

        def send(self, path, title, cat, before_publish=None, **k):
            if self.cid == "号B":
                raise omapi.ApiError("换 token 失败", "40001", fatal=True)
            before_publish("V")
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}
    real = omapi.Client
    omapi.Client = C
    try:
        tab.run_publish()
        _settle(tab)
    finally:
        omapi.Client = real
    sts = {v.title: tab.state.get(i, "") for i, v in enumerate(tab.videos)}
    check(sts["号A1"] == "已发" and sts["号B1"].startswith("没发成:号的密钥") and sts["号B2"].startswith("没发:号的密钥坏了"),
          f"号密钥坏了:第一条没发成,剩下的标「没轮到」-> {sts['号B1']} / {sts['号B2']}")
    check("号B" in tab.groups and "能发 0" in tab.groups["号B"].text(3), "号那一行:没发成的不算「能发」")
    check(tab.logbox.toPlainText().count("发完了:成功 3 · 没发成 1") == 1 and "号的密钥坏了没发 2" in tab.logbox.toPlainText(),
          "发完总结:成功 / 没发成 / 没轮到都写")

    # 窗口能缩窄:接口页最小宽度不超过 1000
    tab.tree.clearSelection()
    tab.groups["号A"].setSelected(True)
    QApplication.processEvents()
    check(tab.minimumSizeHint().width() <= 1000, f"接口页选中号以后最小宽度 {tab.minimumSizeHint().width()} ≤ 1000 -> "
          + str([(x.text()[:14], x.minimumSizeHint().width()) for x in tab.hero.findChildren(QPushButton) if x.isVisibleTo(tab)]))


def test_squarish(tmp: Path):
    print("[偏方的(3:4 / 正方形)也转成标准竖屏]")
    from qe import checks, vertical

    class V:
        title = "a"
    for (w, h), want in (((1080, 1920), []), ((720, 1280), []), ((1080, 1440), ["偏方要转竖屏"]),
                         ((1080, 1080), ["偏方要转竖屏"]), ((1920, 1080), ["横屏"])):
        V.checks = {"w": w, "h": h, "landscape": vertical.is_landscape({"w": w, "h": h}),
                    "squarish": vertical.is_squarish({"w": w, "h": h}), "tags": []}
        check(checks.problems(V) == want and checks.needs_vertical(V.checks) == bool(want), f"{w}×{h} -> {want or '能发'}")
    ff = vertical.ffmpeg_path()
    if ff:
        import subprocess
        d = tmp / "sq"
        d.mkdir()
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=720x960:rate=25", "-t", "1",
                        "-c:v", "libx264", str(d / "方.mp4")], check=True)
        out, _orig = vertical.to_vertical(str(d / "方.mp4"), ff)
        info = vertical.probe(out, ff)
        check((info["w"], info["h"]) == (1080, 1920), f"3:4 转成 1080×1920 -> {info['w']}×{info['h']}")


def test_audit_r3(tmp: Path):
    print("[复审修的:GBK 设置不打不开 / 撤销遇到占用不出错 / 转竖屏撞名序号位置 / C# 不吃 / 话题间的 n / 挪文件不覆盖]")
    from qe import fsops, topic, vertical
    from qe.config import Settings

    gbk = tmp / "st_gbk.json"
    gbk.write_bytes("{'素材': 乱码}".encode("gbk"))
    st = Settings(gbk)
    check(st.broken and os.path.exists(st.broken), "设置文件是 GBK / 乱码:改名留着,照样打开")

    d = tmp / "r3_undo"
    d.mkdir()
    (d / "甲 #a #b.mp4").write_bytes(b"x")
    (d / "甲 #a #b.png").write_bytes(b"x")
    led = topic.Ledger(tmp / "r3_led.json")
    mv = topic.rename_logged(str(d / "甲 #a #b.mp4"), "甲 #a", led, "t¦整理话题¦1")
    fh = open(mv[0][1], "rb")
    try:
        ok, bad, skipped = led.undo_last()
    finally:
        fh.close()
    names = sorted(os.listdir(d))
    check(bad >= 1 and names == ["甲 #a.mp4", "甲 #a.png"] and led.last(), f"撤销遇到被占用:不出错、封面不乱改、留着下次撤 -> {names}")
    ok, bad, skipped = led.undo_last()
    check(sorted(os.listdir(d)) == ["甲 #a #b.mp4", "甲 #a #b.png"], "关掉占用再撤一次:改回来了")

    check(topic.tidy("学习C#") is None and topic.split_topics("打法 #王者荣耀n#英雄联盟")["tags"] == ["王者荣耀", "英雄联盟"],
          "「学习C#」不吃掉 #;两个话题中间粘的 n 去掉")

    ff = vertical.ffmpeg_path()
    if ff:
        import subprocess
        vd = tmp / "r3_vert"
        vd.mkdir()
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-t", "1",
                        "-c:v", "libx264", str(vd / "同名 #原神.mov")], check=True)
        (vd / "同名 #原神.mp4").write_bytes(b"keep")
        out, _o = vertical.to_vertical(str(vd / "同名 #原神.mov"), ff)
        check(os.path.basename(out) == "同名 (2) #原神.mp4", f"转竖屏撞名:序号在话题前面 -> {os.path.basename(out)}")

    a, b = tmp / "r3_a.txt", tmp / "r3_b.txt"
    a.write_bytes(b"A")
    b.write_bytes(b"B")
    try:
        fsops._replace(str(a), str(b))
        over = b.read_bytes() == b"A"
    except OSError:
        over = False
    check(not over and b.read_bytes() == b"B", "挪文件:目标已经有了就报错,不会盖掉")


def test_bughunt(tmp: Path):
    print("[挖 bug 修的:不重复发 / 不发错号 / 不发错日期 / 不卡死 / 不丢设置]")
    import json
    import time as _t
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from qe import apiacc, apisrc, fsops, omapi, topic
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T

    # 4. 「10.15已发」「10.15不能用」「10.150」不算 10.15
    check([apisrc.date_key(x) for x in ("10.15", "10.15(20)", "10.14（星期三）", "10.15已发", "10.15不能用", "10.15 脏话", "10.150")]
          == ["10.15", "10.15", "10.14", "", "", "", ""], "日期文件夹:后面跟了别的字的不算(已发 / 不能用 / 脏话)")

    # 15. 同一个号两个文件夹:都扫
    r2 = tmp / "bh_dup"
    for f in ("CSGO--号X/10.15/甲.mp4", "CS2--号X/10.15/乙.mp4"):
        (r2 / f).parent.mkdir(parents=True, exist_ok=True)
        (r2 / f).write_bytes(b"x")
    check(sorted(v.title for v in apisrc.videos([str(r2)], "10.15", {"号X"})) == ["乙", "甲"], "同一个号两个文件夹:都列出来")

    # 6. 话题正则不卡死
    t0 = _t.time()
    topic.split_topics("标题 " + " ".join(f"#t{i}" for i in range(40)) + " 来学")
    check(_t.time() - t0 < 0.2, f"40 个话题后面跟字:不卡({(_t.time() - t0) * 1000:.0f} 毫秒)")

    # 7. 设置文件带 BOM:保存不丢别的键
    sp = tmp / "bh_bom.json"
    sp.write_bytes(b"\xef\xbb\xbf" + json.dumps({"material_root": "Q:\\x", "api_interval": 9}).encode("utf-8"))
    st = Settings(sp)
    st["last_date"] = "10.16"
    st.save()
    d = json.loads(sp.read_text("utf-8-sig"))
    check(d.get("material_root") == "Q:\\x" and d.get("api_interval") == 9, "设置带 BOM:保存后别的设置还在")

    # 2. 平台回话格式怪:发布已收下就算已发,绝不记成没发成
    class R:
        def __init__(self, j):
            self.j, self.status_code = j, 200

        def json(self):
            return self.j

    class S:
        def __init__(self, pub, poll):
            self.pub, self.poll, self.pubs, self.toks = pub, poll, 0, 0

        def post(self, url, **kw):
            self.toks += 1
            return R({"code": "0", "data": {"access_token": f"T{self.toks}", "expires_in": 7200}})

        def request(self, method, url, **kw):
            path = url.split("api.om.qq.com")[1]
            if path == "/video/clientuploadready":
                return R({"code": "0", "data": {"transaction_id": "U"}})
            if path == "/video/clientuploadtrunk":
                return R({"code": "0", "data": {"end_offset": len(kw["files"]["mediatrunk"][1]) - 1}})
            if path == "/transaction/infoclient":
                if kw["params"]["transaction_id"] == "U":
                    return R({"code": "0", "data": {"vid": "V"}})
                return R(self.poll)
            if path == "/articlev2/clientpubvideo":
                self.pubs += 1
                return R(self.pub)
    f = tmp / "bh.mp4"
    f.write_bytes(b"abc")

    class FakeTime:                                  # 假时钟:查进度的「等一会儿」瞬间过去
        now = 1000.0

        @classmethod
        def time(cls):
            return cls.now

        @classmethod
        def sleep(cls, x):
            cls.now += x
    real_time = omapi.time
    omapi.time = FakeTime
    for pub, poll in (({"code": "0", "data": "weird"}, {"code": "0", "data": {}}),
                      ({"code": "0", "data": {"transaction_id": "P"}}, {"code": "0", "data": ["x"]}),
                      ({"code": "0", "data": {"transaction_id": "P"}}, None)):
        r = omapi.Client("c", "s", session=S(pub, poll)).send(str(f), "t", 1)
        check(r["status"] in ("处理中", "成功"), f"平台收下了但回话怪({str(pub)[:30]} / {str(poll)[:20]}):算已发,不记没发成")
    try:
        omapi.Client("c", "s", session=S(None, None)).send(str(f), "t", 1)
        kind = "none"
    except omapi.Unknown:
        kind = "unknown"
    except omapi.ApiError:
        kind = "apierror"
    check(kind == "unknown", f"发布回了 null:结果不明(不能自动再发)-> {kind}")

    # 12. token 中途作废:重换一次再试
    class S2(S):
        def request(self, method, url, **kw):
            if "articlev2" in url and kw["params"]["access_token"] == "T1":
                return R({"code": "40001", "msg": "access_token expired"})
            return super().request(method, url, **kw)
    s2 = S2({"code": "0", "data": {"transaction_id": "P"}}, {"code": "0", "data": {"transaction_status": "成功"}})
    r = omapi.Client("c", "s", session=s2).send(str(f), "t", 1)
    check(r["status"] == "成功" and s2.toks == 2, "token 中途作废:重换一次就发出去了")
    omapi.time = real_time

    # 1 / 3. 两个号同一个文件:记「发布中」原子地拦;发完挪走的新位置也认得出发过
    rec = Records(tmp / "bh.sqlite")

    class V:
        date, account, title = "10.15", "号A", "v1"
        path = str(tmp / "same" / "v1.mp4")
    rid = rec.try_begin(V)
    check(rid is not None and rec.try_begin(V) is None, "同一个文件:第二个号要发时拦住(不发两遍)")
    rec.api_finish(rid, "已发")
    rec.mark_moved(rid, str(tmp / "归档" / "CSGO" / "v1.mp4"))
    check(rec.blocked(str(tmp / "归档" / "CSGO" / "v1.mp4")), "发完挪走了:在归档里再看到它也认得出发过")

    # 1. 界面:两个号手选同一个文件夹 → 发布前拦住;手选文件夹不许选别的号在用的 / 归档 / 已发
    same = tmp / "bh_same"
    same.mkdir()
    for n in ("s1.mp4", "s2.mp4"):
        (same / n).write_bytes(b"x")
    apiacc.save([{"nick": n, "source_name": n, "client_id": n, "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                 for n in ("号A", "号B")])
    st = Settings(tmp / "bh_st.json")
    st["api_roots"] = [str(tmp / "none")]
    st["api_archive"] = str(tmp / "归档")
    st["api_picked"] = {"1.1": {"号A": str(same), "号B": str(same)}}
    tab = T.ApiTab(st, Records(tmp / "bh2.sqlite"), lambda m: None)
    tab.c_date.addItem("1.1")
    tab.c_date.setCurrentText("1.1")
    tab.load()
    _settle(tab)
    for v in tab.videos:
        v.checks = {"w": 1080, "h": 1920, "tags": [], "platform": []}
    sent = []

    class C:
        def __init__(self, *a, **k):
            pass

        def send(self, path, title, cat, before_publish=None, **k):
            before_publish("V")
            sent.append(title)
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}
    real = omapi.Client
    omapi.Client = C
    try:
        tab.run_publish()
        _settle(tab)
    finally:
        omapi.Client = real
    check(sent == [], "两个号手选同一个文件夹:发布前就拦住,一条都不发")
    check("已经给「号A」用了" in tab._bad_pick("1.1", "号B", str(same)) and "已发布" in tab._bad_pick("1.1", "号A", str(tmp / "归档" / "CSGO"))
          and "不发的" in tab._bad_pick("1.1", "号A", str(tmp / "x" / "10.15已发")), "手选文件夹:别的号在用的 / 归档 / 已发的都不让选")

    # 10. 挪一天时有文件没挪走:要报出来
    src = tmp / "bh_mv" / "a"
    src.mkdir(parents=True)
    (src / "x.mp4").write_bytes(b"x")
    (tmp / "bh_mv_dst" / "存在").mkdir(parents=True)        # 目标已经有(合并):一个文件一个文件挪
    real_safe = fsops.safe_move_file
    fsops.safe_move_file = lambda a, b: b                    # 模拟:挪的时候文件被改名,实际没挪走、也没报错
    try:
        n, errs = fsops.move_tree(str(tmp / "bh_mv"), str(tmp / "bh_mv_dst" / "存在"))
    finally:
        fsops.safe_move_file = real_safe
    check(errs, "挪一天:有文件没挪走就报出来,不说「挪到了」")

    # 11. 改名时封面被占着:视频照样改、照样进台账
    d = tmp / "bh_side"
    d.mkdir()
    (d / "标题 #a #b.mp4").write_bytes(b"x")
    (d / "标题 #a #b.jpg").write_bytes(b"x")
    led = topic.Ledger(tmp / "bh_led.json")
    fh = open(d / "标题 #a #b.jpg", "rb")
    try:
        mv = topic.rename_logged(str(d / "标题 #a #b.mp4"), "标题 #a", led, "t¦x¦1")
    finally:
        fh.close()
    check(os.path.basename(mv[0][1]) == "标题 #a.mp4" and led.last(), "改名时封面被占着:视频照样改、能撤销")

    # 17. 账号文件手改坏了:坏的那条跳过,不让软件打不开
    apiacc.store_path().write_text(json.dumps([{"client_id": "x"}, {"nick": "好的", "category": 1604}]), "utf-8")
    check([a["nick"] for a in apiacc.load()] == ["好的"], "账号文件里少 nick 的那条跳过")
    check(T._num("5秒", 1) == 5 and T._num(None, 3) == 3, "设置里数字写坏了:有兜底")


def test_daily_plan(tmp: Path):
    print("[每日规定条数 / 实际(按发出去的日期)/ 复制到企业微信表]")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from qe import plan
    from qe.config import Settings
    from qe.records import Records
    from gui.daily_tab import DailyTab
    st = Settings(tmp / "plan_st.json")
    check(plan.quota_of(st, "和平精英上分君") == 80 and plan.quota_of(st, "国漫解说馆z") == 50 and plan.quota_of(st, "家长小天地") == 54
          and plan.quota_of(st, "家长研习空间") == 102, "每个号每天规定条数(教育 2026-10-19 定了:小天地 54、研习空间 102)")
    check([plan.cat_total(g) for g in plan.plan(st)] == [380, 156, 200], "游戏 380 / 教育 156 / 动漫 200")
    rec = Records(tmp / "plan.sqlite")

    class V:
        date, employee, cover_kind = "10.15", "接口", ""

        def __init__(self, acc, title):
            self.account, self.title, self.path = acc, title, str(tmp / acc / f"{title}.mp4")
    day = plan.today()
    for k in range(3):
        rec.api_result(V("PUBG突围君", f"p{k}"), "已发")
    rec.api_result(V("国漫趣谈", "g1"), "已发")
    rec.api_result(V("国漫趣谈", "g2"), "没发成")
    rec.manual_published(type("E", (), {"date": "10.14", "account": "家长小天地", "employee": "张三", "title": "e1",
                                        "path": str(tmp / "e1.mp4"), "cover_kind": ""}))
    s1 = plan.summary(rec, st, day)
    check(s1["rows_rule"] == [380, 156, 200] and s1["total_rule"] == 736 and [r[2] for r in s1["rows"]] == [3, 1, 1]
          and s1["rows"][0][1] == 300 and s1["total"] == (656, 5),        # 电竞情报馆有多少发多少:这天要发的不算它的 80

          f"按发出去的日期数(素材文件夹 10.14 / 10.15 都算今天发的)-> {s1['rows']}")
    tsv = plan.as_tsv(s1).splitlines()
    check(tsv[0].split("\t") == [plan.short_day(day), "发布企鹅号视频", "游戏总发布量", "380", "3"] and tsv[3].split("\t")[2:] == ["总发布量", "736", "5"],
          "复制的文字:日期 / 工作内容 / 品类 / 规定 / 实际,和企业微信表一样")
    short = plan.short_list(s1)
    check(short[0][0] == "家长研习空间" and ("家长小天地", 54, 1) in short, f"没达标的号:差得多的排前面;教育也按号算 -> {short[:3]}")
    dt = DailyTab(st, rec)
    check("3 / 380" in dt.cards.vals["游戏"].text() and dt.t.rowCount() == 18, "每日汇总页:卡片和每个号的表")
    dt.copy_day()
    check(QApplication.clipboard().text().startswith(plan.short_day(day)), "点「复制」:剪贴板里就是那四行")


def test_login_key(tmp: Path):
    print("[登录取钥匙:从回话里读钥匙 / 挑对号 / 换上]")
    from qe import apiacc, omlogin
    data = {"response": {"code": 0}, "data": {"list": [
        {"mediaName": "CSGO战神君", "mediaId": "245", "clientID": "CID1", "clientSecret": "SEC1"},
        {"mediaName": "PUBG突围君", "mediaId": "246", "clientID": "CID2", "clientSecret": "SEC2"}]}}
    keys = omlogin.extract(data)
    check([(k["name"], k["client_id"], k["secret"]) for k in keys] == [("CSGO战神君", "CID1", "SEC1"), ("PUBG突围君", "CID2", "SEC2")],
          "回话里读出每个号的钥匙(字段名大小写不同也认)")
    check(omlogin.match(keys, "PUBG突围君")["client_id"] == "CID2" and omlogin.match(keys, "x", source_id="245")["client_id"] == "CID1"
          and omlogin.match(keys, "没有的号") is None, "按名字 / 号 ID 挑对号;挑不出来让人选")
    check("SEC1" not in omlogin.shape(data) and "clientSecret" in omlogin.shape(data), "读不到时日志里只写字段名,不写钥匙")
    apiacc.save([{"nick": "PUBG突围君", "client_id": "old", "secret": apiacc.seal("oldsec"), "category": 1604, "enabled": True}])
    check(apiacc.set_key("PUBG突围君", "CID2", "SEC2") and apiacc.find("PUBG突围君")["client_id"] == "CID2"
          and apiacc.unseal(apiacc.find("PUBG突围君")["secret"]) == "SEC2" and "SEC2" not in apiacc.store_path().read_text("utf-8"),
          "换上新钥匙:加密存")
    check("relogin" not in omlogin.OAUTH_LIST, "读钥匙不带 relogin=1(不让平台重发钥匙)")


def test_api_cover(tmp: Path):
    print("[接口封面:找同名封面 / 裁比例 / 原图不动]")
    from PIL import Image
    from qe import apicover
    d = tmp / "cover"
    d.mkdir()
    (d / "a.mp4").write_bytes(b"x")
    Image.new("RGB", (1280, 720), "red").save(d / "a.PNG")
    Image.new("RGB", (1280, 720), "blue").save(d / "a.jpg")
    check(apicover.find(str(d / "a.mp4")).endswith("a.jpg") and apicover.find(str(d / "b.mp4")) == "", "同名封面 jpg 优先;没有给空")
    check(apicover.crop_box(1280, 720, (3, 4)) == (370, 0, 910, 720) and apicover.crop_box(720, 1280, (16, 9)) == (0, 437, 720, 842),
          "从中间裁成 3:4 / 16:9")
    src = str(d / "a.jpg")
    check(apicover.prepare(src, "原样") == src and apicover.prepare(src, "不传") == "", "原样就传原图;不传给空")
    out = apicover.prepare(src, "3:4")
    with Image.open(out) as im:
        check(im.size == (540, 720) and out != src, "裁好的放临时文件夹")
    with Image.open(src) as im:
        check(im.size == (1280, 720), "原图没动")
    apicover.cleanup(out, src)
    check(not os.path.exists(out) and os.path.exists(src), "传完删临时的,原图还在")


def test_scan_any_layout(tmp: Path):
    print("[教育素材:每个人怎么放都认]")
    from qe import folders as F
    a = tmp / "scan" / "小天地"
    mk = lambda rel: ((a / rel).parent.mkdir(parents=True, exist_ok=True), (a / rel).write_bytes(b"x"))
    mk("卓金凤13/视频/卓的一条.mp4")                       # 卓金凤:视频\ + 封面\_标题.jpg
    mk("卓金凤13/封面/_卓的一条.jpg")
    mk("51/横/五一 的一条.mp4")                            # 51:原\ 横\(51 是人名,纯数字也算人)
    mk("51/横/五一 的一条.png")
    mk("51/原/五一 的一条.jpg")
    mk("梁/1/梁的一条.mp4")                                # 梁:梁/1/视频 + 封面/标题/横屏16x9.png
    _img(a / "梁/封面/梁的一条/横屏16x9.png", 320, 180)
    _img(a / "梁/封面/梁的一条/竖屏9x16.png", 180, 320)
    _img(a / "梁/封面/梁的一条/原图.png", 180, 320)
    mk("直接放 的一条.mp4")                                # 10.16:没有员工层,封面工具文件夹堆在旁边 / 套在封面\里
    mk("直接放 的一条.jpg")
    _img(a / "直接放_的一条/横屏16x9.png", 320, 180)
    mk("另一条.mp4")
    _img(a / "封面/另一条/横屏16x9.png", 320, 180)
    mk("31条（已发）/发过的.mp4")                          # 名字带「已发」的文件夹不碰
    mk("快捷方式 - 快捷方式.lnk")
    vs = {v.title: v for v in F.scan_account(str(a), "家长小天地", "10.16")}
    check(sorted(vs) == sorted(["卓的一条", "五一 的一条", "梁的一条", "直接放 的一条", "另一条"]), f"视频都翻出来,「已发」文件夹和快捷方式不算 -> {sorted(vs)}")
    check(vs["卓的一条"].employee == "卓金凤13" and vs["五一 的一条"].employee == "51" and vs["梁的一条"].employee == "梁"
          and vs["直接放 的一条"].employee == F.NO_EMP and vs["另一条"].employee == F.NO_EMP, "按第一层人名分人;直接放的算「没分人」")
    check(not vs["卓的一条"].png and vs["卓的一条"].jpg.endswith("_卓的一条.jpg"), "卓:只有带字 jpg(前面多个下划线也配得上),不当无字封面")
    check(vs["五一 的一条"].png.endswith("五一 的一条.png") and vs["五一 的一条"].jpg.endswith("五一 的一条.jpg"), "51:横里的 png 当无字封面")
    check(vs["梁的一条"].png.endswith("横屏16x9.png") and vs["梁的一条"].vpng.endswith("竖屏9x16.png"), "梁:封面工具的横屏 / 竖屏")
    check(vs["直接放 的一条"].png.endswith(os.path.join("直接放_的一条", "横屏16x9.png")) and vs["另一条"].png.endswith(os.path.join("另一条", "横屏16x9.png")),
          "10.16:封面工具文件夹在旁边 / 在封面里都配得上(空格变下划线也认)")
    check([g for g, _n in F.groups(list(vs.values()))][-1] == F.NO_EMP, "「没分人」排最后")


def test_scan_tricky(tmp: Path):
    print("[教育素材:漏不漏、多不多、分错人没有(刁钻放法)]")
    from qe import folders as F
    from qe.config import Settings
    a = tmp / "tricky" / "小天地"

    def mk(rel, data=b"x"):
        (a / rel).parent.mkdir(parents=True, exist_ok=True)
        (a / rel).write_bytes(data)
    mk("张/a/b/c/d/e/f/g/h/深.mp4")                         # 超深
    (a / "空的" / "更空").mkdir(parents=True)                # 空文件夹
    _img(a / "只有封面的人" / "只有封面.png", 64, 36)          # 只有封面没视频:不成为员工
    mk("张/大写.MP4")
    mk("张/._苹果垃圾.mp4")                                  # 苹果电脑留下的小文件,不是视频
    mk("张/Thumbs.db")
    mk("张/快捷 - 快捷方式.lnk")
    mk("张/零字节.mp4", b"")
    mk("张/表情😀🎉.mp4")
    mk("张/没脏话/干净的.mp4")                               # 脏话检测工具分出来的「没脏话」正是要发的
    mk("张/1疑似脏话/脏的.mp4")                              # 这个要跳过
    mk("张/31条（已发）/发过的.mp4")
    mk("张/31条（已发）/发过的2.mp4")
    mk("张/不重复/不重复的那条.mp4")
    mk("张/孩子不要玩手机，三招管住.mp4")                     # 封面工具按标题建的文件夹,标题里有「不要」
    _img(a / "张/封面/孩子不要玩手机，三招管住/横屏16x9.png", 320, 180)
    mk("张/孩子已发烧三天怎么办.mp4")                         # 标题里有「已发」(已发烧),里面是封面工具的图
    _img(a / "张/封面/孩子已发烧三天怎么办/横屏16x9.png", 320, 180)
    mk("10.16/李/嵌套日期.mp4")                              # 号文件夹里又套了一层日期:人还是「李」
    mk("一条放在自己标题文件夹里/一条放在自己标题文件夹里.mp4")  # <号>\<标题>\<标题>.mp4:不是人名
    mk("下载中.mp4.part")
    rep = []
    vs = {v.title: v for v in F.scan_account(str(a), "家长小天地", "10.16", rep)}
    want = ["深", "大写", "零字节", "表情😀🎉", "干净的", "不重复的那条", "孩子不要玩手机，三招管住", "孩子已发烧三天怎么办",
            "嵌套日期", "一条放在自己标题文件夹里"]
    check(sorted(vs) == sorted(want), f"该有的都有,._ / .lnk / Thumbs.db / .part / 疑似脏话 / 已发 的没有 -> {sorted(vs)}")
    check(vs["深"].employee == "张" and vs["大写"].employee == "张" and vs["干净的"].employee == "张", "超深 / 大写 .MP4 / 没脏话 里的都算张的")
    check(vs["嵌套日期"].employee == "李", f"号文件夹里套了日期:人是李 -> {vs['嵌套日期'].employee}")
    check(vs["一条放在自己标题文件夹里"].employee == F.NO_EMP, "放在自己标题文件夹里的不把标题当人名")
    check("只有封面的人" not in {v.employee for v in vs.values()} and "空的" not in {v.employee for v in vs.values()},
          "只有封面 / 空文件夹不成为员工")
    check(vs["零字节"].skip.startswith("0 字节") and not vs["深"].skip, "0 字节的列出来但写明不发")
    check(vs["孩子不要玩手机，三招管住"].png.endswith("横屏16x9.png") and vs["孩子已发烧三天怎么办"].png.endswith("横屏16x9.png"),
          "标题里有「不要 / 已发」的封面文件夹照样配得上")
    skipped = {os.path.basename(p): n for k, p, n in rep if k == "跳过"}
    check(skipped == {"1疑似脏话": 1, "31条（已发）": 2}, f"跳过的文件夹和条数报出来 -> {skipped}")
    # 路径太长(超过 260 字):文件夹打不开 → 报出来;文件夹打得开、文件太长 → 列出来写明不发
    P = "\\\\?\\"
    b = tmp / "tricky2" / "小天地"
    b.mkdir(parents=True)
    cur = str(b / "王")
    os.mkdir(cur)
    while len(cur) < 190:
        cur = os.path.join(cur, "长" * 40)
        os.mkdir(P + cur)
    cur = os.path.join(cur, "补" * (240 - len(cur) - 1))           # 文件夹正好 240 字:打得开,里面的文件超过 260
    os.mkdir(P + cur)
    open(P + os.path.join(cur, "文件名也很长" * 8 + ".mp4"), "wb").write(b"x")
    deep = os.path.join(cur, "再深一层" * 10)
    os.mkdir(P + deep)
    open(P + os.path.join(deep, "更深.mp4"), "wb").write(b"x")
    rep = []
    vs2 = F.scan_account(str(b), "家长小天地", "10.16", rep)
    long_ok = os.path.isdir(deep)                             # 这台电脑开了长路径就都读得到
    check(long_ok or (len(vs2) == 1 and "路径太长" in vs2[0].skip and any(k == "打不开" for k, _p, _n in rep)),
          f"路径太长:读不到的文件夹报「打不开」,文件写明不发 -> {[(v.title[:6], v.skip) for v in vs2]} {[k for k, _p, _n in rep]}")
    shutil.rmtree(P + str(tmp / "tricky2"), ignore_errors=True)
    # 同号同名只发一条(两个人各放一份);0 字节的那份不把好的挤掉;已经发了的那条优先留着
    mk2 = lambda p, t, e, s="": F.Video(path=p, title=t, account="家长小天地", employee=e, date="10.16", skip=s)
    x1, x2, x3 = mk2("z/同名.mp4", "同名", "卓金凤13"), mk2("y/同名.mp4", "同名 ", "卓金凤14"), mk2("w/同名.mp4", "同名", "51", "0 字节")
    check(F.mark_dups([x2, x1, x3]) == 1 and not x1.skip and "卓金凤13" in x2.skip and x3.skip == "0 字节",
          "同号同名:留卓金凤13 的,卓金凤14 的写明不发")
    y1, y2 = mk2("z/同名.mp4", "同名", "卓金凤13"), mk2("y/同名.mp4", "同名", "卓金凤14")
    F.mark_dups([y1, y2], prefer=lambda v: v is y2)
    check(not y2.skip and y1.skip, "已经发了 / 摆了的那条优先留着")
    other = F.Video(path="q/同名.mp4", title="同名", account="家长研习空间", employee="卓", date="10.16")
    check(F.mark_dups([mk2("z/同名.mp4", "同名", "卓"), other]) == 0, "不同号同名不算重复")
    # 号文件夹:套了一层同名日期、名字多了空格 / 多了字、两个号都像的、带「已发」的
    st = Settings(tmp / "s_tricky.json")
    r = tmp / "tricky_root"
    for rel in ("10.15/10.15/研习空间/梁/x.mp4", "10.16/ 家长小天地/x.mp4", "10.18/小天地 /x.mp4", "10.16/研习空间2/x.mp4",
                "10.16/研习空间和小天地/x.mp4", "10.16/小天地（已发）/x.mp4", "10.16/别的/x.mp4"):
        p = Path(P + str(r / rel))                           # 「小天地 」末尾带空格:Windows 只有加 \\?\ 才建得出来(苹果电脑能建)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"x")
    ad = F.account_dirs(str(r), "10.15", st)
    check(ad["家长研习空间"] == [str(r / "10.15" / "10.15" / "研习空间")] and not F.unknown_dirs(str(r), "10.15", st),
          f"10.15\\10.15\\研习空间:里面那层照样认 -> {ad}")
    ad = F.account_dirs(str(r), "10.16", st)
    check(ad["家长小天地"] == [str(r / "10.16" / " 家长小天地")] and ad["家长研习空间"] == [str(r / "10.16" / "研习空间2")],
          f"「 家长小天地」「研习空间2」认得出 -> {ad}")
    ad = F.account_dirs(str(r), "10.18", st)
    rep = []
    got = F.scan_account(ad["家长小天地"][0], "家长小天地", "10.18", rep) if ad["家长小天地"] else None
    check(got == [] and [k for k, _p, _n in rep] == ["打不开"] or got and len(got) == 1,
          f"「小天地 」末尾带空格:认得出是哪个号,读不了就报「打不开」(不悄悄漏) -> {ad} {rep}")
    check(sorted(F.unknown_dirs(str(r), "10.16", st)) == sorted(["研习空间和小天地", "小天地（已发）", "别的"]),
          f"两个号都像的 / 带已发的 / 不认识的不分给号,报出来 -> {F.unknown_dirs(str(r), '10.16', st)}")
    check(F.account_dirs(str(r), "9.9", st) == {}, "没有这天的文件夹:空")
    shutil.rmtree(P + str(r), ignore_errors=True)


def test_scan_cover_match(tmp: Path):
    print("[教育素材:封面配对不能配错(A 的视频配上 B 的封面)]")
    import time
    from qe import folders as F
    a = tmp / "covmatch"

    def mk(rel):
        (a / rel).parent.mkdir(parents=True, exist_ok=True)
        (a / rel).write_bytes(b"x")

    def im(rel, w=320, h=180, t=None):
        _img(a / rel, w, h)
        if t:
            os.utime(a / rel, (t, t))

    def scan(sub):
        return {os.path.relpath(v.path, a / sub): v for v in F.scan_account(str(a / sub), "号", "10.16")}
    rel = lambda p, sub: os.path.relpath(p, a / sub) if p else ""
    # ① 两个标题去掉标点后一样(「三个方法，」「三个方法。」):各配各的;只有一张时另一条不能拿
    mk("p1/人/三个方法，.mp4"); mk("p1/人/三个方法。.mp4"); im("p1/人/三个方法，.png"); im("p1/人/三个方法。.png")
    vs = scan("p1")
    check(vs[r"人\三个方法，.mp4"].png.endswith("三个方法，.png") and vs[r"人\三个方法。.mp4"].png.endswith("三个方法。.png"),
          "只差标点的两条:各配各的 png")
    mk("p2/人/三个方法，.mp4"); mk("p2/人/三个方法。.mp4"); im("p2/人/三个方法，.png")
    vs = scan("p2")
    check(vs[r"人\三个方法，.mp4"].png.endswith("三个方法，.png") and not vs[r"人\三个方法。.mp4"].png,
          f"只差标点、只有一张:另一条不给(宁可没封面) -> {rel(vs[r'人\三个方法。.mp4'].png, 'p2')}")
    mk("p3/三个方法，A.mp4"); mk("p3/三个方法。A.mp4"); im("p3/封面/三个方法，A/横屏16x9.png"); im("p3/封面/三个方法。A/横屏16x9.png")
    vs = scan("p3")
    check(rel(vs["三个方法，A.mp4"].png, "p3") == os.path.join("封面", "三个方法，A", "横屏16x9.png")
          and rel(vs["三个方法。A.mp4"].png, "p3") == os.path.join("封面", "三个方法。A", "横屏16x9.png"), "封面工具文件夹只差标点:各配各的")
    # ② 两个员工都有同名的:各用各的;乙没有封面不能拿甲的
    mk("p4/甲/视频/同一个标题.mp4"); im("p4/甲/封面/同一个标题/横屏16x9.png"); mk("p4/乙/同一个标题.mp4"); im("p4/乙/封面/同一个标题/横屏16x9.png")
    vs = scan("p4")
    check(rel(vs[r"甲\视频\同一个标题.mp4"].png, "p4").startswith("甲") and rel(vs[r"乙\同一个标题.mp4"].png, "p4").startswith("乙"),
          "甲、乙各有同名封面工具文件夹:各用各的(原来两条都配到乙的)")
    mk("p5/甲/同一个标题.mp4"); im("p5/甲/同一个标题.png"); mk("p5/乙/同一个标题.mp4")
    vs = scan("p5")
    check(vs[r"甲\同一个标题.mp4"].png and not vs[r"乙\同一个标题.mp4"].png, "乙的视频不拿甲的同名 png")
    mk("p5b/甲/视频/只此一条.mp4"); im("p5b/封面/只此一条/横屏16x9.png")
    check(scan("p5b")[r"甲\视频\只此一条.mp4"].png.endswith("横屏16x9.png"), "号\\封面\\ 里的(没分人)封面,员工的视频也能用")
    # ③ 前缀、全是表情、全角数字
    mk("p6/三个方法.mp4"); mk("p6/三个方法帮孩子.mp4"); im("p6/三个方法帮孩子.png")
    mk("p6/🔥🔥.mp4"); mk("p6/❤❤.mp4"); im("p6/❤❤.png")
    mk("p6/好习惯１２招.mp4"); im("p6/封面/好习惯12招/横屏16x9.png")
    vs = scan("p6")
    check(not vs["三个方法.mp4"].png, "一个标题是另一个的前缀:不配")
    check(not vs["🔥🔥.mp4"].png and vs["❤❤.mp4"].png.endswith("❤❤.png"), "全是表情的标题不互相配")
    check(vs["好习惯１２招.mp4"].png.endswith("横屏16x9.png"), "全角数字「１２」和「12」算一样")
    # ④ 封面工具文件夹里好几张:用最后改的;只有「修改 (1).png」没有横屏竖屏也认;带字 jpg 不当无字封面
    now = time.time()
    mk("p7/多版本.mp4")
    im("p7/封面/多版本/横屏16x9.png", t=now - 300); im("p7/封面/多版本/横屏16x9 (2).png", t=now - 200); im("p7/封面/多版本/横屏16x9 (3).png", t=now - 100)
    mk("p7/改过.mp4"); im("p7/封面/改过/横屏16x9.png", t=now - 300); im("p7/封面/改过/修改 (1).png", 400, 225, t=now - 100)
    mk("p7/只有修改.mp4"); im("p7/封面/只有修改/修改 (1).png"); im("p7/封面/只有修改/原图.png", 180, 320); mk("p7/封面/只有修改/_thumb.jpg")
    mk("p7/带字.mp4"); im("p7/封面/带字/原图.png"); im("p7/封面/带字/带字.jpg"); im("p7/封面/带字/竖屏9x16.png", 180, 320)
    mk("p7/人/修改作文的三个技巧.mp4"); im("p7/人/修改作文的三个技巧.png")
    vs = scan("p7")
    check(vs["多版本.mp4"].png.endswith("横屏16x9 (3).png"), f"横屏跑了几遍:用最新的 -> {os.path.basename(vs['多版本.mp4'].png)}")
    check(vs["改过.mp4"].png.endswith("修改 (1).png"), "他自己改的「修改 (1).png」比横屏晚:用修改的")
    check(vs["只有修改.mp4"].png.endswith("修改 (1).png"), "只有「修改 (1).png + 原图.png」(没横屏竖屏)也认(真实素材梁有 9 条)")
    check(not vs["带字.mp4"].png and vs["带字.mp4"].vpng.endswith("竖屏9x16.png"), "工具文件夹里的带字 jpg / 原图 不当无字横封面")
    check(vs[r"人\修改作文的三个技巧.mp4"].png.endswith("修改作文的三个技巧.png"), "标题以「修改」开头的同名 png 照样配")


def test_window_dups(tmp: Path):
    print("[界面:同号同名只摆一条、挪过地方的发过的不再摆]")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from gui.window import MainWindow
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records
    app = QApplication.instance() or QApplication([])
    root = tmp / "素材dup"
    a = root / "10.17" / "小天地"
    for rel in ("张/同名.mp4", "李/同名.mp4", "李/挪过的.mp4", "李/新的.mp4"):
        (a / rel).parent.mkdir(parents=True, exist_ok=True)
        (a / rel).write_bytes(b"x")
        _img(str(a / rel)[:-4] + ".png", 640, 360)
    rec = Records(tmp / "rdup.sqlite")
    rec.manual_published(F.Video(path=str(tmp / "早就挪走了" / "挪过的.mp4"), title="挪过的", account="家长小天地",
                                 employee="李", date="10.15"))
    st = Settings(tmp / "sdup.json")
    st["material_root"] = str(root)
    st["api_roots"] = [str(tmp / "api_none")]
    st["ai_key"] = "test"
    w = MainWindow(st, rec)
    w.c_date.setCurrentIndex(w.c_date.findData("10.17"))
    top = next(w.tree.topLevelItem(i) for i in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(i).text(0) == "家长小天地")
    kids = {top.child(k).data(0, Qt.ItemDataRole.UserRole + 3): top.child(k) for k in range(top.childCount())}
    check(kids["张"].text(1) == "1 条" and kids["李"].text(1).startswith("3 条") and "2 条不发" in kids["李"].text(1),
          f"树:李 3 条里 2 条不发 -> 张「{kids['张'].text(1)}」李「{kids['李'].text(1)}」")
    for it in kids.values():
        it.setCheckState(0, Qt.CheckState.Checked)
    t0 = time.time()
    while (w.cover_worker and w.cover_worker.isRunning()) and time.time() - t0 < 30:
        app.processEvents()
        time.sleep(0.05)
    app.processEvents()
    _acc, vids = w.next_batch()
    check(sorted((v.employee, v.title) for v in vids) == sorted([("张", "同名"), ("李", "新的")]),
          f"这批只摆 张的同名 + 李的新的 -> {[(v.employee, v.title) for v in vids]}")
    st4 = {w.table.item(i, 0).text() + "/" + w.table.item(i, 2).text(): w.table.item(i, 4).text() for i in range(w.table.rowCount())}
    check(st4["同名/李"].startswith("不发:和「张」") and st4["挪过的/李"].startswith("不发:这个号发过同名的") and st4["新的/李"] == "还没发",
          f"表里写清楚为什么不发 -> {st4}")
    check("另有 2 条不发" in w._nothing_left(), "没得摆时也说还有几条不发")
    w.close()


def test_window_renamed(tmp: Path):
    print("[界面:日期 / 号文件夹被人改了名(10.14 → 10.14（101）),发过的照样算已发,不用再标手动发过]")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from gui.window import MainWindow
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records, pkey
    app = QApplication.instance() or QApplication([])
    check(F.day_of("10.14（101）") == "10.14" and F.day_of("10-14") == "10.14" and F.day_of("10.145") == ""
          and F.month_of("10.14（101）") == "10", "改了名的日期文件夹也认得出哪天、几月(挪走要用月)")
    root = tmp / "素材ren"
    a = root / "10.14（101）" / "小天地（33）"
    for rel in ("李/发过的.mp4", "李/没发的.mp4", "李/别天发过的.mp4", "李/归档过的.mp4"):
        (a / rel).parent.mkdir(parents=True, exist_ok=True)
        (a / rel).write_bytes(b"x")
        _img(str(a / rel)[:-4] + ".png", 640, 360)
    rec = Records(tmp / "rren.sqlite")
    old = tmp / "素材ren" / "10.14" / "小天地" / "李"
    for t, d in (("发过的", "10.14"), ("别天发过的", "10.12"), ("归档过的", "10.14")):
        rec.manual_published(F.Video(path=str(old / f"{t}.mp4"), title=t, account="家长小天地", employee="李", date=d))
    arch = tmp / "归档ren" / "归档过的.mp4"
    arch.parent.mkdir(parents=True)
    arch.write_bytes(b"x")
    rid = next(r["id"] for r in rec.rows() if r["title"] == "归档过的")
    rec.mark_moved(rid, str(arch))
    st = Settings(tmp / "sren.json")
    st["material_root"] = str(root)
    st["api_roots"] = [str(tmp / "api_none")]
    st["ai_key"] = "test"
    w = MainWindow(st, rec)
    w.c_date.setCurrentIndex(w.c_date.findData("10.14（101）"))
    top = next(w.tree.topLevelItem(i) for i in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(i).text(0) == "家长小天地")
    li = top.child(0)
    check("发了 1" in li.text(1) and "2 条不发" in li.text(1), f"同一天改了名的算已发,别天的 / 归档过的照旧不发 -> {li.text(1)}")
    check(rec.status_of(str(a / "李" / "发过的.mp4")) == "已发", "记下了新位置,按路径就认得出")
    check(rec.status_of(str(arch)) == "已发", "已经挪到归档的那条,归档位置没被盖掉")
    li.setCheckState(0, Qt.CheckState.Checked)
    t0 = time.time()
    while (w.cover_worker and w.cover_worker.isRunning()) and time.time() - t0 < 30:
        app.processEvents()
        time.sleep(0.05)
    app.processEvents()
    _acc, vids = w.next_batch()
    check([v.title for v in vids] == ["没发的"], f"下一批只摆没发的 -> {[v.title for v in vids]}")
    n = sum(1 for r in rec.rows() if r["status"] == "已发")
    check(n == 3 and pkey(str(old / "发过的.mp4")) in {r["pkey"] for r in rec.rows()}, "没多记一笔(每日汇总不会多算)")
    w.close()


def test_at_titles(tmp: Path):
    print("[标题带 @某某 / @某平台 的不能发(16 个号算有问题,教育号跳过)]")
    from types import SimpleNamespace as NS
    from qe import checks
    check(checks.at_hits("好看的视频@腾讯游戏 #英雄联盟") == ["@腾讯游戏"] and checks.at_hits("联系＠小王 有惊喜") == ["＠小王"]
          and checks.at_hits("没有艾特的标题#英雄联盟") == [], "认得出半角 @ 和全角 ＠,没有的不乱报")
    v = NS(title="五杀集锦 @抖音小助手 #英雄联盟", checks={"w": 1080, "h": 1920, "tags": ["英雄联盟"], "platform": []}, account="CSGO战神君")
    check(any("@抖音小助手" in p and "整理标题" in p for p in checks.problems(v)), f"16 个号:带 @ 的算要处理,点「整理标题」删掉 -> {checks.problems(v)}")
    from qe import titleclean as TC
    for t, want in (("今天三角洲打大坝#443241gfsdfq123@!#!%$#fa", "今天三角洲打大坝"), ("五杀集锦 @抖音小助手 #英雄联盟", "五杀集锦 #英雄联盟"),
                    ("iPhone16ProMax评测 #数码", None), ("AK-12突击步枪改枪 #三角洲行动", None), ("S14总决赛 #LOL", None),
                    ("联系 a@b.com 吧", None), ("@某平台官方", None)):
        check(TC.clean(t) == want, f"整理标题:「{t}」→ {want!r}(实际 {TC.clean(t)!r})")
    ad = NS(title="点击下方小黄车领取同款 #好物", checks={"w": 1080, "h": 1920, "tags": ["好物"], "platform": []}, account="CSGO战神君")
    check(any("广告" in p for p in checks.problems(ad)), "标题有广告词:算有问题(视频里八成也有广告,要人工剪)")
    TC.store().set("CSGO战神君", ad.title, True)
    check(not any("广告" in p for p in checks.problems(ad)), "人点了「剪好了」就能发(按 号 + 标题 记)")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from gui.window import MainWindow
    from qe.config import Settings
    from qe.records import Records
    root = tmp / "素材at"
    d = root / "10.20" / "小天地" / "张"
    d.mkdir(parents=True)
    for t in ("正常的一条视频标题", "带艾特的视频@某平台官方"):
        (d / f"{t}.mp4").write_bytes(b"x")
        _img(d / f"{t}.png", 640, 360)
    st = Settings(tmp / "s_at.json")
    st["material_root"] = str(root)
    st["api_roots"] = [str(tmp / "api_none")]
    st["ai_key"] = "test"
    w = MainWindow(st, Records(tmp / "r_at.sqlite"))
    w.c_date.setCurrentIndex(w.c_date.findData("10.20"))
    vs = [v for vv in w._scan.values() for v in vv]
    bad = next(v for v in vs if "艾特" in v.title)
    ok = next(v for v in vs if "正常" in v.title)
    check(not bad.skip and not ok.skip, f"教育号:带 @ 的不跳过(摆的时候在罐头里删掉 @ 那段)-> {bad.skip}")
    top = next(w.tree.topLevelItem(k) for k in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(k).text(0) == "家长小天地")
    from PySide6.QtCore import Qt
    top.child(0).setCheckState(0, Qt.CheckState.Checked)
    row = next(k for k, v in enumerate(w.videos) if "艾特" in v.title)
    check("→ 带艾特的视频" in w.table.item(row, 0).text(), f"表里写着摆的时候标题会改成什么 -> {w.table.item(row, 0).text()}")
    w.close()


def test_schedule(tmp: Path):
    print("[定时发布:定好 / 到点 / 软件没开着错过了 / 别让电脑睡]")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from gui import schedule as S
    from qe.config import Settings
    st = Settings(tmp / "s_sched.json")
    said, fired = [], []
    s = S.Scheduler(st, "edu", said.append, fired.append)
    check(s.pending() is None and s.label() == "", "没定的时候什么都没有")
    import datetime as dt
    soon = (dt.datetime.strptime(S.now_str(), "%Y-%m-%d %H:%M") + dt.timedelta(hours=13, minutes=5)).strftime("%Y-%m-%d %H:%M")
    s.set(soon, "10.18")
    check(s.pending() == {"at": soon, "date": "10.18"} and "定时-edu" in S.awake_reasons() and "13 小时 5 分" in s.label(),
          f"定好了:存进设置、叫电脑别睡、写着还有多久 -> {s.label()}")
    check(Settings(tmp / "s_sched.json").get("schedules", {}).get("edu", {}).get("date") == "10.18", "重开软件还在(存在设置文件里)")
    s.tick()
    check(not fired, "没到点不发")
    sc = dict(st["schedules"])
    sc["edu"] = {"at": "2000-01-01 09:00", "date": "10.18"}
    st["schedules"] = sc
    s.tick()
    check(fired == ["10.18"] and s.pending() is None and "定时-edu" not in S.awake_reasons(), "到点了:开始发这天的,定时清掉")
    st["schedules"] = {"api": {"at": "2000-01-01 09:00", "date": "10.17"}}
    said.clear()
    s2 = S.Scheduler(st, "api", said.append, fired.append)
    check(s2.pending() is None and fired == ["10.18"] and any("没开着" in m for m in said), "到点时软件没开着:不补发,说一声、清掉")
    s.set(soon, "10.19")
    s.cancel()
    check(s.pending() is None and "定时-edu" not in S.awake_reasons(), "取消定时:清掉,电脑能睡了")


def test_prepare_retry(tmp: Path):
    print("[定时发布:没摆好的再试一次,还不行就算没摆好]")
    from types import SimpleNamespace as NS
    from qe.prepare import Runner
    r = object.__new__(Runner)
    r.log, r.stopped = (lambda m: None), (lambda: False)
    r.open_account = r.go_create = lambda *a: None
    r.upload = lambda vs: None
    r.pending_videos = lambda: []
    tries: dict[str, int] = {}

    def one(v):
        tries[v.title] = tries.get(v.title, 0) + 1
        ok = v.title == "好" or (v.title == "第二次好" and tries[v.title] == 2)
        return ok, "" if ok else "封面:尺寸过小"
    r._prepare_one = one
    vs = [NS(title=t) for t in ("好", "坏", "第二次好")]
    got = []
    res = r.prepare_batch("号", vs, on_one=lambda v, ok, why: got.append((v.title, ok)), retry=True)
    check(tries == {"好": 1, "坏": 2, "第二次好": 2}, f"没摆好的原地再试一次 -> {tries}")
    check([(v.title, ok) for v, ok, _w in res] == [("好", True), ("坏", False), ("第二次好", True)]
          and sorted(got) == sorted([("好", True), ("坏", False), ("第二次好", True)]), "每条只报一次最后的结果")
    tries.clear()
    r.prepare_batch("号", vs, retry=False)
    check(tries == {"好": 1, "坏": 1, "第二次好": 1}, "平时(人看着)不重试")
    from qe.prepare import FORBID
    check(not FORBID.search("放弃编辑"), "「放弃编辑」能点(不是发布)")
    import datetime as dt
    from qe.config import BJ
    from qe.prepare import sched_local
    check(sched_local("2026-10-18 09:00") == dt.datetime(2026, 10, 18, 9, 0) and BJ,
          f"定时框里直接填北京时间,不按电脑时钟换算(2026-10-17 换算错了一回,一定上就发了)-> {sched_local('2026-10-18 09:00')}")
    r2 = object.__new__(Runner)
    try:
        r2.schedule_all(["x"], dt.datetime(2000, 1, 1, 9, 0))
        check(False, "定在过去的不能定")
    except Exception as e:
        check("不定" in str(e), "定时时间已经过了 / 太近:不定(不然企鹅号会马上发)")
    try:
        r2._set_sched_time(dt.datetime(2026, 10, 18, 9, 3))
        check(False, "定时只能 5 分钟一档")
    except Exception as e:
        check("5 分钟" in str(e), "定时不是 5 分钟一档的直接拒绝(罐头只能选 00 / 05 … 55)")


def test_sched_edu(tmp: Path):
    print("[定时发布 · 教育号:到点勾上所有员工、一批一批发到完;没摆好的跳过;标黄没看过的不发;「全部调好了」]")
    from PySide6.QtCore import QObject, QTimer, Signal
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    import gui.window as GW
    from qe import folders as F
    from qe.adjust import Adjustments
    from qe.config import Settings
    from qe.records import Records
    root = tmp / "素材sched"
    d = root / "10.18" / "小天地" / "张"
    d.mkdir(parents=True)
    for t in ("好的一条视频", "好的二条视频", "坏的一条视频", "黄的一条视频", "黄的二条视频"):
        (d / f"{t}.mp4").write_bytes(b"x")
        _img(d / f"{t}.png", 640, 360)
    d2 = root / "10.19" / "小天地" / "李"
    d2.mkdir(parents=True)
    for t in ("定时的一条视频", "定时的二条视频", "定时坏的视频"):
        (d2 / f"{t}.mp4").write_bytes(b"x")
        _img(d2 / f"{t}.png", 640, 360)
    calls, ats = [], []

    class FakeBW(QObject):
        log = Signal(str)
        one = Signal(object, bool, str)
        finished_batch = Signal(object, str, str)

        def __init__(self, acc, vids, shots, parent=None, auto_publish=False, unattended=False, sched_at=""):
            super().__init__(parent)
            self.vids, self.unattended = vids, unattended
            calls.append(([v.title for v in vids], unattended))
            ats.append(sched_at)

        def isRunning(self):
            return False

        def stop(self):
            pass

        def start(self):
            def go():
                res = []
                for v in self.vids:
                    ok = "坏" not in v.title
                    self.one.emit(v, ok, "" if ok else "封面:尺寸过小")
                    res.append((v, ok, "" if ok else "封面:尺寸过小"))
                self.finished_batch.emit(res, "", "ok" if any(o for _v, o, _w in res) else "none")
            QTimer.singleShot(0, go)

    def fake_decide(v):
        v.cover, v.cover_kind, v.cover_why, v.plan = v.png, "png", "测试", {}
        v.warn = "黄" in v.title
    real = (GW.BatchWorker, GW.C.decide, GW.C.note_wide, QMessageBox.information)
    boxes = []
    GW.BatchWorker, GW.C.decide, GW.C.note_wide = FakeBW, fake_decide, (lambda v: None)
    QMessageBox.information = staticmethod(lambda _p, t, x, *a, **k: boxes.append(x))
    try:
        rec = Records(tmp / "r_sched.sqlite")
        Adjustments(tmp / "封面调整.json").set_ok(str(d / "黄的一条视频.mp4"), str(d / "黄的一条视频.png"))
        st = Settings(tmp / "s_sched_edu.json")
        st["material_root"] = str(root)
        st["api_roots"] = [str(tmp / "api_none")]
        st["ai_key"] = "test"
        st["batch_size"] = 2
        w = GW.MainWindow(st, rec)
        w._sched_fire("10.18")
        t0 = time.time()
        while (w._unattended or not boxes) and time.time() - t0 < 60:
            app.processEvents()
            time.sleep(0.05)
        sent = sorted(v.title for v in w.videos if rec.status_of(v.path) == "已发")
        check(all(u for _t, u in calls) and len(calls) >= 2, f"到点后自己一批一批摆(没人看着的模式)-> {calls}")
        check(sent == ["好的一条视频", "好的二条视频", "黄的一条视频"],
              f"能发的都发了(标黄但点过「全部调好了」的也发)-> {sent}")
        check(not any("黄的二条" in t for ts, _u in calls for t in ts), "还标黄、人没看过的不摆")
        check(sum(1 for ts, _u in calls for t in ts if "坏" in t) == 1, "没摆好的跳过,不反复摆")
        check(not w._unattended and boxes and "这次发了 3 条" in boxes[-1] and "坏的一条视频" in boxes[-1] and "黄的二条视频" in boxes[-1],
              f"发完写总结:发了几条、哪几条没发 -> {boxes[-1][:120] if boxes else ''}")
        from gui.schedule import awake_reasons
        check(not any(x.startswith("定时-edu") for x in awake_reasons()), "发完了电脑能睡了")
        y2 = next(v for v in w.videos if v.title == "黄的二条视频")
        w._ask = lambda *a, **k: True
        w.all_covers_ok()
        check(not y2.warn and Adjustments(tmp / "封面调整.json").is_ok(y2.path, y2.cover), "「全部调好了」:标黄的不再标黄,记在硬盘上")
        # 罐头自带的定时(用户 2026-10-17):现在就摆,每批「全部定时发布」到同一个北京时间
        calls.clear()
        ats.clear()
        boxes.clear()
        w._sched_fire("10.19", "2026-10-19 09:00")
        t0 = time.time()
        while (w._unattended or not boxes) and time.time() - t0 < 60:
            app.processEvents()
            time.sleep(0.05)
        rows = {r["title"]: r for r in rec.rows("10.19")}
        check(ats and all(a == "2026-10-19 09:00" for a in ats), f"每批都定时到同一个点 -> {ats}")
        check(rows["定时的一条视频"]["status"] == "已发" and rows["定时的一条视频"]["published_at"] == "2026-10-19 09:00:00"
              and "罐头定时发布" in rows["定时的一条视频"]["note"] and (rows.get("定时坏的视频") or {}).get("status") != "已发",
              "定好的记成已发,发布时间 = 定的那个点(每日汇总算到那天);没摆好的不算")
        check(boxes and "定好了 2 条" in boxes[-1] and "电脑可以关" in boxes[-1] and not w._sched_at, f"写总结:定好了几条、电脑能关 -> {boxes[-1][:80] if boxes else ''}")
        w.close()
    finally:
        GW.BatchWorker, GW.C.decide, GW.C.note_wide, QMessageBox.information = real


def test_sched_api(tmp: Path):
    print("[定时发布 · 16 个号:到点不弹框直接发;没有能发的就结束]")
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "api_sched"
    (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18").mkdir(parents=True)
    for f in ("s1.mp4", "s2.mp4"):
        (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18" / f).write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_api_sched.json")
    st["api_roots"] = [str(root / "游戏视频")]
    st["api_archive"] = str(tmp / "发布sched" / "{date}")
    st["api_interval"] = 0
    rec = Records(tmp / "rec_api_sched.sqlite")
    logs = []
    tab = T.ApiTab(st, rec, logs.append)
    sent = []

    class FakeClient:
        def __init__(self, *a, **k):
            pass

        def send(self, path, title, cat, before_publish=None, **k):
            if before_publish:
                before_publish("V")
            sent.append(title)
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}

    def nope(*a, **k):
        raise AssertionError("定时发布不该弹框问")
    real = (omapi.Client, T.ask, QMessageBox.question, QMessageBox.information, QMessageBox.warning)
    omapi.Client, T.ask, QMessageBox.question = FakeClient, nope, staticmethod(nope)
    QMessageBox.information = QMessageBox.warning = staticmethod(lambda *a, **k: None)
    try:
        tab._sched_fire("10.18")                    # 假视频读不出来 → 检查有问题 → 没有能发的
        t0 = time.time()
        while tab._auto and time.time() - t0 < 60:
            app.processEvents()
            time.sleep(0.05)
        check(not tab._auto and not sent and "没有能发的" in tab.logbox.toPlainText(), "到点了:先等自动检查,没有能发的就结束")
        for v in tab.videos:
            v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
        tab.k_bad.setChecked(True)                  # 定时发布不管「只看有问题的」/ 选中
        tab._auto = "等检查"
        tab._auto_poll()
        if tab.job:
            tab.job.wait()
        app.processEvents()
        check(sorted(sent) == ["s1", "s2"] and not tab._auto, f"到点不弹框直接发全部 -> {sent}")
    finally:
        omapi.Client, T.ask, QMessageBox.question, QMessageBox.information, QMessageBox.warning = real


def test_api_wm(tmp: Path):
    print("[高级检查:查水印(记在硬盘上)/ 普通检查过的只查水印 / 右键「没水印」能发]")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from qe import apiacc, checks, watermark, wmstore
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "apiwm" / "游戏视频" / "CSGO--CSGO战神君" / "10.15"
    root.mkdir(parents=True)
    for f in ("w1.mp4", "w2.mp4"):
        (root / f).write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_wm.json")
    st["api_roots"] = [str(tmp / "apiwm" / "游戏视频")]
    tab = T.ApiTab(st, Records(tmp / "rec_wm.sqlite"), lambda m: None)
    tab.c_date.setCurrentText("10.15")
    tab.load()
    _settle(tab)
    for v in tab.videos:                               # 假视频:当成竖屏的好视频
        v.checks.update(bad_read=False, landscape=False, squarish=False, w=1080, h=1920)
        tab._remember_check(v)
    calls = []

    def fake_detect(path, ff, info, stop=None, **k):
        calls.append(os.path.basename(path))
        if path.endswith("w1.mp4"):
            return {"hits": [{"label": "B站", "kind": "logo", "score": 0.9, "box": [600, 40, 60, 30],
                              "ratio": [0.83, 0.03, 0.08, 0.02], "where": "右上角", "t": 1.0}], "shot": "", "frames": 10}
        return {"hits": [], "shot": "", "frames": 10}
    real_detect, real_check = watermark.detect, checks.check_one
    normal = []
    watermark.detect = fake_detect
    checks.check_one = lambda v, *a, **k: normal.append(v.title) or real_check(v, *a, **k)
    try:
        tab.run_check(deep=True)
        _settle(tab)
        check(sorted(calls) == ["w1.mp4", "w2.mp4"] and not normal, "高级检查:普通检查过的只查水印,不再过一遍普通检查")
        i1 = next(i for i, v in enumerate(tab.videos) if v.title == "w1")
        i2 = next(i for i, v in enumerate(tab.videos) if v.title == "w2")
        check("B站水印(右上角)" in checks.problems(tab.videos[i1]) and not tab._can_send(i1) and tab._can_send(i2),
              "查到水印:标「B站水印(右上角)」,不发;没水印的能发")
        calls.clear()
        tab.load()
        _settle(tab)
        tab.run_check(deep=True)
        _settle(tab)
        i1 = next(i for i, v in enumerate(tab.videos) if v.title == "w1")
        check(not calls and not tab._can_send(i1), "关了再开 / 刷新:水印结果记在硬盘上,不再查,照样拦着")
        tab.set_wm_ok(i1, True)
        check(tab._can_send(i1) and wmstore.get(tab.videos[i1].path)["ok"], "右键「没水印 / 不用管」:能发了,也记在硬盘上")
        tab.set_wm_ok(i1, False)
        check(not tab._can_send(i1), "「其实有水印」:又不发了")
    finally:
        watermark.detect, checks.check_one = real_detect, real_check
    f = str(root / "w2.mp4")
    wmstore.put(f, {"hits": [{"label": "x"}], "shot": ""})
    wmstore.set_ok(f, True)
    (root / "w2.mp4").write_bytes(b"xy")              # 文件变了(模糊过 / 换了一条):当没查过
    check(wmstore.get(f) is None, "文件变了:水印结果作废,要重查")
    from gui.widgets import MenuButton
    b = MenuButton("检查")
    b.add("普通检查", lambda: None)
    b.add("高级检查(再查水印)", lambda: None)
    check([a.text() for a in b.menu().actions()] == ["普通检查", "高级检查(再查水印)"], "检查按钮带下拉:普通 / 高级")


def test_edu_wm(tmp: Path):
    print("[教育号查水印:查到的标黄不摆、右键「没水印」能摆、查过的不再查]")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from gui.window import MainWindow
    from qe import watermark, wmstore
    from qe.config import Settings
    from qe.records import Records
    app = QApplication.instance() or QApplication([])
    a = tmp / "素材wm" / "10.16" / "小天地"
    for rel in ("有水印.mp4", "干净的.mp4"):
        (a / rel).parent.mkdir(parents=True, exist_ok=True)
        (a / rel).write_bytes(b"x")
        _img(str(a / rel)[:-4] + ".png", 640, 360)
    st = Settings(tmp / "swm.json")
    st["material_root"] = str(tmp / "素材wm")
    st["api_roots"] = [str(tmp / "api_none")]
    st["ai_key"] = "test"
    w = MainWindow(st, Records(tmp / "rwm.sqlite"))

    def wait():
        t0 = time.time()
        while ((w.cover_worker and w.cover_worker.isRunning()) or (w.wm_worker and w.wm_worker.isRunning())) and time.time() - t0 < 30:
            app.processEvents()
            time.sleep(0.05)
        app.processEvents()
    w.c_date.setCurrentIndex(w.c_date.findData("10.16"))
    top = w.tree.topLevelItem(next(i for i in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(i).text(0) == "家长小天地"))
    top.child(0).setCheckState(0, Qt.CheckState.Checked)
    wait()
    calls = []

    def fake(path, ff, info, stop=None, **k):
        calls.append(os.path.basename(path))
        hit = [{"label": "抖音", "kind": "logo", "score": 0.9, "box": [10, 10, 40, 20], "ratio": [0.01, 0.01, 0.05, 0.02],
                "where": "左上角", "t": 1.0}] if "有水印" in path else []
        return {"hits": hit, "shot": "", "frames": 5}
    real = watermark.detect
    watermark.detect = fake
    try:
        w.run_wm()
        wait()
        check(sorted(calls) == ["干净的.mp4", "有水印.mp4"], "查水印:勾上的都查")
        i = next(k for k, v in enumerate(w.videos) if v.title == "有水印")
        check("有水印" in w.table.item(i, 4).text() and "抖音" in w.table.item(i, 4).text(), "查到的:状态写「有水印:抖音水印(左上角)」")
        _acc, vids = w.next_batch()
        check([v.title for v in vids] == ["干净的"], "摆下一批:有水印的不摆")
        w.table.selectRow(i)
        w.set_wm_ok(True)
        _acc, vids = w.next_batch()
        check(sorted(v.title for v in vids) == ["干净的", "有水印"], "右键「没水印,能摆」:能摆了")
        calls.clear()
        w._load_tree()
        wait()
        top = w.tree.topLevelItem(next(i for i in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(i).text(0) == "家长小天地"))
        if top.child(0).checkState(0) != Qt.CheckState.Checked:
            top.child(0).setCheckState(0, Qt.CheckState.Checked)
            wait()
        w.run_wm()
        wait()
        check(not calls, "查过的(记在硬盘上)不再查")
    finally:
        watermark.detect = real
    w.close()


def test_blur(tmp: Path):
    print("[模糊水印:只糊出现的那几段、原片进「水印原片」、教育号不会把原片再读进来、16 个号糊完能发]")
    import subprocess
    from qe import blur, folders as F, vertical, wmstore
    ff = vertical.ffmpeg_path()
    check(blur._enable([[0, 3], [10, 12.5]]) == ":enable='between(t,0.00,3.00)+between(t,10.00,12.50)'" and blur._enable([]) == "",
          "只在水印出现的那几段糊")
    px, py, pw, ph = blur.pixel_box([0.8, 0.02, 0.15, 0.03], 1080, 1920)
    check(px % 2 == 0 and py % 2 == 0 and px + pw <= 1080 and py + ph <= 1920 and pw > 0.15 * 1080, "框放宽一点、不出界、偶数")
    if not ff:
        print("  (没有 ffmpeg,跳过真糊)")
        return
    d = tmp / "blur" / "10.16" / "小天地"
    d.mkdir(parents=True)
    src = d / "一条.mp4"
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "testsrc=size=360x640:rate=10:duration=3",
                    "-pix_fmt", "yuv420p", str(src)], check=True)
    new, orig = blur.blur(str(src), [{"ratio": [0.6, 0.02, 0.3, 0.05], "spans": [[0, 1.5]]}], ff)
    check(os.path.normcase(new) == os.path.normcase(str(src)) and os.path.isfile(orig)
          and os.path.basename(os.path.dirname(orig)) == "水印原片" and vertical.probe(new, ff)["w"] == 360,
          "糊好的放回原名,原片进「水印原片」")
    check([v.title for v in F.scan_account(str(d), "家长小天地", "10.16")] == ["一条"], "教育号不会把「水印原片」里的原片再当成一条")
    wmstore.put(new, {"hits": [{"label": "B站", "ratio": [0.6, 0.02, 0.3, 0.05]}]}, ok=True, blurred=True)
    w = wmstore.get(new)
    check(w["blurred"] and not wmstore.blocks(w), "糊过的记下来,不再拦")


def test_ads(tmp: Path):
    print("[查口播广告 / 剪掉广告:挨着的广告句连成一段、只拦很像广告的、剪完变短、原片进「广告原片」]")
    import subprocess
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from qe import ads, apiacc, asr, checks, cut, vertical, wmstore
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    kw = ads.split_keywords(ads.DEFAULT_KEYWORDS)
    sents = [{"start": 1, "end": 3, "text": "这把枪很好用,"}, {"start": 40, "end": 42, "text": "关注我,"},
             {"start": 42.5, "end": 45, "text": "点击下方链接领取福利。"}, {"start": 60, "end": 61, "text": "限时活动开始了"}]
    g = ads.groups(sents, kw)
    check(len(g) == 1 and g[0]["start"] == 40 and g[0]["end"] == 45, "「关注我 + 点击下方链接领取」连成一段;单独一个「限时」不算")
    check(asr.sentences(["一", "个", "，", "好", "。"], [0.1, 0.2, 0.3, 0.5, 0.6], 10) ==
          [{"start": 10.1, "end": 10.5, "text": "一个，"}, {"start": 10.5, "end": 10.8, "text": "好。"}], "每个字的时间 → 一句一句几秒到几秒")
    check(cut.spans_to_cut([{"start": 5, "end": 6}, {"start": 6.2, "end": 7}], 20) == [(4.7, 7.3)], "要剪的段:前后多 0.3 秒、挨着的并起来")
    ff = vertical.ffmpeg_path()
    if ff:
        d = tmp / "adcut"
        d.mkdir()
        src = d / "广告.mp4"
        subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "testsrc=size=320x568:rate=10:duration=6",
                        "-f", "lavfi", "-i", "sine=frequency=440:duration=6", "-shortest", "-pix_fmt", "yuv420p", str(src)], check=True)
        new, orig, gone = cut.cut(str(src), [{"start": 2, "end": 3.4}], ff)
        dur = vertical.probe(new, ff)["dur"]
        check(abs(gone - 2.0) < 0.05 and 3.6 < dur < 4.4 and os.path.basename(os.path.dirname(orig)) == "广告原片",
              f"剪掉 2 秒:6 秒变 {dur:.1f} 秒,原片进「广告原片」")
    root = tmp / "apiad" / "游戏视频" / "CSGO--CSGO战神君" / "10.15"
    root.mkdir(parents=True)
    for f in ("a1.mp4", "a2.mp4"):
        (root / f).write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_ad.json")
    st["api_roots"] = [str(tmp / "apiad" / "游戏视频")]
    tab = T.ApiTab(st, Records(tmp / "rec_ad.sqlite"), lambda m: None)
    tab.c_date.setCurrentText("10.15")
    tab.load()
    _settle(tab)
    for v in tab.videos:
        v.checks.update(bad_read=False, landscape=False, squarish=False, w=1080, h=1920)
        tab._remember_check(v)

    class FakeT:
        n = 0

        def transcribe(self, path, stop=None):
            FakeT.n += 1
            return sents if path.endswith("a1.mp4") else [{"start": 1, "end": 2, "text": "今天打一把排位。"}]
    real_shared, real_dir = asr.shared, asr.model_dir
    asr.shared, asr.model_dir = (lambda ff: FakeT()), (lambda: "x")
    try:
        tab.run_ads()
        _settle(tab)
        i1 = next(i for i, v in enumerate(tab.videos) if v.title == "a1")
        i2 = next(i for i, v in enumerate(tab.videos) if v.title == "a2")
        check(any(p.startswith("口播广告 1 处(0:40") for p in checks.problems(tab.videos[i1])) and not tab._can_send(i1)
              and tab._can_send(i2), "查到口播广告:标「口播广告 1 处(0:40…)」不发;没广告的能发")
        FakeT.n = 0
        tab.load()
        _settle(tab)
        tab.run_ads()
        _settle(tab)
        check(FakeT.n == 0, "查过的(记在硬盘上)不再转文字")
        i1 = next(i for i, v in enumerate(tab.videos) if v.title == "a1")
        tab.set_ad_ok(i1, True)
        check(tab._can_send(i1) and wmstore.ADS.get(tab.videos[i1].path)["ok"], "右键「不是广告,能发」:能发了")
    finally:
        asr.shared, asr.model_dir = real_shared, real_dir


def test_cover_cache_unique(tmp: Path):
    print("[封面缓存按完整路径分:每条都叫「横屏16x9.png」也不会串]")
    from qe import cover as C
    a = C.cache_file("2.35专用封面", str(tmp / "甲" / "横屏16x9.png"), "_2.35_AI.jpg")
    b = C.cache_file("2.35专用封面", str(tmp / "乙" / "横屏16x9.png"), "_2.35_AI.jpg")
    check(a != b and a.name.startswith("横屏16x9_") and a.name.endswith("_2.35_AI.jpg"), "两条视频的同名封面:AI 扩图各存各的")
    check(C.cache_file("放大的封面", str(tmp / "甲" / "横屏16x9.png"), "_x.jpg") != C.cache_file("放大的封面", str(tmp / "乙" / "横屏16x9.png"), "_x.jpg"),
          "放大的封面也各存各的")


def test_cover_dialog_src(tmp: Path):
    print("[封面预览:显示封面的标题和位置、对不上标红、自己换封面 / 用回自动的]")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from gui.cover_dialog import CoverDialog
    from qe import cover as C
    from qe.adjust import Adjustments
    from qe.folders import Video
    d = tmp / "cdlg"
    _img(d / "封面" / "别的视频" / "横屏16x9.png", 640, 360)
    _img(d / "我选的.png", 640, 360)
    (d / "这条视频.mp4").write_bytes(b"x")
    v = Video(path=str(d / "这条视频.mp4"), title="这条视频", account="号", employee="甲", date="10.16",
              png=str(d / "封面" / "别的视频" / "横屏16x9.png"))
    adj = Adjustments(tmp / "adj_dlg.json")
    adj.use_own(v)
    C.decide(v)
    dlg = CoverDialog([v], 0, adj)
    t = dlg.lab_src.text()
    check("别的视频" in t and "对不上" in t and "横屏16x9.png" in t, "封面标题 = 封面工具的文件夹名;和视频标题对不上标红;写位置")
    dlg.set_cover(str(d / "我选的.png"))
    check(v.cover.endswith("我选的.png") and "你自己换的" in dlg.lab_src.text() and adj.own_cover(v.path).endswith("我选的.png"),
          "更换封面:用你选的,记下来")
    v2 = Video(path=v.path, title=v.title, account="号", employee="甲", date="10.16", png=str(d / "封面" / "别的视频" / "横屏16x9.png"))
    check(Adjustments(tmp / "adj_dlg.json").use_own(v2) and v2.png.endswith("我选的.png"), "关了再开:还是用你选的")
    dlg.set_cover("")
    check(v.cover.endswith("横屏16x9.png") and not adj.own_cover(v.path), "用回自动配的")
    # 用户 2026-10-19:员工忘了处理封面 → 复制封面图 / 去豆包处理 / 打开封面文件夹
    from PySide6.QtGui import QDesktopServices
    import subprocess
    opened, ran = [], []
    real = (QDesktopServices.openUrl, subprocess.Popen)
    QDesktopServices.openUrl = staticmethod(lambda u: opened.append(u.toString()) or True)
    subprocess.Popen = lambda args, *a, **k: ran.append(args)
    try:
        QApplication.clipboard().clear()
        check(dlg.copy_cover() and QApplication.clipboard().image().width() == 640, "复制封面图:剪贴板里是这张图")
        QApplication.clipboard().clear()
        dlg.to_doubao()
        check(opened == [dlg.DOUBAO] and QApplication.clipboard().image().width() == 640 and "Ctrl+V" in dlg.lab_toast.text(),
              f"去豆包处理:先复制封面,再在浏览器里打开豆包 -> {opened}")
        dlg.open_cover_dir()
        check(ran and ran[-1][0] == "explorer" and ran[-1][-1].endswith("横屏16x9.png"), f"打开封面文件夹并选中它 -> {ran}")
        dlg._src_path = ""
        QApplication.clipboard().clear()
        check(not dlg.copy_cover() and QApplication.clipboard().image().isNull(), "没有封面:复制不了,写一句")
    finally:
        QDesktopServices.openUrl, subprocess.Popen = real
    dlg.close()


def test_api_direct(tmp: Path):
    print("[接口发直连企鹅号(不走代理,快 7 倍);直连不通才走代理]")
    import requests
    from qe import omapi

    class R:
        def __init__(self, j):
            self.j, self.status_code = j, 200

        def json(self):
            return self.j

    class S:
        trust_env = True

        def __init__(self):
            self.tries = []

        def post(self, url, **kw):
            self.tries.append(self.trust_env)
            if not self.trust_env:
                raise requests.ConnectionError("直连不通")
            return R({"code": "0", "data": {"access_token": "T", "expires_in": 7200}})
    s = S()
    c = omapi.Client("a", "b", session=s)
    check(s.trust_env is False, "默认直连(不用电脑上的代理)")
    check(c.token() == "T" and s.tries == [False, True] and s.trust_env is True, "直连不通:换 token 时改走代理再试一次")


def test_api_cover_pic(tmp: Path):
    print("[接口发带封面:先传图拿地址、发布带 cover_pic;平台不收封面就不带封面再发一次]")
    from qe import omapi

    class R:
        def __init__(self, j):
            self.j, self.status_code = j, 200

        def json(self):
            return self.j

    class S:
        def __init__(self, reject_cover=False):
            self.bodies, self.reject = [], reject_cover

        def post(self, url, **kw):
            return R({"code": "0", "data": {"access_token": "T", "expires_in": 7200}})

        def request(self, method, url, **kw):
            path = url.split("api.om.qq.com")[1]
            if path == "/video/clientuploadready":
                self.size = kw["json"]["size"]
                return R({"code": "0", "data": {"transaction_id": "U1"}})
            if path == "/video/clientuploadtrunk":
                return R({"code": "0", "data": {"start_offset": self.size - 1, "end_offset": self.size - 1}})
            if path == "/transaction/infoclient":
                return R({"code": "0", "data": {"vid": "V7", "transaction_status": "成功"}})
            if path == "/articlev2/clientpubvideo":
                self.bodies.append(kw["json"])
                if self.reject and "cover_pic" in kw["json"]:
                    return R({"code": "40002", "msg": "invalid ClientPubVideoReq.CoverPic: value must be absolute"})
                return R({"code": "0", "data": {"transaction_id": "P1", "article_id": "A9"}})
            return R({"code": "-1"})
    f = tmp / "cv.mp4"
    f.write_bytes(b"0123456789")
    real_sleep = omapi.time.sleep
    omapi.time.sleep = lambda x: None
    try:
        s1 = S()
        r = omapi.Client("a", "b", session=s1).send(str(f), "t", 1604, cover="x.jpg", cover_upload=lambda img: "https://inews.gtimg.com/om_bt/abc")
        check(s1.bodies[0].get("cover_pic") == "https://inews.gtimg.com/om_bt/abc" and r["cover"] == "传了", "发布带上 cover_pic(传图拿到的地址)")
        s2 = S(reject_cover=True)
        r = omapi.Client("a", "b", session=s2).send(str(f), "t", 1604, cover="x.jpg", cover_upload=lambda img: "https://x/y")
        check(len(s2.bodies) == 2 and "cover_pic" not in s2.bodies[1] and r["article_id"] == "A9" and r["cover"].startswith("平台不收"),
              "平台不收封面:不带封面再发一次(视频照样发出去)")
        s3 = S()

        def boom(img):
            raise RuntimeError("没登录")
        r = omapi.Client("a", "b", session=s3).send(str(f), "t", 1604, cover="x.jpg", cover_upload=boom)
        check("cover_pic" not in s3.bodies[0] and r["cover"].startswith("没传上"), "封面传不上:照样发,平台自己截")
    finally:
        omapi.time.sleep = real_sleep


def test_bj_time():
    print("[北京时间]")
    import datetime as dt
    from qe.config import bj_now
    check(bj_now().utcoffset() == dt.timedelta(hours=8), "记录时间按北京时间(不管电脑时区)")


def test_guard():
    print("[不点发布]")
    from qe.prepare import PrepError, _guard
    for t in ("发布", "定时发布", "存草稿", "全部发布", "单条发布"):
        try:
            _guard(t)
            check(False, f"「{t}」应该拒绝")
        except PrepError:
            pass
    check(True, "发布 / 定时发布 / 存草稿 / 全部发布 / 单条发布 一律不点")
    for t in ("下一步", "完成", "确认", "本地上传", "添加内容自主声明"):
        _guard(t)
    check(True, "下一步 / 完成 / 确认 这些能点")
    src = (Path(__file__).resolve().parent.parent / "qe" / "prepare.py").read_text("utf-8")
    hits = [m for m in re.findall(r"click_text\(\"([^\"]+)\"", src) if re.search("发布|定时|存草稿", m)]
    check(not hits, f"代码里没有去点发布的地方 -> {hits}")


def test_window(tmp: Path):
    print("[界面]")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from gui.window import MainWindow
    from qe.config import Settings
    from qe.records import Records
    app = QApplication.instance() or QApplication([])
    root = tmp / "素材"
    e = root / "10.16" / "研习空间" / "张三"
    e.mkdir(parents=True)
    (root / "10.17").mkdir()
    for i in range(12):
        (e / f"{i:02d}.第{i}条视频标题够长.mp4").write_bytes(b"x")
        _img(e / f"{i:02d}.第{i}条视频标题够长.png", 640, 360)
        _img(e / f"{i:02d}.第{i}条视频标题够长.jpg", 270, 480)
    st = Settings(tmp / "sw.json")
    st["material_root"] = str(root)
    st["api_roots"] = [str(tmp / "api_none")]
    st["ai_key"] = "test"
    w = MainWindow(st, Records(tmp / "rw.sqlite"))
    check(w.c_date.currentData() == "" and w.c_date.count() == 3, "打开时不自动选日期,要人选")
    w.c_date.setCurrentIndex(w.c_date.findData("10.16"))
    top = w.tree.topLevelItem(0)
    check(top.text(0) == "家长研习空间" and top.childCount() == 1 and top.child(0).text(1) == "12 条", "左边:号下面员工,写着几条")
    top.child(0).setCheckState(0, Qt.CheckState.Checked)
    t0 = time.time()
    while (w.cover_worker and w.cover_worker.isRunning()) and time.time() - t0 < 30:
        app.processEvents()
        time.sleep(0.05)
    app.processEvents()
    check(w.table.rowCount() == 12, "勾上员工:中间列出 12 条")
    check(all(v.cover_kind for v in w.videos), "每条都定了封面")
    acc, vids = w.next_batch()
    check(acc == "家长研习空间" and len(vids) == 10, f"一批最多 10 条 -> {len(vids)}")
    w.rec.prepared(vids[0], "X")
    w.rec.mark_batch_published("X")
    acc, vids2 = w.next_batch()
    check(vids2[0].title != vids[0].title and len(vids2) == 10, "发过的下一批不再拿")
    w.close()


def test_window_scan(tmp: Path):
    print("[界面:号文件夹整个翻(没分人 / 有人名)、勾选三元组、缓存、刷新、发了几条]")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from gui.window import MainWindow
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records
    app = QApplication.instance() or QApplication([])
    root = tmp / "素材2"
    a = root / "10.16" / "小天地"
    for rel in ("直接放一.mp4", "直接放二.mp4", "张三/视频/张三的一.mp4"):
        (a / rel).parent.mkdir(parents=True, exist_ok=True)
        (a / rel).write_bytes(b"x")
        _img(str(a / rel)[:-4] + ".png", 640, 360)
    st = Settings(tmp / "sw2.json")
    st["material_root"] = str(root)
    st["api_roots"] = [str(tmp / "api_none")]
    st["ai_key"] = "test"
    w = MainWindow(st, Records(tmp / "rw2.sqlite"))

    def wait():
        t0 = time.time()
        while (w.cover_worker and w.cover_worker.isRunning()) and time.time() - t0 < 30:
            app.processEvents()
            time.sleep(0.05)
        app.processEvents()

    def kid(emp):
        top = next(w.tree.topLevelItem(i) for i in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(i).text(0) == "家长小天地")
        return next((top.child(k) for k in range(top.childCount()) if top.child(k).data(0, Qt.ItemDataRole.UserRole + 3) == emp), None)
    w.c_date.setCurrentIndex(w.c_date.findData("10.16"))
    zs, no = kid("张三"), kid(F.NO_EMP)
    check(zs is not None and no is not None and zs.text(1) == "1 条" and no.text(1) == "2 条", "树:张三 1 条、没分人 2 条")
    check(os.path.normcase(zs.data(0, Qt.ItemDataRole.UserRole)) == os.path.normcase(str(a / "张三"))
          and os.path.normcase(no.data(0, Qt.ItemDataRole.UserRole)) == os.path.normcase(str(a)), "打开文件夹:张三 → 张三\\,没分人 → 号文件夹")
    no.setCheckState(0, Qt.CheckState.Checked)
    wait()
    check(w.checked_dirs() == [("家长小天地", str(a), F.NO_EMP)], f"checked_dirs = (号, 号文件夹, 员工) -> {w.checked_dirs()}")
    check(w.table.rowCount() == 2 and all(v.employee == F.NO_EMP for v in w.videos), "勾没分人:只列直接放的 2 条,张三的不混进来")
    # 软件开着时员工加了一条:重新勾不重新翻(缓存),「刷新」才看得到,勾选还在
    (a / "直接放三.mp4").write_bytes(b"x")
    no.setCheckState(0, Qt.CheckState.Unchecked)
    no.setCheckState(0, Qt.CheckState.Checked)
    wait()
    check(w.table.rowCount() == 2, "重新勾:用缓存,不翻文件夹")
    w._load_dates()                                   # 「刷新」
    wait()
    no = kid(F.NO_EMP)
    check(no.checkState(0) == Qt.CheckState.Checked and w.table.rowCount() == 3 and "1 条没无字封面" in no.text(1),
          f"刷新:新加的看得到、勾着的还勾着 -> {w.table.rowCount()} 行,{no.text(1)}")
    # 标成已发:树上「发了 N」跟着变
    w.table.selectRow(0)
    w.mark_manual()
    check("发了 1" in kid(F.NO_EMP).text(1), f"标成已发后树上写「发了 1」-> {kid(F.NO_EMP).text(1)}")
    check(w.rec.rows("10.16")[0]["employee"] == F.NO_EMP, "记录里员工 = 没分人")
    # 摆之前文件被改名了(缓存里还是旧路径):挡下来,不开始摆
    import gui.window as GW
    real_bw, started = GW.BatchWorker, []
    GW.BatchWorker = lambda *a_, **k_: started.append(1) or (_ for _ in ()).throw(RuntimeError("不该开始摆"))
    try:
        acc, vids = w.next_batch()
        os.rename(vids[0].path, vids[0].path + ".bak")
        w.start_batch()
        check(not started and not w.batch_id, "这批里有文件找不到了:不开始摆,叫人刷新")
        os.rename(vids[0].path + ".bak", vids[0].path)
    finally:
        GW.BatchWorker = real_bw
    # 换日期:勾选不带过去
    w.c_date.setCurrentIndex(0)
    check(not w.checked_dirs() and w.table.rowCount() == 0, "换日期:什么都不勾")
    w.close()


def test_title_zone(tmp: Path):
    print("[2.35:1 标题区不挡脸]")
    from qe import cover as C
    # 一张 1600×900 的图,2.35:1 框高 = 1600/2.35 ≈ 681,能上下挪 219
    small = {"cx": 0.5, "cy": 0.4, "faces": [[0.45, 0.25, 0.1, 0.18]]}        # 小脸在中间偏上
    ok, ty = C.title_fit(1600, 900, small)
    x, y, w, h = C.crop_rect(1600, 900, 2.35, small)
    check(ok and y == ty and (0.25 + 0.18 * 1.03) * 900 <= y + h * C.TITLE_SAFE, "小脸:挪框就能让脸在标题区上面")
    low = {"cx": 0.5, "cy": 0.7, "faces": [[0.45, 0.55, 0.1, 0.15]]}          # 脸偏下:框挪到底
    ok, ty = C.title_fit(1600, 900, low)
    check(ok and abs(ty - (900 - 1600 / 2.35)) < 1, f"脸偏下:框挪到最底下 -> {ty:.0f}")
    # 大头照(像今天 16 号):挪框躲不开 → 切掉一点头顶(办法一)
    big = {"cx": 0.5, "cy": 0.6, "faces": [[0.38, 0.15, 0.26, 0.63, 0.40]]}
    mode, ty = C.wide_mode(1600, 900, big)
    check(mode == "crop" and ty > 0, f"大头照:切掉一点头顶 -> {mode}")
    x, y, w, h = C.crop_rect(1600, 900, 2.35, big)
    chin = (0.15 + 0.63 * 1.03) * 900
    check(y + h * C.TITLE_TEXT >= chin - 1 and y <= (0.40 - 0.63 * 0.22) * 900, "切完:下巴不碰标题字,眼睛眉毛都在")
    # 脸在图最下面:切头顶也躲不开 → AI 扩图(办法三)
    low2 = {"cx": 0.5, "cy": 0.9, "faces": [[0.4, 0.7, 0.12, 0.25, 0.8]]}
    check(C.wide_mode(1600, 900, low2)[0] == "ai" and C.needs_wide(1600, 900, low2), "脸贴着图片底下:要 AI 扩图")
    check(C.wide_mode(1600, 900, dict(low2, no_ai=True))[0] == "crop", "人点了「不用 AI」:改成切头顶")
    check(C.wide_mode(1600, 900, dict(big, ai235=True))[0] == "ai", "人点了「用 AI 扩」:就用 AI")
    check(C.wide_mode(1600, 900, dict(big, manual={C.rkey(2.35): [0.5, 0.5]}))[0] == "manual", "人手动拖过 2.35:1:按人的来")
    check(C.wide_mode(1600, 900, {"cx": 0.5, "cy": 0.5})[0] == "none", "没人:不管标题区")
    # AI 画布:脸落在标题区上面,原图该往上推就往上推
    import numpy as np
    img = np.zeros((900, 1600, 3), np.uint8)
    canvas, (left, up, w2, h2, src0) = C.ai_canvas(img, low2)
    H2, W2 = canvas.shape[:2]
    check(abs(W2 / H2 - 2.35) < 0.01, f"AI 画布就是 2.35:1 -> {W2}×{H2}")
    check(up == 0 and src0 > 0 and h2 < 900 * 1424 / 1600, "脸在最底下:原图往上推(上面切掉一块),下面空出来给 AI 补")
    old_key = C.AI_KEY
    C.AI_KEY = ""
    _img(tmp / "big.png", 1600, 900)
    try:
        C.wide_cover(str(tmp / "big.png"), low2)
        check(False, "没填 Key 要说清楚")
    except RuntimeError as e:
        check("Key" in str(e), f"没填阿里百炼 Key:说清楚 -> {e}")
    C.AI_KEY = old_key


def test_group(tmp: Path):
    print("[两个人挨着:框得下就框两个人中间]")
    from qe import cover as C
    # 2848×1600,小孩脸在 0.33,大人脸在 0.55(像用户截图那张)
    kid, mom = [0.30, 0.25, 0.07, 0.13, 0.30], [0.52, 0.12, 0.08, 0.15, 0.17]
    plan = dict(C.focus_of([mom]), faces=[mom], all=[kid, mom])
    for tag, r in (("3:4", 0.75), ("1:1", 1.0)):
        x, y, w, h = C.crop_rect(2848, 1600, r, plan)
        check(x <= (kid[0] - 0.02) * 2848 and x + w >= (mom[0] + mom[2] + 0.02) * 2848, f"{tag}:两个人都在框里")
    x, y, w, h = C.crop_rect(2848, 1600, 9 / 16, plan)
    check(x <= mom[0] * 2848 and x + w >= (mom[0] + mom[2]) * 2848, "9:16 太窄框不下两个人:对准脸大的那个")


def test_preview(tmp: Path):
    print("[封面预览 / 手动拖框]")
    from PySide6.QtCore import QPointF, Qt
    from PySide6.QtGui import QMouseEvent
    from PySide6.QtWidgets import QApplication
    from gui.cover_dialog import CoverDialog
    from qe import cover as C
    from qe.adjust import Adjustments
    from qe.folders import Video
    app = QApplication.instance() or QApplication([])
    d = tmp / "pv"
    d.mkdir()
    _img(d / "a.png", 1600, 900)
    _img(d / "a.jpg", 720, 1280)
    (d / "a.mp4").write_bytes(b"x")
    # crop_rect:手动的 > 自动的 > 正中
    check(C.crop_rect(1600, 900, 0.75, {}) == (462.5, 0.0, 675.0, 900.0), "3:4 框:图那么高,放正中")
    x, y, w, h = C.crop_rect(1600, 900, 0.75, {"cx": 0.9, "cy": 0.5, "manual": {C.rkey(0.75): [0.25, 0.5]}})
    check(abs((x + w / 2) / 1600 - 0.25) < 1e-6, "手动拖过的位置优先")
    x, y, w, h = C.crop_rect(1600, 900, 0.75, {"cx": 0.05, "cy": 0.5})
    check(x == 0.0, "拖到头就停在图边上")
    adj = Adjustments(tmp / "adj.json")
    v = Video(path=str(d / "a.mp4"), title="a", account="号", employee="员", date="1.1", png=str(d / "a.png"), jpg=str(d / "a.jpg"))
    C.decide(v)
    adj.apply(v)
    dlg = CoverDialog([v], 0, adj)
    dlg.resize(1100, 680)
    dlg.show()
    app.processEvents()
    cv = dlg.canvas
    before = cv.rect_img
    s, ox, oy = cv._geom()
    cx, cy = ox + (before[0] + before[2] / 2) * s, oy + (before[1] + before[3] / 2) * s
    for typ, pos in ((QMouseEvent.Type.MouseButtonPress, (cx, cy)), (QMouseEvent.Type.MouseMove, (cx + 120, cy)),
                     (QMouseEvent.Type.MouseButtonRelease, (cx + 120, cy))):
        ev = QMouseEvent(typ, QPointF(*pos), QPointF(*pos), Qt.MouseButton.LeftButton, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
        {QMouseEvent.Type.MouseButtonPress: cv.mousePressEvent, QMouseEvent.Type.MouseMove: cv.mouseMoveEvent,
         QMouseEvent.Type.MouseButtonRelease: cv.mouseReleaseEvent}[typ](ev)
    app.processEvents()
    check(cv.rect_img[0] > before[0] + 50, f"在大图上拖框:框往右挪了 -> {before[0]:.0f} → {cv.rect_img[0]:.0f}")
    man = Adjustments(tmp / "adj.json").manual(v.path, v.cover)
    check(C.rkey(0.75) in man and man[C.rkey(0.75)][0] > 0.5, f"拖的位置记住了(关了软件再开还在)-> {man}")
    sx, sy, ex, ey = C.box_drag({"left": 0, "top": 0, "width": 1600, "height": 900}, 1600, 900, 0.75, v.plan)
    check(abs(ex - (man[C.rkey(0.75)][0] * 1600)) < 1, "摆的时候按拖过的位置去拖(和预览一样)")
    v2 = Video(path=v.path, title="a", account="号", employee="员", date="1.1", png=v.png, jpg=v.jpg)
    C.decide(v2)
    Adjustments(tmp / "adj.json").apply(v2)
    check(C.rkey(0.75) in v2.plan.get("manual", {}), "重新打开:拖过的位置还在")
    check(v.cover.endswith(".png") and not hasattr(dlg, "b_v"), "预览里只有 png,没有「换竖封面」")
    dlg.reset()
    check(not Adjustments(tmp / "adj.json").manual(v.path, v.cover) and "manual" not in v.plan, "恢复自动")
    dlg.close()


def test_wm_detect_unit(tmp: Path):
    """查水印的算法本身(合成图,不用 X 盘):找 logo、框的位置、出现在哪几段、认字规则、比例框、一句话。"""
    print("[查水印:合成图找 logo / 认字规则 / 比例框]")
    import cv2
    import numpy as np
    from qe import watermark as W

    rng = np.random.default_rng(3)
    cw, ch = W.CANVAS_V
    rect = (0, 0, cw, ch)
    # 一个「水印」:白色细笔画拼的小图案(像 bilibili 那种字)
    mark = np.zeros((40, 110), np.uint8)
    cv2.putText(mark, "bili", (4, 32), cv2.FONT_HERSHEY_SIMPLEX, 1.1, 255, 3)

    def frame(with_mark: bool, x=560, y=60):
        f = cv2.resize(rng.integers(0, 140, (ch // 4, cw // 4), dtype=np.uint8), (cw, ch))   # 会动的画面
        if with_mark:
            reg = f[y:y + 40, x:x + 110].astype(np.float32)
            f[y:y + 40, x:x + 110] = np.clip(reg * (1 - mark / 255 * 0.8) + mark * 0.8, 0, 255).astype(np.uint8)
        return f

    ys, xs = np.where(mark > 0)
    gx, gy, gw, gh = 560 + xs.min(), 60 + ys.min(), xs.max() - xs.min() + 1, ys.max() - ys.min() + 1   # 真实位置
    crops = [frame(True)[60:100, 560:670] for _ in range(3)]
    lg = W.Logo("测试台", W.make_logo(crops), 0.75, "合成.png")
    # 20 秒:前 6 秒有水印(右上角),后面没有
    frames = [frame(i < 6) for i in range(20)]
    times = [float(i) for i in range(20)]
    maps = W.still_maps(frames)
    check(len(maps) == 18 and maps[0][1].shape == (ch, cw), "不动的图:窗口挪 1 帧,20 帧出 18 张")
    hits = W.match_logos(maps, [lg], times, rect)
    check(len(hits) == 1 and hits[0]["label"] == "测试台" and hits[0]["score"] >= 0.75, f"找到合成的水印:{hits[:1]}")
    if hits:
        bx, by, bw, bh = hits[0]["box"]
        check(abs(bx - gx) <= 5 and abs(by - gy) <= 5 and 0.8 < bw / gw < 1.25 and 0.8 < bh / gh < 1.25,
              f"框的位置对:{hits[0]['box']} ≈ {[int(gx), int(gy), int(gw), int(gh)]}")
        sp = W.spans_of(hits[0]["wins"], maps, times, 20.0)
        check(len(sp) == 1 and sp[0][0] <= 0.5 and 5 <= sp[0][1] <= 9, f"只在前几秒有:spans={sp}")
    clean = [frame(False) for _ in range(20)]
    check(not W.match_logos(W.still_maps(clean), [lg], times, rect), "没水印的(画面一直在动):不报")
    # 整条都有 → [[0, 时长]]
    allm = W.still_maps([frame(True) for _ in range(12)])
    h2 = W.match_logos(allm, [lg], times[:12], rect)
    check(h2 and W.spans_of(h2[0]["wins"], allm, times[:12], 12.0) == [[0.0, 12.0]], "整条都有:spans=[[0, 时长]]")
    # 认字规则:水印写法才算,夹在句子里 / 玩家名字不算
    for txt, want in (("抖音号：jiangjiangACG", "抖音"), ("抖音", "抖音"), ("快手号:12345", "快手"),
                      ("bilibili", "B站"), ("小红书号：99", "小红书"), ("得物", "得物"),
                      ("抖音熊二大王", ""), ("开局不到30秒抖音直接上线", ""), ("我们下期教学再会", ""),
                      ("快手", "快手"), ("微博视频", "微博"), ("抖音|悦策电影", "抖音"), ("三斤西瓜", ""),
                      ("直播版号：Z1904073", "")):
        check(W.text_label(txt) == want, f"认字规则:「{txt}」→「{want}」")
    # 比例框:竖画布里放横视频(上下黑边),框伸进黑边的部分切掉
    r = W.content_rect(1920, 1080)
    check(r == (0, 0, 1280, 720), f"横视频画布:{r}")
    r = (0, 437, 720, 406)
    check(W.to_ratio([360, 437, 72, 40.6], r) == [0.5, 0.0, 0.1, 0.1], "比例框:画布 → 原画面比例")
    check(W.to_ratio([700, 420, 40, 40], r)[1] == 0.0 and W.to_ratio([700, 420, 40, 40], r)[2] <= 0.0278,
          "比例框:伸到黑边 / 画外的部分切掉")
    check(W.where([600, 450, 80, 30], r) == "右上角", "位置说法:右上角")
    res = {"hits": [{"kind": "logo", "label": "B站", "where": "右上角"}, {"kind": "text", "label": "抖音", "where": "右下角"},
                    {"kind": "still", "label": "疑似水印", "where": "左边"}]}
    check(W.label_of(res) == "B站水印(右上角)、抖音字样(右下角)、疑似水印(左边)", f"一句话:{W.label_of(res)}")
    check(W.label_of({"hits": []}) == "", "没查到:空")
    # 整条 detect(合成一个小视频;要 ffmpeg)
    from qe import vertical
    ff = vertical.ffmpeg_path()
    if ff:
        p = str(tmp / "wm_synth.mp4")
        vw = cv2.VideoWriter(p, cv2.VideoWriter_fourcc(*"mp4v"), 2, (cw, ch), False)
        for i in range(2 * 16):
            vw.write(frame(i < 2 * 8))                     # 16 秒,前 8 秒有
        vw.release()
        info = vertical.probe(p, ff)
        res = W.detect(p, ff, info, use_ocr=False, logos=[lg])
        hs = res["hits"]
        check(res["frames"] >= 15 and len(hs) == 1 and hs[0]["label"] == "测试台", f"整条 detect 找到:{hs}")
        if hs:
            x = hs[0]
            check(all(k in x for k in ("label", "kind", "score", "box", "ratio", "where", "t", "spans")), "返回的字段都在")
            check(x["spans"] and x["spans"][0][0] <= 0.5 and 7 <= x["spans"][-1][1] <= 11.5, f"出现的那几段:{x['spans']}")
            check(abs(x["ratio"][0] - gx / cw) < 0.02 and abs(x["ratio"][1] - gy / ch) < 0.02, f"比例框:{x['ratio']}")
            check(res["shot"] and os.path.isfile(res["shot"]), "存了截图")


def _wait(app, cond, secs=60):
    t0 = time.time()
    while not cond() and time.time() - t0 < secs:
        app.processEvents()
        time.sleep(0.05)
    return cond()


def test_ui_g(tmp: Path):
    print("[新界面(用户 2026-10-18 选的 G 一键驾驶舱):大字 + 大按钮、问题卡片、明细展开收起、能发 / 已发筛选、「今天」页]")
    from PySide6.QtWidgets import QApplication, QHBoxLayout
    app = QApplication.instance() or QApplication([])
    from qe import apiacc
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "api_g"
    (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18").mkdir(parents=True)
    for f in ("g1.mp4", "g2.mp4", "g3.mp4"):
        (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18" / f).write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_g.json")
    st["api_roots"] = [str(root / "游戏视频")]
    st["api_archive"] = str(tmp / "发布g" / "{date}")
    tab = T.ApiTab(st, Records(tmp / "rec_g.sqlite"), lambda m: None)
    tab.c_date.setCurrentText("10.18")
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 3)
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
    tab.videos[0].checks["landscape"] = True
    tab._update_scope()
    tab._update_sum()
    check(tab.hero.big.text() == "2 条可以发了" and "1 条有点小问题" in tab.hero.sub.text(),
          f"大字写能发几条、下面写有几条小问题 -> {tab.hero.big.text()} / {tab.hero.sub.text()}")
    check(tab.card_v.title.text() == "1 条横屏 / 偏方" and tab.card_t.title.text() == "标题 / 话题都没问题",
          f"问题卡片按种类数 -> {tab.card_v.title.text()} / {tab.card_t.title.text()}")
    check(tab.b_vert in tab.card_v.buttons and tab.b_topic in tab.card_t.buttons, "卡片里就是原来的「转标准竖屏 / 整理标题」按钮")
    check(tab.b_pub.text() == "现在发布 2 条" and tab.b_pub.objectName() == "primary", "大按钮还是「现在发布 N 条」")
    bar = next(lay for lay in tab.findChildren(QHBoxLayout) if lay.indexOf(tab.b_pub) >= 0)
    check(bar.indexOf(tab.b_sched) >= 0 and bar.indexOf(tab.b_pub) - bar.indexOf(tab.b_sched) >= 3, "定时按钮还是离「现在发布」远一点")
    tab.set_detail(False, animate_=False)
    check(tab.detail.isHidden() and "看每个号的明细" in tab.b_detail.text() and tab.hero.property("open") == "false",
          "收起明细:列表藏起来、大字变大居中")
    tab.set_detail(True, animate_=False)
    check(not tab.detail.isHidden() and "收起明细" in tab.b_detail.text() and tab.hero.property("open") == "true",
          "看明细:列表出来、大字缩成一条")
    tab._seg_changed("ok")
    vis = [tab.videos[i].title for i, it in tab.items.items() if not it.isHidden()]
    check(sorted(vis) == ["g2", "g3"] and "过滤中" in tab.lab_scope.text(), f"点「能发」:只看能发的 -> {vis}")
    tab._seg_changed("bad")
    vis = [tab.videos[i].title for i, it in tab.items.items() if not it.isHidden()]
    check(vis == ["g1"] and tab.k_bad.isChecked(), f"点「要处理」= 原来的「只看有问题的」 -> {vis}")
    tab.k_bad.setChecked(False)
    check(tab.seg.key() == "all" and len([1 for it in tab.items.values() if not it.isHidden()]) == 3, "取消「只看有问题的」:回到全部")
    from gui.widgets import tone_of
    check([tone_of(x) for x in ("能发", "要处理", "没发成:平台不收", "已发", "结果不明", "")] == ["acc", "warn", "bad", "ok", "warn", ""],
          "状态胶囊的颜色:能发蓝 / 要处理橙 / 没发成红 / 已发绿")
    # 主窗口:「今天」页在最前面,下一步能跳过去
    import gui.window as GW
    st2 = Settings(tmp / "st_g2.json")
    st2["material_root"] = str(tmp / "edu_none")
    st2["api_roots"] = [str(tmp / "api_none")]
    w = GW.MainWindow(st2, Records(tmp / "rec_g2.sqlite"))
    check(w.stack.currentWidget() is w.today and w.nav.button(0).text() == "今天", "打开先看「今天」页")
    w.today.refresh(animate_=False)
    check("今天还差" in w.today.big.text() and "736" in w.today.lab_sent.text(), f"今天页写还差多少 -> {w.today.big.text()}")
    check(w.today.next_v.count() >= 1, "「下一步」至少有一条")
    w.goto("api")
    check(w.stack.currentWidget() is w.api_tab, "跳到 16 个号页")
    w.goto("edu")
    check(w.hero.big.text() == "先选要发哪天的", f"教育号页大字:先选日期 -> {w.hero.big.text()}")
    w.play_splash()
    _wait(app, lambda: False, 0.3)
    from gui.widgets import Splash
    sp = w.centralWidget().findChild(Splash)
    check(sp is not None and "欢迎使用本工作台" in sp.hi.text(), "开场动画:Hi,欢迎使用本工作台")
    sp.finish()
    w.close()


def test_pub_progress(tmp: Path):
    print("[发布进度(用户 2026-10-19 选的 B):大字写第几条、纸飞机进度条、每秒网速、正在发哪条、小字轮换、发完总结]")
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    import gui.widgets as W
    root = tmp / "api_pp"
    (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18").mkdir(parents=True)
    for f in ("p1.mp4", "p2.mp4", "p3.mp4"):
        (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18" / f).write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_pp.json")
    st["api_roots"] = [str(root / "游戏视频")]
    st["api_archive"] = str(tmp / "发布pp" / "{date}")
    st["api_interval"] = 0
    st["api_cover"] = "不传"
    tab = T.ApiTab(st, Records(tmp / "rec_pp.sqlite"), lambda m: None)
    tab.c_date.setCurrentIndex(tab.c_date.findData("10.18"))
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 3)
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
    tab._update_scope()
    seen = {}

    class Slow:
        def __init__(self, *a, **k):
            pass

        def send(self, path, title, cat, before_publish=None, on_progress=None, **k):
            if on_progress:
                on_progress(5 << 20, 10 << 20)
            time.sleep(0.4)
            if before_publish:
                before_publish("V")
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}
    real = (omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning)
    omapi.Client, T.ask = Slow, (lambda *a, **k: True)
    QMessageBox.information = QMessageBox.warning = staticmethod(lambda *a, **k: 0)
    try:
        tab.run_publish()
        _wait(app, lambda: tab._pub and tab._pub["done"] >= 1, 20)
        seen["big"] = tab.hero.big.text()
        seen["strip"] = not tab.hero._extra_w.isHidden() and not tab.b_stop.isHidden() and tab.b_pub.isHidden()
        seen["fun"] = tab.hero.sub.text()
        _wait(app, lambda: tab._pub and tab._pub.get("cur") is not None, 5)
        tab._pub_tick()
        seen["cur"] = tab.pub_strip.cur.text()
        _wait(app, lambda: not tab.busy(), 30)
        app.processEvents()
    finally:
        omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning = real
    check(seen.get("big", "").startswith("正在发布 · ") and seen["big"].endswith("/ 3"), f"发的时候大字写第几条 / 一共几条 -> {seen.get('big')}")
    check(seen.get("strip") is True, "发的时候露出进度条和「停」,灰掉的「现在发布」藏起来")
    pool = W.FUN_TIPS + W.FUN_ROAST + [f"{a}:{x}    {b}:{y}" for a, x, b, y in W.FUN_TALK]
    check(seen.get("fun") in pool, f"大字下面是小科普 / 段子 / 两个小人对话 -> {seen.get('fun')}")
    check("正在发" in seen.get("cur", "") and "CSGO战神君" in seen.get("cur", ""), f"写着正在发哪个号哪条 -> {seen.get('cur')}")
    p = tab._pub
    check(p["end"] and p["ok"] == 3 and "收工" in tab.pub_strip.cur.text() and "用时" in tab.pub_strip.eta.text(),
          f"发完:进度停下、写总结和用时 -> {tab.pub_strip.cur.text()} / {tab.pub_strip.eta.text()}")
    check(not tab.b_pub.isHidden() and tab.b_stop.isHidden(), "发完「现在发布」回来、「停」藏起来")
    lines = [W.fun_line() for _ in range(8)]
    check(len(set(lines)) == 8, "小字连着换不重复")
    tab.pub_strip.set_speed(2.345)
    check(tab.pub_strip.spd.text() == "↑ 2.3 MB/s", f"网速写成 MB/s -> {tab.pub_strip.spd.text()}")
    tab.pub_strip.set_speed(0.05)
    check(tab.pub_strip.spd.text() == "↑ 51 KB/s", f"慢的时候写 KB/s -> {tab.pub_strip.spd.text()}")


def test_awake(tmp: Path):
    print("[罐头最小化了任务不动(2026-10-19):摆的时候一直叫醒页面;最小化了悄悄恢复(测试里不动真窗口)]")
    from qe import prepare, winfix
    calls = []

    class FakePage:
        def cmd(self, method, params=None, timeout=20):
            calls.append((method, params))
            if method == "Page.setWebLifecycleState":
                raise prepare.CdpError("不支持")         # 叫不醒也不耽误
            return {}
    r = prepare.Runner.__new__(prepare.Runner)
    r.log, r.stopped, r.page, r._awake_at, r._told_min = (lambda m: None), (lambda: False), FakePage(), 0.0, False
    r.awake()
    check(("Emulation.setFocusEmulationEnabled", {"enabled": True}) in calls
          and ("Page.setWebLifecycleState", {"state": "active"}) in calls, f"叫页面别省电 -> {calls}")
    n = len(calls)
    r.awake()
    check(len(calls) == n, "3 秒内不重复叫")
    r.awake(force=True)
    check(len(calls) == n + 2, "换号的时候马上叫")
    check(winfix.unminimize() == 0, "测试里不去动真罐头的窗口")
    import threading
    r._closed = threading.Event()
    t = threading.Thread(target=r._window_guard, daemon=True)
    t.start()
    r.page = None
    r.close()
    t.join(3)
    check(not t.is_alive(), "摆完 close():看窗口的小哨兵跟着收工(平时你最小化罐头它不管)")
    check(isinstance(winfix.czgt_windows(), list), "能找罐头窗口(找不到就是空的)")


def test_yellow_manual(tmp: Path):
    print("[教育号不定时也能去掉标黄(2026-10-19):A 开始前问一次 / B 停下来能接着自动发布 / C 预览里「这张没问题」]")
    from PySide6.QtCore import QObject, QTimer, Signal
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    import gui.api_tab as A
    import gui.window as GW
    from qe.adjust import Adjustments
    from qe.config import Settings
    from qe.records import Records
    root = tmp / "素材yellow"
    for emp, ts in (("张", ("黄的一条视频", "好的一条视频")), ("李", ("黄的三条视频", "好的三条视频"))):
        d = root / "10.18" / "小天地" / emp
        d.mkdir(parents=True)
        for t in ts:
            (d / f"{t}.mp4").write_bytes(b"x")
            _img(d / f"{t}.png", 640, 360)
    calls = []

    class FakeBW(QObject):
        log = Signal(str)
        one = Signal(object, bool, str)
        finished_batch = Signal(object, str, str)

        def __init__(self, acc, vids, shots, parent=None, auto_publish=False, unattended=False, sched_at="", publish_only=False):
            super().__init__(parent)
            self.vids, self.auto, self.only = vids, auto_publish, publish_only
            calls.append(("只点发布" if publish_only else "摆", [v.title for v in vids]))

        def isRunning(self):
            return False

        def stop(self):
            pass

        def start(self):
            def go():
                if self.only:
                    self.finished_batch.emit([(v, True, "") for v in self.vids], "", "ok")
                    return
                res = []
                for v in self.vids:
                    self.one.emit(v, True, "")
                    res.append((v, True, ""))
                ylw = [v for v in self.vids if v.warn]
                pub = ("ok" if not ylw else f"有 {len(ylw)} 条封面标黄(要人看一眼),没自动发布") if self.auto else ""
                self.finished_batch.emit(res, "", pub)
            QTimer.singleShot(0, go)

    def fake_decide(v):
        v.cover, v.cover_kind, v.cover_why, v.plan = v.png, "png", "测试", {}
        v.warn = "黄" in v.title
    real = (GW.BatchWorker, GW.C.decide, GW.C.note_wide, QMessageBox.information, A.ask)
    boxes, asked = [], []
    GW.BatchWorker, GW.C.decide, GW.C.note_wide = FakeBW, fake_decide, (lambda v: None)
    QMessageBox.information = staticmethod(lambda _p, t, x, *a, **k: boxes.append(x))
    try:
        rec = Records(tmp / "r_yellow.sqlite")
        st = Settings(tmp / "s_yellow.json")
        st["material_root"] = str(root)
        st["api_roots"] = [str(tmp / "api_none")]
        st["ai_key"] = "test"
        st["batch_size"] = 2
        w = GW.MainWindow(st, rec)
        w.c_date.setCurrentIndex(w.c_date.findData("10.18"))
        from PySide6.QtCore import Qt
        top = next(w.tree.topLevelItem(k) for k in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(k).text(0) == "家长小天地")
        kids = {top.child(k).data(0, Qt.ItemDataRole.UserRole + 3): top.child(k) for k in range(top.childCount())}
        kids["张"].setCheckState(0, Qt.CheckState.Checked)
        _wait(app, lambda: not (w.cover_worker and w.cover_worker.isRunning()))
        w.k_auto.setChecked(True)
        w._ask = lambda *a, **k: True

        # A「还没看」:照旧摆,摆到有标黄的那批停下来
        def not_yet(*a, **k):
            asked.append(a[1])
            k["on_extra"]()
            return False
        A.ask = not_yet
        w.start_batch()
        _wait(app, lambda: bool(boxes), 10)
        y1 = next(v for v in w.videos if v.title == "黄的一条视频")
        check(asked and asked[0] == "摆下一批" and calls == [("摆", ["好的一条视频", "黄的一条视频"])] and y1.warn,
              f"A:勾了自动发布、有标黄 → 开始前问一次;选「还没看」照旧摆 -> {calls}")
        check(w._yellow_stop and not w.b_pubgo.isHidden() and "标黄" in w.hero.big.text()
              and "接着自动发布" in boxes[-1], f"摆好了因为标黄停下来:大字写标黄,露出「封面看过了,接着自动发布」-> {w.hero.big.text()}")
        # B:点「封面看过了,接着自动发布」→ 软件只点全部发布,记成已发,不再标黄
        A.ask = real[4]
        w.continue_publish()
        _wait(app, lambda: not w.batch_id, 10)
        st_ = {v.title: rec.status_of(v.path) for v in w.videos}
        check(calls[-1] == ("只点发布", ["好的一条视频", "黄的一条视频"]) and st_["黄的一条视频"] == "已发" and st_["好的一条视频"] == "已发"
              and not y1.warn and Adjustments(tmp / "封面调整.json").is_ok(y1.path, y1.cover) and w.b_pubgo.isHidden(),
              f"B:接着自动发布 = 只点「全部发布」、这批记成已发、封面记成看过了 -> {calls[-1]} {st_}")
        # A「都看过了」:开始前就去黄,一路发,不停
        w.k_auto.setChecked(False)
        kids["李"].setCheckState(0, Qt.CheckState.Checked)
        _wait(app, lambda: not (w.cover_worker and w.cover_worker.isRunning()))
        w.k_auto.setChecked(True)
        A.ask = lambda *a, **k: asked.append(a[1]) or True
        boxes.clear()
        w.start_batch()
        _wait(app, lambda: not w.batch_id and calls[-1][0] == "摆", 10)
        app.processEvents()
        y3 = next(v for v in w.videos if v.title == "黄的三条视频")
        check(not y3.warn and rec.status_of(y3.path) == "已发" and not w._yellow_stop,
              f"A:选「都看过了」→ 去黄、这批自动发出去,中间不停 -> {calls[-1]} {rec.status_of(y3.path)}")
        w.k_auto.setChecked(False)
        w.close()
    finally:
        GW.BatchWorker, GW.C.decide, GW.C.note_wide, QMessageBox.information, A.ask = real
    # C:预览里「这张没问题」
    from gui.cover_dialog import CoverDialog
    from qe.folders import Video
    d = tmp / "yc"
    _img(d / "这条.png", 640, 360)
    (d / "这条.mp4").write_bytes(b"x")
    v = Video(path=str(d / "这条.mp4"), title="这条", account="号", employee="甲", date="10.18", png=str(d / "这条.png"))
    v.cover, v.cover_kind, v.cover_why, v.plan, v.warn = v.png, "png", "没认出人,框放中间(看一眼)", {}, True
    adj = Adjustments(tmp / "adj_yc.json")
    dlg = CoverDialog([v], 0, adj)
    got = []
    dlg.changed.connect(got.append)
    vis0 = not dlg.b_okc.isHidden()
    dlg.mark_ok()
    check(vis0 and not v.warn and adj.is_ok(v.path, v.cover) and got == [0] and dlg.b_okc.isHidden() and "人看过" in v.cover_why,
          "C:标黄的封面在预览里点「这张没问题」→ 不再标黄、记在硬盘上")
    dlg.close()


def test_api_per_acc(tmp: Path):
    print("[同一个号一次传几条(2026-10-19):一组同时传、组间歇间隔;默认 1 = 一条一条;号的密钥坏了后面的不发]")
    import threading
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "api_pa"
    apiacc.save([{"nick": n, "source_name": g, "client_id": g, "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                 for g, n in (("CSGO", "CSGO战神君"), ("PUBG", "PUBG突围君"))])
    st = Settings(tmp / "st_pa.json")
    check(st.get("api_per_acc") == 1, "默认同一个号一次传 1 条(一条一条发)")
    st["api_roots"] = [str(root / "游戏视频")]
    st["api_archive"] = str(tmp / "发布pa" / "{date}")
    st["api_interval"] = 0
    st["api_cover"] = "不传"
    lock = threading.Lock()
    live = {"now": 0, "max": {}, "starts": []}

    class Slow:
        def __init__(self, cid, sec, **k):
            self.cid = cid

        def send(self, path, title, cat, before_publish=None, **k):
            if self.cid == "PUBG" and live.get("bad_pubg"):
                raise omapi.ApiError("换 token 失败", "40001", fatal=True)
            with lock:
                live["now"] += 1
                live["max"][self.cid] = max(live["max"].get(self.cid, 0), live["now"])
                live["starts"].append((time.time(), title))
            time.sleep(0.3)
            with lock:
                live["now"] -= 1
            if before_publish:
                before_publish("V")
            return {"vid": "V", "article_id": "A", "transaction_id": "P", "status": "成功"}
    real = (omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning)
    omapi.Client, T.ask = Slow, (lambda *a, **k: True)
    QMessageBox.information = QMessageBox.warning = staticmethod(lambda *a, **k: 0)
    try:
        def run(per, threads, only=None):
            st["api_per_acc"], st["api_threads"] = per, threads
            for g, n in (("CSGO", "CSGO战神君"), ("PUBG", "PUBG突围君")):    # 每轮一份新视频(发完会挪走)
                d = root / "游戏视频" / f"{g}--{n}" / "10.18"
                d.mkdir(parents=True, exist_ok=True)
                for k in range(7 if g == "CSGO" else 3):
                    (d / f"{g}{k}_{per}_{threads}.mp4").write_bytes(b"x")
                for old in d.glob("*.mp4"):
                    if not old.stem.endswith(f"_{per}_{threads}"):
                        old.unlink()
            rec = Records(tmp / f"rec_pa_{per}_{threads}_{bool(only)}.sqlite")
            tab = T.ApiTab(st, rec, lambda m: None)
            tab.set_day("10.18")
            _wait(app, lambda: not tab.busy() and len(tab.videos) == 10)
            for v in tab.videos:
                v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
            tab._update_scope()
            live["max"], live["starts"] = {}, []
            tab.run_publish(only=[i for i, v in enumerate(tab.videos) if only is None or v.account == only])
            _wait(app, lambda: not tab.busy(), 60)
            app.processEvents()
            return tab, rec
        tab, rec = run(1, 1, only="CSGO战神君")
        check(live["max"].get("CSGO") == 1 and sum(1 for v in tab.videos if rec.status_of(v.path) == "已发") == 7,
              f"一次传 1 条:一条一条发,7 条都发了 -> 同时最多 {live['max']}")
        tab, rec = run(3, 1, only="CSGO战神君")
        t = sorted(x for x, _ in live["starts"])
        groups = [len([x for x in t if t[g] <= x < t[g] + 0.2]) for g in (0, 3, 6)]
        check(live["max"].get("CSGO") == 3 and groups == [3, 3, 1] and sum(1 for v in tab.videos if rec.status_of(v.path) == "已发") == 7,
              f"一次传 3 条:3 + 3 + 1 一组组发,7 条都发了 -> 同时最多 {live['max']},每组 {groups}")
        live["bad_pubg"] = True
        tab, rec = run(5, 2)
        pubg = [i for i, v in enumerate(tab.videos) if v.account == "PUBG突围君"]
        sts = [tab.state.get(i, "") for i in pubg]
        check(all(x.startswith(("没发成", "没发:")) for x in sts) and all(not rec.status_of(tab.videos[i].path) for i in pubg)
              and sum(1 for v in tab.videos if rec.status_of(v.path) == "已发") == 7 and live["max"].get("CSGO", 0) <= 6,
              f"号的密钥坏了:这个号这组 + 后面的都不发,别的号照发 -> {sts}")
        live["bad_pubg"] = False
        tab, rec = run(5, 1)
        check(sum(1 for v in tab.videos if rec.status_of(v.path) == "已发") == 10 and len(live["starts"]) == 10,
              "一次传 5 条:全部一条不落、每条只发一次")
        tab.run_publish()
        _wait(app, lambda: not tab.busy(), 30)
        check(len(live["starts"]) == 10, "再点一次发布:发过的不再发(防重发照旧)")
    finally:
        omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning = real


def test_quick_review(tmp: Path):
    print("[快速过封面(2026-10-19):一屏铺开标黄的、回车 = 没问题跳下一张、剩下的都没问题、脸被切了标出来]")
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QKeyEvent
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    import gui.api_tab as A
    import gui.cover_review as R
    from qe.adjust import Adjustments
    from qe.folders import Video
    d = tmp / "qr"
    vs = []
    for k, fx in enumerate((0.45, 0.5, 0.97)):
        _img(d / f"q{k}.png", 1600, 900)
        v = Video(path=str(d / f"q{k}.mp4"), title=f"q{k}", account="家长小天地", employee="甲", date="10.19", png=str(d / f"q{k}.png"))
        v.cover, v.cover_kind, v.warn, v.cover_why = v.png, "png", True, "看一眼"
        v.plan = {"cx": 0.5, "cy": 0.35, "faces": [[fx, 0.2, 0.08, 0.2]], "all": [[fx, 0.2, 0.08, 0.2]]}
        vs.append(v)
    adj = Adjustments(tmp / "adj_qr.json")
    dlg = R.QuickReview(vs, adj)
    got = []
    dlg.changed.connect(got.append)
    _wait(app, lambda: len(dlg.cards) == 3, 10)
    check(len(dlg.cards) == 3 and dlg.cur == 0 and "剩 3" in dlg.lab_n.text(), f"一屏铺开 3 张标黄的 -> {dlg.lab_n.text()}")
    check(dlg.cards[2].cut and not dlg.cards[0].cut and "脸好像被切了" in dlg.cards[2].why.text(), "脸跑到框外的标出来「脸好像被切了」")
    dlg.keyPressEvent(QKeyEvent(QKeyEvent.Type.KeyPress, Qt.Key.Key_Return, Qt.KeyboardModifier.NoModifier))
    check(not vs[0].warn and adj.is_ok(vs[0].path, vs[0].cover) and dlg.cur == 1 and got == [vs[0]] and "看过 1" in dlg.lab_n.text(),
          "回车:这张没问题(去黄、记在硬盘上),跳到下一张")
    dlg.keyPressEvent(QKeyEvent(QKeyEvent.Type.KeyPress, Qt.Key.Key_Right, Qt.KeyboardModifier.NoModifier))
    check(dlg.cur == 2, "→:换下一张")
    real = A.ask
    A.ask = lambda *a, **k: True
    try:
        dlg.all_ok()
    finally:
        A.ask = real
    check(all(not v.warn for v in vs) and "剩 0" in dlg.lab_n.text() and "都看完了" in dlg.tip.text(), "「剩下的都没问题」:全去黄")
    dlg.close()


def test_game_guard(tmp: Path):
    print("[游戏号的标题像别的游戏(2026-10-20 和平精英传成了 PUBG):一条不发等确认;3 条以上整批先不发;关键词看不出来问 AI]")
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, checks, gamecheck
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    check(gamecheck.by_words("CSGO战神君", "今天和嫂子一起打和平精英") == ("和平精英", True), "CSGO 号里「和平精英」→ 像和平精英")
    check(gamecheck.by_words("CSGO战神君", "下班啦来整一把#英雄联盟#lol") == ("英雄联盟", True), "CSGO 号里「#英雄联盟#lol」→ 像英雄联盟")
    check(gamecheck.by_words("CSGO战神君", "cs2 急停教学,和平年代也要练") == ("", True), "提到了自己的游戏(cs2)就不算串")
    check(gamecheck.by_words("LOL玩家聚集地", "云顶之弈新阵容") == ("", True), "英雄联盟 / 云顶之弈算一家")
    check(gamecheck.by_words("LOLM峡谷日常", "lolm 打野") == ("", True) and gamecheck.hits("lolm 打野") == ["英雄联盟手游"],
          "英文词要单独出现:「lolm」不算「lol」")
    check(gamecheck.by_words("电竞情报馆", "和平精英决赛") == ("", True), "电竞情报馆(各种游戏)不查")
    check(gamecheck.by_words("PUBG突围君", "哈哈哈") == ("", False), "关键词看不出来的 → 要问 AI")
    root = tmp / "api_gg"
    apiacc.save([{"nick": n, "source_name": g, "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                 for g, n in (("CSGO", "CSGO战神君"), ("和平精英", "和平精英上分君"))])
    (root / "游戏视频" / "和平精英--和平精英上分君").mkdir(parents=True, exist_ok=True)
    st = Settings(tmp / "st_gg.json")
    st["api_roots"] = [str(root / "游戏视频")]
    st["quota_open"] = True
    st["ai_key"] = "fake"
    d = root / "游戏视频" / "CSGO--CSGO战神君" / "10.20"

    def load(names):
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
        for nm in names:
            (d / f"{nm}.mp4").write_bytes(b"x")
        tab = T.ApiTab(st, Records(tmp / "rgg.sqlite"), lambda m: None)
        tab.set_day("10.20")
        _wait(app, lambda: not tab.busy() and len(tab.videos) == len(names))
        tab.run_check()
        _wait(app, lambda: not tab.busy())
        for v in tab.videos:
            v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
        tab._update_scope()
        tab._update_sum()
        return tab

    def idx(tab, w):
        return next(i for i, v in enumerate(tab.videos) if w in v.title)
    AI_GAME_CALLS.clear()
    tab = load(["cs2 急停教学", "今天和嫂子一起打和平精英", "G镇六个苟分点", "残局一打三", "大狙甩枪", "沙二A点",
                "经济局", "道具教学", "残局二", "残局三", "残局四", "残局五"] + [f"日常{k}" for k in range(10)])
    asked = [t for _o, ts in AI_GAME_CALLS for t in ts]
    check(AI_GAME_CALLS and AI_GAME_CALLS[0][0] == "CSGO" and "G镇六个苟分点" in asked and "cs2 急停教学" not in asked
          and "今天和嫂子一起打和平精英" not in asked,
          f"检查时只把关键词看不出来的问 AI(按号一批)-> {AI_GAME_CALLS[:1]}")
    a, b, ok = idx(tab, "和平精英"), idx(tab, "苟分"), idx(tab, "急停")
    check("标题像是《和平精英》的" in checks.problems(tab.videos[a]) and not tab._can_send(a),
          f"一条像和平精英:标「标题像是《和平精英》的」,不发 -> {checks.problems(tab.videos[a])}")
    check(not checks.problems(tab.videos[b]) and tab._can_send(b) and "AI 觉得像《PUBG》,不拦" in tab.items[b].text(2),
          f"AI 说像 PUBG 的:只提示、不拦(用户 2026-10-22:AI 误判太多)-> {tab.items[b].text(2)}")
    check(not tab._blocked and tab._can_send(ok), f"22 条里 2 条(不到一成):不算整批放错,别的照常能发 -> {tab._blocked}")
    AI_GAME_CALLS.clear()
    tab.run_check()
    _wait(app, lambda: not tab.busy())
    check(not AI_GAME_CALLS, "问过的标题记着,下次检查不再问(不重复花钱)")
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
    tab.tree.clearSelection()
    tab.set_game_ok(a, True)
    check(not checks.problems(tab.videos[a]) and tab._can_send(a),
          f"右键「是这个号的,能发」→ 能发了 -> {checks.problems(tab.videos[a])} {tab._blocked}")
    tab.set_game_ok(a, False)
    check(not tab._can_send(a), "再右键「其实不是这个号的」→ 又不发了")
    tab = load(["cs2 急停教学", "今天和嫂子一起打和平精英", "超体新皮肤", "下班啦来整一把#英雄联盟#lol", "残局一打三", "大狙甩枪",
                "沙二A点", "经济局", "道具教学", "残局二"])
    nick = "CSGO战神君"
    check(tab._blocked.get(nick) == 3 and not any(tab._can_send(i) for i in range(len(tab.videos))),
          f"一个号 3 条像别的游戏 → 整批先不发 -> {tab._blocked}")
    g = tab.groups.get(nick) if hasattr(tab, "groups") else None
    gtext = g.text(2) if g is not None else ""
    check("放错" in tab.hero.sub.text() and (g is None or "可能放错了 3" in gtext),
          f"大字下面 / 号那一行写「可能整批放错了」-> {tab.hero.sub.text()[:80]} / {gtext}")
    check(tab.b_pub.text() == "现在发布 0 条" or not tab.b_pub.isEnabled(), f"发布按钮不算这个号 -> {tab.b_pub.text()}")
    tab.set_batch_ok(nick, True)
    ok = idx(tab, "急停")
    check(not tab._blocked and tab._can_send(ok) and not tab._can_send(idx(tab, "和平精英")),
          "右键号「这批确认没放错」→ 这批能发了,像别的游戏的那几条还是不发")
    tab.set_batch_ok(nick, False)
    check(tab._blocked.get(nick) == 3, "取消确认 → 又整批不发")
    # 真放错了:右键「放错了:挪到「和平精英上分君」的 10.20」→ 视频连同名封面挪过去(新员工试用 2026-10-21)
    (d / "今天和嫂子一起打和平精英 .jpg").write_bytes(b"j")
    check(tab._game_home("和平精英") and tab._game_home("和平精英")[0] == "和平精英上分君", "找得到和平精英的号和文件夹")
    tab.tree.clearSelection()
    tab.move_to_game(idx(tab, "和平精英"))
    _wait(app, lambda: not tab.busy())
    hp = root / "游戏视频" / "和平精英--和平精英上分君" / "10.20"
    got = sorted(os.listdir(hp)) if hp.exists() else []
    check("今天和嫂子一起打和平精英.mp4" in got and "今天和嫂子一起打和平精英 .jpg" in got
          and not any("嫂子" in v.title for v in tab.videos if v.account == nick),
          f"放错了的视频和封面挪到和平精英号同一天的文件夹,列表里没了 -> {got}")
    check(tab._blocked.get(nick) == 2 or not tab._blocked.get(nick), f"挪走以后少了一条像别的游戏的 -> {tab._blocked}")
    # AI 的判断要「说得出标题里哪个词」、两遍一样才算(10-21 和平精英号「3v3无限火力」「冰河禁区」被误判)
    import json as _json
    import requests
    replies = iter([
        [{"i": 1, "game": "英雄联盟", "word": ""}, {"i": 2, "game": "暗区突围", "word": "冰河"},
         {"i": 3, "game": "王者荣耀", "word": "鲁班七号"}, {"i": 4, "game": "无畏契约", "word": "卡德尔"},
         {"i": 5, "game": "英雄联盟", "word": "亚索"}],
        [{"i": 1, "game": "英雄联盟", "word": ""}, {"i": 2, "game": "暗区突围", "word": "冰河"},
         {"i": 3, "game": "王者荣耀", "word": "鲁班七号"}, {"i": 4, "game": "看不出", "word": ""},
         {"i": 5, "game": "英雄联盟", "word": "不在标题里"}]])

    class _R:
        status_code, text = 200, ""

        def __init__(self, arr):
            self.arr = arr

        def json(self):
            return {"choices": [{"message": {"content": _json.dumps(self.arr, ensure_ascii=False)}}]}
    real_post = requests.post
    requests.post = lambda *a, **k: _R(next(replies))
    try:
        ts = ["3v3无限火力最新彩蛋密码", "不是吧冰河禁区提前进", "鲁班七号一套带走", "AMR淘汰卡德尔需要多少发子弹", "亚索快乐风男"]
        got = gamecheck._ask_once("和平精英", ts, "k", "m", "https://x", 5)
        got2 = {t: "" for t in ts}
        requests.post = lambda *a, **k: _R(next(replies))
        got2 = gamecheck._ask_once("和平精英", ts, "k", "m", "https://x", 5)
    finally:
        requests.post = real_post
    both = {t: (g if g and got2.get(t) == g else "") for t, g in got.items()}
    check(not got[ts[0]] and got[ts[2]] == "王者荣耀" and got[ts[4]] == "英雄联盟",
          f"说不出是哪个词的不算,说得出的算 -> {got}")
    check({t: g for t, g in both.items() if g} == {ts[1]: "暗区突围", ts[2]: "王者荣耀"} or
          {t: g for t, g in both.items() if g} == {ts[2]: "王者荣耀"},
          f"两遍一样才算:「卡德尔」一遍说是一遍说看不出 → 不算 -> {both}")
    check(gamecheck.AI_KEY == "ai2" and gamecheck.store().ai("任何标题") is None, "旧版 AI 答案(ai)不再用,换了新问法重新问")


def test_ux_fixes(tmp: Path):
    print("[五个人试用后改的(2026-10-21):会真发的确认框默认「先不」、选中了按钮写明白、默认选今天、卡片别假装没事、角标两位数]")
    import datetime
    import inspect
    from PySide6.QtWidgets import QApplication, QInputDialog, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc
    from qe.config import Settings, bj_now
    from qe.records import Records
    import gui.api_tab as T
    import gui.window as GW
    # A1:会真发出去的确认框:回车 = 先不,「好」那个画红
    seen = []
    real_exec = QMessageBox.exec
    QMessageBox.exec = lambda self: seen.append((self.defaultButton().text(), [b.objectName() for b in self.buttons()
                                                                               if b.text() == "开始摆,并自动发布"])) or 0
    try:
        _REAL_ASK(None, "摆下一批", "x", "开始摆,并自动发布", "先不", default_no=True, danger=True)
    finally:
        QMessageBox.exec = real_exec
    check(seen and seen[0] == ("先不", ["danger"]), f"自动点发布的确认框:默认「先不」、开始按钮红 -> {seen}")
    src = inspect.getsource(GW.MainWindow)
    check('"开始摆,并自动发布" if auto else "开始摆"' in src and "default_no=auto, danger=auto" in src,
          "教育号勾了「摆好后自动点发布」:摆下一批的确认框默认「先不」(原来回车就真发)")
    # B1:今天的文件夹在就默认选今天;昨天还剩的单独说
    today = f"{bj_now().month}.{bj_now().day}"
    y = bj_now() - datetime.timedelta(days=1)
    yday = f"{y.month}.{y.day}"
    root = tmp / "api_ux" / "游戏视频" / "CSGO--CSGO战神君"
    for dd, names in ((today, ["今天一", "今天二"]), (yday, ["昨天剩一", "昨天剩二"])):
        (root / dd).mkdir(parents=True, exist_ok=True)
        for nm in names:
            (root / dd / f"{nm}.mp4").write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"),
                  "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_ux.json")
    st["api_roots"] = [str(tmp / "api_ux" / "游戏视频")]
    st["api_last_date"] = yday
    st["quota_open"] = True
    tab = T.ApiTab(st, Records(tmp / "r_ux.sqlite"), lambda m: None)
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 2)
    check(tab.day() == today, f"打开就选今天(上次选的是昨天也一样)-> {tab.day()}")
    check(tab._prev == (yday, 2) and tab.b_prev.isVisibleTo(tab) and "还剩 2 条" in tab.b_prev.text(),
          f"昨天还剩的单独写一行 -> {tab._prev} / {tab.b_prev.text()}")
    # A2:选中了几条:大按钮写「只发选中的」、旁边「✕ 取消选中」、大字下面写清楚
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
    tab._update_sum()
    tab._update_scope()
    check(tab.b_pub.text() == "现在发布 2 条" and not tab.b_unsel.isVisibleTo(tab), f"没选:照旧 -> {tab.b_pub.text()}")
    tab.items[0].setSelected(True)
    app.processEvents()
    check(tab.b_pub.text() == "只发选中的 1 条" and tab.b_pub.property("sel") and tab.b_unsel.isVisibleTo(tab)
          and "选中了 1 条" in tab.hero.sub.text() and tab.b_vert.text() == "转标准竖屏(选中 1 条)",
          f"选了 1 条:按钮写「只发选中的 1 条」换颜色,旁边能取消 -> {tab.b_pub.text()} / {tab.hero.sub.text()[:30]}")
    tab.b_unsel.click()
    check(not tab.selected_rows() and tab.b_pub.text() == "现在发布 2 条", "点「✕ 取消选中」回到全部")
    tab.items[0].setSelected(True)
    tab._seg_changed("ok")
    check(not tab.selected_rows(), "切「全部 / 要处理 / 能发…」自动取消选中")
    tab._seg_changed("all")
    tab.items[0].setSelected(True)
    tab._ask_title = lambda *a, **k: ("改过的标题", True)
    tab.edit_title(0)
    del tab._ask_title
    check(not tab.selected_rows(), "F2 改完标题自动取消选中")
    # 老手包:批量替换标题里的字(先预览、改完能撤销)
    from PySide6.QtWidgets import QDialog, QLineEdit as _QLE, QPlainTextEdit as _QPT
    real_dexec = QDialog.exec
    seen_pv = []

    def fake_dexec(dlg):
        eds = dlg.findChildren(_QLE)
        if len(eds) >= 2:
            eds[0].setText("今天")
            eds[1].setText("明天")
            seen_pv.append(dlg.findChild(_QPT).toPlainText())
            return QDialog.DialogCode.Accepted
        return real_dexec(dlg)
    QDialog.exec = fake_dexec
    try:
        tab.replace_titles()
    finally:
        QDialog.exec = real_dexec
    titles = sorted(v.title for v in tab.videos)
    check(seen_pv and "→ 明天" in seen_pv[0] and titles == ["改过的标题", "明天二"],
          f"批量替换:先预览,「今天」换成「明天」-> {titles}")
    real_ask = T.ask
    T.ask = lambda *a, **k: True
    try:
        tab.undo_topic()
    finally:
        T.ask = real_ask
    _settle(tab)
    check(sorted(v.title for v in tab.videos) == ["今天二", "改过的标题"], f"批量替换能撤销 -> {[v.title for v in tab.videos]}")
    tab._seg_changed("bad")
    check(tab.b_pub.text() == "点「全部」再发", f"看「要处理」时按钮说清楚为什么不能发 -> {tab.b_pub.text()}")
    tab._seg_changed("all")
    # D:没视频的日子:卡片不打绿勾;角标两位数不被切
    st2 = Settings(tmp / "st_ux2.json")
    st2["api_roots"] = [str(tmp / "api_ux_empty")]                # 别去读真的素材盘
    (tmp / "api_ux_empty").mkdir(exist_ok=True)
    empty = T.ApiTab(st2, Records(tmp / "r_ux2.sqlite"), lambda m: None)
    _wait(app, lambda: not empty.busy())
    check("还没查" in empty.card_v.title.text() and "还没查" in empty.card_t.title.text(),
          f"没视频 / 没检查:卡片写「还没查」,不打绿勾 -> {empty.card_v.title.text()}")
    from gui.widgets import set_badge, shell
    from PySide6.QtWidgets import QWidget
    root_w, grp, _stk = shell([("today", "今天", QWidget())], lambda: None)
    set_badge(grp, 0, "12")
    check(grp.badges[0].width() >= 24, f"角标两位数够宽 -> {grp.badges[0].width()}")


def test_packs(tmp: Path):
    print("[试用后的三包(2026-10-21):老板看数(出问题卡 / 多发变黄 / 两边各算各的批 / 摆好写号名)、新员工(小问号 / 没 Key 先改写法)、老手(点号全勾)]")
    from PySide6.QtWidgets import QApplication, QLabel, QTreeWidgetItem
    app = QApplication.instance() or QApplication([])
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    import gui.api_tab as T
    import gui.window as GW
    (tmp / "pk_mat").mkdir(exist_ok=True)
    (tmp / "pk_api").mkdir(exist_ok=True)
    st = Settings(tmp / "st_pk.json")
    st["material_root"] = str(tmp / "pk_mat")
    st["api_roots"] = [str(tmp / "pk_api")]
    st["splash"] = False
    st["daily_plan"] = [{"cat": "教育", "label": "教育总发布量", "accounts": {"家长研习空间": 2}},
                        {"cat": "游戏", "label": "游戏总发布量", "accounts": {"CSGO战神君": 5}}]
    rec = Records(tmp / "r_pk.sqlite")
    md = f"{bj_now().month}.{bj_now().day}"
    for k in range(3):                                            # 教育号发了 3 条(规定 2 条 → 多发 1)
        rec.manual_published(F.Video(path=str(tmp / f"e{k}.mp4"), title=f"e{k}", account="家长研习空间", employee="梁", date=md))
        rec.api_result(F.Video(path=str(tmp / f"a{k}.mp4"), title=f"a{k}", account="CSGO战神君", employee="", date=md), "已发")
    rid = rec.api_begin(F.Video(path=str(tmp / "bad.mp4"), title="bad", account="CSGO战神君", employee="", date=md))
    rec.api_finish(rid, "没发成", "测试")
    rec.prepared(F.Video(path=str(tmp / "r1.mp4"), title="r1", account="家长小天地", employee="51", date=md), "批-1")
    p = rec.problems_on(bj_now().strftime("%Y-%m-%d"))
    check(p == {"fail": 1, "unsure": 0, "ready": 1}, f"今天出问题:没发成 1、摆好没确认 1 -> {p}")
    w = GW.MainWindow(st, rec)
    w.today.refresh(animate_=False)
    check(w.today.prob.isVisibleTo(w.today) and "出问题 2 条" in w.today.prob_t.text() and "没发成 1" in w.today.prob_s.text(),
          f"首页「今天出问题」卡 -> {w.today.prob_t.text()} / {w.today.prob_s.text()}")
    edu = w.today.rings["教育"][2].text()
    check(edu == "多发 1 条" and w.today.rings["教育"][2].property("tone") == "warn", f"多发了:写「多发 1 条」变黄,不是绿的「够了」-> {edu}")
    eta = w.today.lab_eta.text()
    check("教育号 1 批 3 条" in eta and "16 个号 1 批 3 条" in eta, f"今天发了这几批:教育号、16 个号各算各的 -> {eta}")
    w.show_problems()
    st_ = sorted(w.t_rec.item(r, 4).text() for r in range(w.t_rec.rowCount()))
    check(w.k_rec_bad.isChecked() and len(st_) == 2 and "没发成" in st_, f"点「出问题」:数据记录只看出问题的 -> {st_}")
    w.k_rec_bad.setChecked(False)
    w.daily.refresh()
    diffs = {w.daily.t.item(r, 1).text(): w.daily.t.item(r, 4).text() for r in range(w.daily.t.rowCount())}
    check(diffs.get("家长研习空间") == "多发 1" and diffs.get("CSGO战神君") == "还差 2", f"每日汇总:多发 / 还差 -> {diffs}")
    # 摆好的那批:软件重开以后也写号名
    w.batch_id, w.batch_videos, w._batch_acc = "批-1", [], "家长小天地"
    w.today._fill_next(False)
    texts = [x.text() for x in w.today.findChildren(QLabel) if "摆好了" in x.text()]
    check(any("「家长小天地」1 条摆好了" in x for x in texts), f"摆好的那批写号名和几条(原来是空的「」)-> {texts}")
    w.batch_id = ""
    # 小问号
    check(len(w.findChildren(QLabel, "helpDot")) >= 5, f"行话旁边有小问号 -> {len(w.findChildren(QLabel, 'helpDot'))}")
    # 老手:点号那一行 = 这个号的员工全勾 / 全不勾,只读一次
    loads = []
    real_load = w._load_videos
    w._load_videos = lambda: loads.append(1)
    top = QTreeWidgetItem(["家长研习空间", ""])
    for e in ("梁", "51"):
        it = QTreeWidgetItem([e, ""])
        it.setData(0, GW.Qt.ItemDataRole.UserRole + 2, "d")
        it.setCheckState(0, GW.Qt.CheckState.Unchecked)
        top.addChild(it)
    w.tree.addTopLevelItem(top)
    w._top_clicked(top)
    on = [top.child(k).checkState(0) == GW.Qt.CheckState.Checked for k in range(2)]
    w._top_clicked(top)
    off = [top.child(k).checkState(0) == GW.Qt.CheckState.Unchecked for k in range(2)]
    w._load_videos = real_load
    check(all(on) and all(off) and len(loads) == 2, f"点号那一行:全勾 / 再点全不勾,一次只读一遍 -> {on} {off} {loads}")
    # 新员工:没填 AI Key 的「整理标题」:不用 AI 的先改,要 AI 的跳过
    d = tmp / "pk_api" / "CSGO--CSGO战神君" / md
    d.mkdir(parents=True, exist_ok=True)
    for nm in ("甲 #a #b", "乙 ##c"):
        (d / f"{nm}.mp4").write_bytes(b"x")
    st2 = Settings(tmp / "st_pk2.json")
    st2["api_roots"] = [str(tmp / "pk_api")]
    st2["quota_open"] = True
    st2.ai_key = lambda: ""
    tab = T.ApiTab(st2, Records(tmp / "r_pk2.sqlite"), lambda m: None)
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 2)
    seen = []
    real_ask = T.ask
    T.ask = lambda _p, _t, text, *a, **k: seen.append(text) or True
    try:
        tab.run_topic()
        _wait(app, lambda: not tab.busy())
    finally:
        T.ask = real_ask
    got = sorted(v.title for v in tab.videos)
    check(got == ["乙 #c", "甲 #a #b"] and seen and "还没填阿里百炼 Key" in seen[0],
          f"没 Key:写法不对的先改,挂好几个话题的跳过并说一声 -> {got}")


def test_clean_dirs(tmp: Path):
    print("[发完文件夹空了自动删(用户 2026-10-21):剩下的图挪进归档再删;手选的 10.2 发空了删 10.2,不动 10.1;没发完的不删]")
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "api_cl" / "游戏视频"
    acc = root / "CSGO--CSGO战神君"
    for d, names in (("10.1", ["一1", "一2"]), ("10.2", ["二1", "二2"])):
        (acc / d).mkdir(parents=True, exist_ok=True)
        for nm in names:
            (acc / d / f"{nm}.mp4").write_bytes(b"x")
            (acc / d / f"{nm}.jpg").write_bytes(b"j")
    (acc / "10.2" / "多的一张图.png").write_bytes(b"p")
    (acc / "10.2" / "Thumbs.db").write_bytes(b"t")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"),
                  "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_cl.json")
    st["api_roots"] = [str(root)]
    st["api_archive"] = str(tmp / "发布cl" / "{date}")
    st["api_interval"] = 0
    st["api_cover"] = "不传"
    st["quota_open"] = True
    st["api_picked"] = {"10.1": {"CSGO战神君": str(acc / "10.2")}}      # 这天手选的是 10.2

    class Fake:
        def __init__(self, *a, **k):
            pass

        def send(self, path, title, cat, before_publish=None, **k):
            if before_publish:
                before_publish("V")
            return {"vid": "V", "article_id": "A" + title, "transaction_id": "P", "status": "成功"}
    real = (omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning)
    omapi.Client, T.ask = Fake, (lambda *a, **k: True)
    QMessageBox.information = QMessageBox.warning = staticmethod(lambda *a, **k: 0)
    try:
        tab = T.ApiTab(st, Records(tmp / "r_cl.sqlite"), lambda m: None)
        tab.set_day("10.1")
        _wait(app, lambda: not tab.busy() and len(tab.videos) == 2)
        for v in tab.videos:
            v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
        tab._update_scope()
        # 先只发 1 条:还剩视频 → 不删
        tab.run_publish([0])
        _wait(app, lambda: not tab.busy())
        app.processEvents()
        check((acc / "10.2").is_dir(), "没发完(还剩视频):文件夹不删")
        for v in tab.videos:
            v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
        tab._update_scope()
        tab.run_publish()
        _wait(app, lambda: not tab.busy())
        _wait(app, lambda: not (acc / "10.2").exists(), 10)
        arc = sorted(os.listdir(tmp / "发布cl")) if (tmp / "发布cl").exists() else []
        moved = sorted(x for dd in arc for x in os.listdir(tmp / "发布cl" / dd / "CSGO"))
        check(not (acc / "10.2").exists() and (acc / "10.1").is_dir() and len(os.listdir(acc / "10.1")) == 4,
              f"手选的 10.2 发空了删 10.2;10.1 不动 -> 10.1 里 {sorted(os.listdir(acc / '10.1'))}")
        check("多的一张图.png" in moved and "二1.jpg" in moved and "Thumbs.db" not in moved,
              f"剩下的图挪进发布归档,Thumbs.db 删掉 -> {moved}")
        tab.load()                                        # 点「刷新」
        _wait(app, lambda: not tab.busy())
        check(not tab.videos and len(os.listdir(acc / "10.1")) == 4,
              f"手选的文件夹删了:这天这个号就没东西了,不会自己去找 10.1 列出来 -> {[v.title for v in tab.videos]}")
        check(acc.is_dir() and root.is_dir(), "号的总文件夹、素材根目录不动")
    finally:
        omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning = real


def test_edit_quota(tmp: Path):
    print("[每天规定条数自己改(用户 2026-10-22:有个号要从 10 条变成 30 条):每日汇总里双击一个号改]")
    from PySide6.QtWidgets import QApplication, QInputDialog
    QApplication.instance() or QApplication([])
    from qe import plan
    from qe.config import Settings
    from qe.records import Records
    from gui.daily_tab import DailyTab
    st = Settings(tmp / "st_eq.json")
    check(plan.quota_of(st, "DNF连招档案馆") == 10, "默认 DNF连招档案馆 每天 10 条")
    tab = DailyTab(st, Records(tmp / "r_eq.sqlite"))
    tab.refresh()
    row = next(r for r in range(tab.t.rowCount()) if tab.t.item(r, 1).text() == "DNF连招档案馆")
    real = QInputDialog.getInt
    QInputDialog.getInt = staticmethod(lambda *a, **k: (30, True))
    try:
        tab.edit_quota(row)
    finally:
        QInputDialog.getInt = real
    st2 = Settings(tmp / "st_eq.json")                               # 重新读一遍:存盘了
    row = next(r for r in range(tab.t.rowCount()) if tab.t.item(r, 1).text() == "DNF连招档案馆")
    check(plan.quota_of(st2, "DNF连招档案馆") == 30 and tab.t.item(row, 2).text() == "30" and tab.t.item(row, 4).text() == "还差 30",
          f"双击改成 30:存下来了,表里马上按 30 算 -> {tab.t.item(row, 2).text()} / {tab.t.item(row, 4).text()}")
    check(plan.quota_of(st2, "家长研习空间") == 102 and plan.quota_of(st2, "和平精英上分君") == 80, "别的号不变")
    check(plan.left_today(Records(tmp / "r_eq.sqlite"), st2, "DNF连招档案馆") == 30, "发够就停的闸也按新数算")


def test_extra_vert_order(tmp: Path):
    print("[今天临时多发 / 教育号转标准竖屏 / 发完的号沉底(用户 2026-10-23)]")
    import datetime
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, plan, vertical
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    import gui.api_tab as T
    import gui.window as GW
    today_md = f"{bj_now().month}.{bj_now().day}"
    tmr = plan.tomorrow_folder()
    # ---- 今天临时多发:上限加 N(只管今天),明天文件夹里挪 N 条回来 ----
    root = tmp / "ev" / "游戏视频"
    for nick in ("CSGO战神君", "PUBG突围君", "DNF连招档案馆"):
        g = nick[:4]
        (root / f"{g}--{nick}" / today_md).mkdir(parents=True, exist_ok=True)
        for k in range(2):
            (root / f"{g}--{nick}" / today_md / f"{nick}{k}.mp4").write_bytes(b"x")
    (root / "CSGO--CSGO战神君" / tmr).mkdir(parents=True, exist_ok=True)
    for k in range(3):
        (root / "CSGO--CSGO战神君" / tmr / f"明天{k}.mp4").write_bytes(b"x")
        (root / "CSGO--CSGO战神君" / tmr / f"明天{k}.jpg").write_bytes(b"j")
    apiacc.save([{"nick": n, "source_name": n[:4], "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                 for n in ("CSGO战神君", "PUBG突围君", "DNF连招档案馆")])
    st = Settings(tmp / "st_ev.json")
    st["api_roots"] = [str(root)]
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"CSGO战神君": 2, "PUBG突围君": 5, "DNF连招档案馆": 5}},
                        {"cat": "教育", "label": "教育总发布量", "accounts": {"家长小天地": 3}}]
    rec = Records(tmp / "r_ev.sqlite")
    for k in range(2):
        rec.api_result(F.Video(path=str(tmp / f"sent{k}.mp4"), title=f"s{k}", account="CSGO战神君", employee="", date=today_md), "已发")
    check(plan.left_today(rec, st, "CSGO战神君") == 0, "CSGO战神君 今天规定 2 条,发够了")
    tab = T.ApiTab(st, rec, lambda m: None)
    _wait(app, lambda: not tab.busy() and tab.videos)
    moved = tab.add_extra_today("CSGO战神君", 2)
    _wait(app, lambda: not tab.busy())
    here = sorted(os.listdir(root / "CSGO--CSGO战神君" / today_md))
    left_tmr = sorted(os.listdir(root / "CSGO--CSGO战神君" / tmr))
    # 打开页面时 CSGO 已经发够了:今天那 2 条先被挪到了明天;临时多发 2 条 → 挪回来的就是它俩(排在最前面)
    check(moved == 2 and plan.left_today(rec, st, "CSGO战神君") == 2 and here == ["CSGO战神君0.mp4", "CSGO战神君1.mp4"]
          and left_tmr == ["明天0.jpg", "明天0.mp4", "明天1.jpg", "明天1.mp4", "明天2.jpg", "明天2.mp4"],
          f"今天临时多发 2 条:今天还能发 2 条,明天文件夹挪回 2 条 -> 今天 {here} / 明天 {left_tmr}")
    tab.add_extra_today("CSGO战神君", 1)
    _wait(app, lambda: not tab.busy())
    here = sorted(os.listdir(root / "CSGO--CSGO战神君" / today_md))
    check("明天0.mp4" in here and "明天0.jpg" in here and plan.quota_on(st, "CSGO战神君") == 5,
          f"再多发 1 条:累加,封面跟着挪回来 -> {here}")
    y = (bj_now() + datetime.timedelta(days=1)).strftime("%Y-%m-%d")
    check(plan.quota_on(st, "CSGO战神君") == 5 and plan.quota_on(st, "CSGO战神君", y) == 2, "只管今天,明天还是原来的条数")
    s = plan.summary(rec, st, plan.today())
    check(("游戏", "CSGO战神君", 5, 2) in s["accounts"], f"每日汇总按 2+2 算 -> {s['accounts']}")
    # ---- 发完的号沉底 ----
    order = [tab.tree.topLevelItem(k).data(0, T.NICK) for k in range(tab.tree.topLevelItemCount())]
    pubg = [i for i, v in enumerate(tab.videos) if v.account == "PUBG突围君"]
    for i in pubg:
        tab.state[i] = "已发"
    tab._reorder()
    order2 = [tab.tree.topLevelItem(k).data(0, T.NICK) for k in range(tab.tree.topLevelItemCount())]
    check(order2[-1] == "PUBG突围君" and tab.groups["PUBG突围君"].text(0).endswith("· 发完了")
          and not tab.groups["PUBG突围君"].isExpanded(), f"发完的号沉到最下面、收起、写「发完了」-> {order} → {order2}")
    tab.k_undone.setChecked(True)
    check(all(tab.items[i].isHidden() for i in pubg) and not tab.groups["CSGO战神君"].isHidden(), "「只看没发完的号」:发完的藏起来")
    tab.k_undone.setChecked(False)
    # ---- 教育号:今天临时多发,从明天的文件夹挪回来(文件夹结构不变)----
    mat = tmp / "ev_mat"
    (mat / today_md / "小天地" / "51").mkdir(parents=True, exist_ok=True)
    (mat / today_md / "小天地" / "51" / "今天的.mp4").write_bytes(b"x")
    (mat / tmr / "小天地" / "51").mkdir(parents=True, exist_ok=True)
    (mat / tmr / "小天地" / "51" / "明天的.mp4").write_bytes(b"x")
    (mat / tmr / "小天地" / "51" / "明天的.png").write_bytes(b"p")
    st["material_root"] = str(mat)
    st["splash"] = False
    w = GW.MainWindow(st, rec)
    w.add_extra("家长小天地", 1)
    got = sorted(os.listdir(mat / today_md / "小天地" / "51"))
    check("明天的.mp4" in got and "明天的.png" in got and plan.quota_on(st, "家长小天地") == 1 + plan.quota_of(st, "家长小天地"),
          f"教育号今天临时多发 1 条:明天的那条连封面挪回今天(结构不变)-> {got}")
    # ---- 教育号转标准竖屏 ----
    ff = vertical.ffmpeg_path()
    if ff:
        import subprocess
        vd = tmp / "ev_vert"
        vd.mkdir(exist_ok=True)
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-t", "1",
                        "-c:v", "libx264", str(vd / "横的.mp4")], check=True)
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=360x640:rate=25", "-t", "1",
                        "-c:v", "libx264", str(vd / "竖的.mp4")], check=True)
        vids = [F.Video(path=str(vd / "横的.mp4"), title="横的", account="家长小天地", employee="51", date=today_md),
                F.Video(path=str(vd / "竖的.mp4"), title="竖的", account="家长小天地", employee="51", date=today_md)]
        res = {}
        wk = GW.VertWorker(vids, [0, 1])
        wk.one.connect(lambda i, r: res.__setitem__(i, r))
        wk.run()
        app.processEvents()
        info0, info1 = vertical.probe(vids[0].path, ff), vertical.probe(vids[1].path, ff)
        check(res.get(0, "").startswith("转好了") and (info0["w"], info0["h"]) == (1080, 1920) and res.get(1) == ""
              and (info1["w"], info1["h"]) == (360, 640) and sorted(os.listdir(vd)) == ["横的.mp4", "竖的.mp4"],
              f"教育号转标准竖屏:横的转成 1080×1920(原片不留),竖的不动 -> {res} {info0} {sorted(os.listdir(vd))}")
    else:
        check(True, "(没有 ffmpeg,转竖屏那段跳过)")


def test_verify(tmp: Path):
    print("[去企鹅号核对(用户 2026-10-23):软件记的已发 vs 企鹅号上的;没发上的拎回来重发(挪回原文件夹、改没发成)]")
    from qe import folders as F
    from qe import verify
    from qe.config import bj_now
    from qe.records import Records
    day = bj_now().strftime("%Y-%m-%d")
    rec = Records(tmp / "r_vf.sqlite")
    src = tmp / "vf" / "游戏视频" / "CSGO--CSGO战神君" / "10.23"
    arc = tmp / "vf" / "发布" / "CSGO"
    src.mkdir(parents=True)
    arc.mkdir(parents=True)
    rows = {}
    for k, (title, aid) in enumerate((("在的", "A1"), ("审核中的", "A2"), ("被删的", "A3"), ("没上去的", "A4"), ("G镇大写", "A5"),
                                      ("后台搜得到", "A6"))):
        v = F.Video(path=str(src / f"{title}.mp4"), title=title, account="CSGO战神君", employee="", date="10.23")
        rid = rec.api_begin(v)
        rec.api_finish(rid, "已发", f"article_id={aid} vid=V{k}")
        (arc / f"{title}.mp4").write_bytes(b"x")
        (arc / f"{title}.jpg").write_bytes(b"j")
        rec.mark_moved(rid, str(arc / f"{title}.mp4"))
        rows[title] = rid
    now = bj_now().strftime("%Y-%m-%d %H:%M:%S")

    class Fake:
        def list_articles(self, page, limit):
            arts = [{"article_url": "http://kuaibao.qq.com/s/A100", "article_pub_flag": "发布成功", "article_pub_time": now,
                     "article_title": "在的", "article_video_info": {"vid": "V0"}},
                    {"article_id": "A2", "article_pub_flag": "审核中", "article_pub_time": now, "article_title": "审核中的"},
                    {"article_id": "A3", "article_pub_flag": "发布成功", "article_del_flag": "已删除", "article_pub_time": now,
                     "article_title": "被删的"},
                    {"article_id": "B9", "article_pub_flag": "发布成功", "article_pub_time": now, "article_title": "手动发的"},
                    {"article_id": "A5", "article_pub_flag": "发布成功", "article_pub_time": "2001-01-01 10:00:00", "article_title": "老的"}]
            return {"articles": arts if page == 1 else [], "total": "5"}

        def article_info(self, aid):
            return {"article_info": {"article_pub_flag": {"A4": "未发布", "A5": "发布成功", "A6": "发布成功"}.get(aid, "未发布"),
                                     "article_del_flag": "已删除" if aid == "A3" else "未删除",
                                     "article_pub_time": now if aid == "A6" else "2001-01-01 10:00:00"}}
    r = verify.check_account(Fake(), rec, "CSGO战神君", day)
    miss = {m["title"]: m["why"] for m in r["missing"]}
    sure = {m["title"]: m["sure"] for m in r["missing"]}
    check(r["ours"] == 6 and r["live"] == 4 and r["ok"] == 3 and r["review"] == 1 and set(miss) == {"被删的", "没上去的", "G镇大写"}
          and "删了" in miss["被删的"] and "未发布" in miss["没上去的"] and not sure["被删的"] and sure["没上去的"],
          f"核对:在的 / 审核中的 算对上;没上去的默认打勾,被删的列出来不打勾(多半是人自己删的)-> {r['ours']} {r['live']} {miss}")
    check(not sure["G镇大写"] and "单查是「发布成功」" in miss["G镇大写"] and "不是这天" in miss["G镇大写"],
          f"单查说发布成功、但发布时间不是这天的:列出来让人看(不默认打勾)-> {miss['G镇大写']}")
    check("后台搜得到" not in miss, "列表里没读到、单查是这天发布成功的:算对上(用户 2026-10-24:后台一搜就在,不是假发成功)")

    class Pages:
        def list_articles(self, page, limit):
            old = {"article_pub_flag": "发布成功", "article_pub_time": "2001-01-01 10:00:00", "article_title": "老"}
            pages = {1: [{"article_id": "P1", "article_pub_flag": "已发布", "article_pub_time": now, "article_title": "一"}],
                     2: [old] * 10, 3: [{"article_id": "P3", "article_pub_flag": "发布成功", "article_pub_time": now, "article_title": "三"}],
                     4: [old] * 10, 5: [old] * 10, 6: [old] * 10, 7: [{"article_id": "P7", "article_pub_time": now}]}
            return {"articles": pages.get(page, []), "total": "999"}
    got = verify.platform_on(Pages(), day)
    check([p["aid"] for p in got] == ["P1", "P3"], f"列表不严格按时间排:中间夹一页老的不停,连着 3 页老的才停 -> {[p['aid'] for p in got]}")
    check([e["title"] for e in r["extra"]] == ["手动发的"], f"企鹅号上有、软件没记的单独列 -> {[e['title'] for e in r['extra']]}")
    m = next(x for x in r["missing"] if x["title"] == "没上去的")
    check(verify.pull_back(rec, m) == "" and (src / "没上去的.mp4").is_file() and (src / "没上去的.jpg").is_file()
          and not (arc / "没上去的.mp4").exists() and rec.status_of(str(src / "没上去的.mp4")) == ""
          and rec.status_of(str(src / "在的.mp4")) == "已发",
          "拎回来重发:视频和封面挪回原来的文件夹(文件夹没了也建),记录改成没发成,能再发;别的不动")
    import shutil
    shutil.rmtree(src)
    m2 = next(x for x in r["missing"] if x["title"] == "被删的")
    check(verify.pull_back(rec, m2) == "" and (src / "被删的.mp4").is_file(), "原来的日期文件夹发空被删了:建回来再挪")
    check(len(rec.sent_on(day).items()) and rec.sent_on(day).get("CSGO战神君") == 4, "拎回来的不再算今天已发(发够就停的闸跟着让出名额)")
    # 窗口:每日汇总「去企鹅号核对」→ 表 + 没发上的打勾 →「拎回来重发」
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi
    from gui.verify_dialog import VerifyDialog
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    real = omapi.Client
    omapi.Client = lambda *a, **k: Fake()
    try:
        rid = rows["G镇大写"]
        rec.mark_unsent(rid, "")
        rec.conn.execute("UPDATE posts SET status='已发' WHERE id=?", (rid,))       # 再造一条「记成已发、平台没有」
        rec.conn.execute("UPDATE posts SET note='article_id=A9 vid=V9', moved_pkey=? WHERE id=?",
                         (os.path.normcase(str(arc / "G镇大写.mp4")), rid))
        rec.conn.commit()
        dlg = VerifyDialog(rec, day)
        _wait(app, lambda: not dlg.w.isRunning(), 20)
        app.processEvents()
        check(dlg.t.rowCount() == 1 and dlg.miss.count() == 1 and "G镇大写" in dlg.miss.item(0).text() and dlg.b_back.isEnabled()
              and dlg.t.horizontalHeaderItem(1).text() == "应发" and dlg.t.horizontalHeaderItem(3).text() == "还差",
              f"核对窗口:每个号一行(应发 / 实发 / 还差 …),假发成功的列出来打勾 -> {[dlg.miss.item(k).text() for k in range(dlg.miss.count())]}")
        dlg.pull_back()
        check((src / "G镇大写.mp4").is_file() and "拎回来 1 条" in dlg.lab_done.text(),
              f"点「拎回来重发」:挪回原文件夹,文件名大小写不变 -> {sorted(os.listdir(src))}")
        # 只查指定的号(右键一个号「去企鹅号核对…」)
        apiacc.save([{"nick": n, "source_name": n, "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                     for n in ("CSGO战神君", "PUBG突围君", "DNF连招档案馆")])
        dlg2 = VerifyDialog(rec, day, nicks=["PUBG突围君"])
        _wait(app, lambda: not dlg2.w.isRunning(), 20)
        app.processEvents()
        check(dlg2.t.rowCount() == 1 and dlg2.t.item(0, 0).text() == "PUBG突围君" and "PUBG突围君" in dlg2.windowTitle(),
              f"指定一个号核对:只查这一个 -> {[dlg2.t.item(k, 0).text() for k in range(dlg2.t.rowCount())]}")
        dlg3 = VerifyDialog(rec, day, nicks=["家长小天地"])
        check(dlg3.w is None and "教育号" in dlg3.lab_done.text(), "教育号指定核对:说清楚查不了")
        # 应发 / 实发 / 还差(带设置)
        from qe.config import Settings
        st = Settings(tmp / "st_vf.json")
        st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"PUBG突围君": 5, "CSGO战神君": 5}}]
        dlg4 = VerifyDialog(rec, day, nicks=["PUBG突围君"], st=st)
        _wait(app, lambda: not dlg4.w.isRunning(), 20)
        app.processEvents()
        row4 = [dlg4.t.item(0, c).text() for c in range(4)]
        check(row4 == ["PUBG突围君", "5", "3", "2"], f"应发 5、企鹅号上实发 3、还差 2 -> {row4}")
    finally:
        omapi.Client = real


def test_unsend(tmp: Path):
    print("[已发的手动改回没发(用户 2026-10-24):视频挪回原文件夹、记录改没发成;视频找不到也改记录并说清楚]")
    from qe import folders as F
    from qe import verify
    from qe.records import Records
    rec = Records(tmp / "r_us.sqlite")
    src = tmp / "us" / "src" / "10.24"
    arc = tmp / "us" / "arc"
    src.mkdir(parents=True)
    arc.mkdir(parents=True)
    v = F.Video(path=str(src / "a.mp4"), title="a", account="CSGO战神君", employee="", date="10.24")
    rid = rec.api_begin(v)
    rec.api_finish(rid, "已发", "article_id=X1 vid=Y1")
    (arc / "a.mp4").write_bytes(b"x")
    rec.mark_moved(rid, str(arc / "a.mp4"))
    row = rec.sent_row(str(arc / "a.mp4"))
    check(row and row["id"] == rid and rec.sent_row(str(src / "a.mp4"))["id"] == rid, "按原位置 / 挪去的位置都找得到这条已发记录")
    check(verify.unsend(rec, row) == "" and (src / "a.mp4").is_file() and rec.status_of(str(src / "a.mp4")) == "",
          "改回没发:视频挪回原来的文件夹,能再发")
    w = F.Video(path=str(src / "b.mp4"), title="b", account="CSGO战神君", employee="", date="10.24")
    rid2 = rec.api_begin(w)
    rec.api_finish(rid2, "已发", "article_id=X2 vid=Y2")
    rec.mark_moved(rid2, str(arc / "不见了.mp4"))
    msg = verify.unsend(rec, rec.sent_row(str(arc / "不见了.mp4")))
    check("视频没找到" in msg and rec.status_of(str(src / "b.mp4")) == "", f"视频找不到:记录照样改回没发,说清楚要自己放回去 -> {msg}")


def test_api_dups(tmp: Path):
    print("[16 个号查重(用户 2026-10-24):号 + 标题(去掉话题空格标点);发过同名的不发、同一天同名只发一条]")
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    root = tmp / "dup" / "游戏视频"
    d = root / "CSGO--CSGO战神君" / "10.24"
    d.mkdir(parents=True)
    for nm in ("旧的", "同名 #CSGO", "同名", "新的"):
        (d / f"{nm}.mp4").write_bytes(b"x")
    (root / "PUBG--PUBG突围君" / "10.24").mkdir(parents=True)
    (root / "PUBG--PUBG突围君" / "10.24" / "旧的.mp4").write_bytes(b"x")
    apiacc.save([{"nick": n, "source_name": n[:4], "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}
                 for n in ("CSGO战神君", "PUBG突围君")])
    rec = Records(tmp / "r_dup.sqlite")
    rec.api_result(F.Video(path=str(tmp / "早就发了" / "旧的 #CSGO.mp4"), title="旧的 #CSGO", account="CSGO战神君", employee="",
                           date="10.20"), "已发")
    st = Settings(tmp / "st_dup.json")
    st["api_roots"] = [str(root)]
    st["quota_open"] = True
    tab = T.ApiTab(st, rec, lambda m: None)
    tab.set_day("10.24")
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 5)
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920, tags=[], platform=[])
    tab._update_sum()

    def i_of(acc, title):
        return next(k for k, v in enumerate(tab.videos) if v.account == acc and v.title == title)
    old, s1, s2, new = (i_of("CSGO战神君", t) for t in ("旧的", "同名 #CSGO", "同名", "新的"))
    other = i_of("PUBG突围君", "旧的")
    check(not tab._can_send(old) and "发过同名的(10.20)" in tab._dups.get(old, ""),
          f"这个号发过同名的(带话题的也认):不发,写清楚哪天发的 -> {tab._dups.get(old)}")
    check(sum(1 for k in (s1, s2) if tab._can_send(k)) == 1 and (s1 in tab._dups) != (s2 in tab._dups),
          f"同一天同一个号同名的两条:只发一条 -> {tab._dups}")
    check(tab._can_send(new) and tab._can_send(other), "不同名的、别的号的同名:照常能发")
    check(tab._status_text(old)[0] == "要处理", "重复的显示「要处理」")
    dup_i = s1 if s1 in tab._dups else s2
    tab._rename_rows([dup_i], lambda t: "改了个名", "改标题")
    check(tab._can_send(dup_i) and dup_i not in tab._dups, "让员工改个名:就能发了")


def test_view_tomorrow(tmp: Path):
    print("[切到明天的文件夹准备定时(用户 2026-10-24):今天发够了的号不算「发完了」,大字提示定到明天]")
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, plan
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    import gui.api_tab as T
    tmr = plan.tomorrow_folder()
    root = tmp / "vt" / "游戏视频"
    (root / "CSGO--CSGO战神君" / tmr).mkdir(parents=True)
    for k in range(3):
        (root / "CSGO--CSGO战神君" / tmr / f"明天{k}.mp4").write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_vt.json")
    st["api_roots"] = [str(root)]
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"CSGO战神君": 2}}]
    rec = Records(tmp / "r_vt.sqlite")
    for k in range(2):
        rec.api_result(F.Video(path=str(tmp / f"今天{k}.mp4"), title=f"今天{k}", account="CSGO战神君", employee="",
                               date=f"{bj_now().month}.{bj_now().day}"), "已发")
    tab = T.ApiTab(st, rec, lambda m: None)
    tab.set_day(tmr)
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 3)
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920, tags=[], platform=[])
    tab._update_sum()
    tab._update_scope()
    g = tab.groups["CSGO战神君"]
    tab._fill_group("CSGO战神君")
    check(not tab._done("CSGO战神君") and "发完了" not in g.text(0) and "今天发够了" not in g.text(0),
          f"看明天的文件夹:不算发完、不沉底,也不写「今天发够了」(用户 2026-10-24:25 号都没发)-> {g.text(0)}")
    check(g.text(3).endswith(f"{tmr} 0/2"), f"号那一行的进度按这天算:{tmr} 发了 0 / 规定 2 -> {g.text(3)}")
    check(tab.hero.big.text().startswith("这天 3 条能发") and "到点自动发" in tab.hero.sub.text() and tab.b_pub.text() == "今天都发够了",
          f"大字:这天的用「到点自动发」;「现在发布」按今天算 -> {tab.hero.big.text()} / {tab.hero.sub.text()} / {tab.b_pub.text()}")
    check(tab._quota_sub() == f"{tmr}还差 2 条", f"还差几条按这天算 -> {tab._quota_sub()}")


def test_cover_move_keep(tmp: Path):
    print("[封面:「图片修改需求 (16).png」也认;挪了文件夹 / 去了黑边,调好的框还在(用户 2026-10-24)]")
    from PIL import Image
    from qe import folders as F
    from qe.adjust import Adjustments
    d = tmp / "cmk" / "10.25" / "研习社"
    sub = d / "家有儿童手表的注意啦，防止被骗要当心"
    sub.mkdir(parents=True)
    (d / "家有儿童手表的注意啦，防止被骗要当心.mp4").write_bytes(b"x")
    (d / "家有儿童手表的注意啦，防止被骗要当心.jpg").write_bytes(b"j")
    Image.new("RGB", (1280, 720), (100, 120, 140)).save(sub / "图片修改需求 (16).png")
    vs = F.scan_account(str(tmp / "cmk" / "10.25" / "研习社"), "家长研习空间", "10.25")
    check(vs and vs[0].png.endswith("图片修改需求 (16).png"), f"封面文件夹里只有「图片修改需求 (16).png」:认成无字封面 -> {[v.png for v in vs]}")
    adj = Adjustments(tmp / "adj_cmk.json")
    old_v, old_c = r"X:\素材\25\研习社\a.mp4", r"X:\素材\25\研习社\a\横屏16x9.png"
    adj.set_box(old_v, old_c, "0.750", 0.3, 0.4)
    adj.set_ok(old_v, old_c)
    new_v, new_c = r"X:\素材\10.25\研习社\a.mp4", r"X:\素材\10.25\研习社\a\横屏16x9.png"
    check(adj.manual(new_v, new_c, "a") == {"0.750": [0.3, 0.4]} and adj.is_ok(new_v, new_c, "a"),
          "视频从 25 挪到 10.25:调好的框、「这张没问题」都还在(按标题 + 封面图名字找回来)")
    cache = r"C:\数据\去黑边的封面\横屏16x9_0123456789_1792837650_12345.png"
    check(adj.manual(new_v, cache, "a") == {"0.750": [0.3, 0.4]}, "去了黑边(封面换成副本路径):框也还在")
    adj.set_box(new_v, new_c, "0.562", 0.5, 0.5, title="a")
    check(adj.manual(new_v, new_c, "a") == {"0.750": [0.3, 0.4], "0.562": [0.5, 0.5]}, "在新位置再调一个比例:原来调过的比例不丢")
    check(adj.manual(r"X:\素材\10.25\研习社\b.mp4", r"X:\素材\10.25\研习社\b\横屏16x9.png", "b") == {}, "别的标题不会串")


def test_export_and_edu_dims(tmp: Path):
    print("[核对能选日期、导出企业微信那样的表;教育号一勾员工就认出横屏 / 偏方(用户 2026-10-25)]")
    import subprocess
    import openpyxl
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import dimcache, vertical, verify
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records
    st = Settings(tmp / "st_ex.json")
    rec = Records(tmp / "r_ex.sqlite")
    day = "2026-10-25"
    rec.conn.execute("INSERT INTO posts(date,account,employee,title,path,status,published_at) VALUES('10.25','家长小天地','51','a','p1','已发','2026-10-25 09:00:00')")
    rec.conn.commit()
    results = {"PUBG突围君": {"live": 19}, "CSGO战神君": {"error": "x"}}
    p = tmp / "导出.xlsx"
    verify.export_xlsx(str(p), st, rec, day, results)
    ws = openpyxl.load_workbook(p).active
    head = [ws.cell(1, c).value for c in range(1, 5)]
    rows = {ws.cell(r, 2).value: [ws.cell(r, c).value for c in range(1, 5)] for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value}
    tot = {ws.cell(r, 1).value: [ws.cell(r, 3).value, ws.cell(r, 4).value] for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value and not ws.cell(r, 2).value}
    check(head == ["品类", "账号名称", "规定视频条数", "10月25日"], f"表头和企业微信表一样 -> {head}")
    check(rows["PUBG突围君"] == ["PUBG", "PUBG突围君", 20, 19] and rows["DNF连招档案馆"][0] == "地下城与勇士"
          and rows["电竞情报馆"][0] == "亚运会电竞赛事", f"16 个号:品类写游戏名,实发用企鹅号上查到的 -> {rows['PUBG突围君']}")
    check(rows["家长小天地"][3] == 1 and rows["CSGO战神君"][3] == 0, "教育号 / 没查成的:按软件记录")
    check(tot.get("游戏总发布量") == [380, 19] and tot.get("教育总发布量", [0])[0] == 156 and tot.get("总发布量", [0])[0] == 736,
          f"小计:游戏 / 教育 / 动漫总发布量、总发布量 -> {tot}")
    check(any(str(m).startswith("A") and ":A" in str(m) for m in ws.merged_cells.ranges), f"教育、动漫的品类格子合并 -> {ws.merged_cells.ranges}")
    # 核对窗口:换日期重查
    from qe import apiacc, omapi
    from gui.verify_dialog import VerifyDialog
    apiacc.save([{"nick": "PUBG突围君", "source_name": "PUBG", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    asked = []

    class Fake:
        def list_articles(self, page, limit):
            asked.append(page)
            return {"articles": [], "total": "0"}

        def article_info(self, aid):
            return {}
    real = omapi.Client
    omapi.Client = lambda *a, **k: Fake()
    try:
        dlg = VerifyDialog(rec, "2026-10-24", st=st)
        _wait(app, lambda: not dlg.w.isRunning(), 20)
        dlg.start("2026-10-20")
        _wait(app, lambda: not dlg.w.isRunning(), 20)
        _wait(app, lambda: dlg.b_xls.isEnabled(), 5)
        check(dlg.day == "2026-10-20" and "10-20" in dlg.windowTitle() and dlg.t.rowCount() == 1 and dlg.b_xls.isEnabled(),
              f"核对窗口:能换日期重查,查完能导出 -> {dlg.windowTitle()}")
    finally:
        omapi.Client = real
    # 教育号:横屏 / 偏方认出来
    ff = vertical.ffmpeg_path()
    if not ff:
        check(True, "(没有 ffmpeg,横屏那段跳过)")
        return
    import gui.window as GW
    vd = tmp / "edu_dims"
    vd.mkdir()
    for nm, size in (("横的", "640x360"), ("方的", "480x480"), ("竖的", "360x640")):
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", f"testsrc=size={size}:rate=25", "-t", "1", "-c:v", "libx264",
                        str(vd / f"{nm}.mp4")], check=True)
    vids = [F.Video(path=str(vd / f"{nm}.mp4"), title=nm, account="家长小天地", employee="51", date="10.25") for nm in ("横的", "方的", "竖的")]
    wk = GW.DimWorker(vids, [0, 1, 2])
    wk.run()
    check([getattr(v, "shape", "") for v in vids] == ["横屏", "偏方", ""], f"教育号读尺寸:横屏 / 偏方 / 竖的 -> {[getattr(v, 'shape', '') for v in vids]}")
    check(dimcache.get(str(vd / "横的.mp4")) == (640, 360), "读过的记着,下次不再读")


def test_black_border(tmp: Path):
    print("[封面去黑边(用户 2026-10-22):上下 / 左右的黑条去掉再裁、再传,原图不动;黑得太多的不去、标黄;暗的夜景不误伤]")
    from PIL import Image, ImageDraw
    from qe import apicover
    from qe import cover as C
    d = tmp / "bb"
    d.mkdir(exist_ok=True)
    im = Image.new("RGB", (2048, 1152), (0, 0, 0))
    ImageDraw.Draw(im).rectangle([0, 82, 2047, 1068], fill=(120, 160, 90))
    p = d / "横屏16x9.png"
    im.save(p)
    out, why = C.deborder(str(p))
    got = Image.open(out).size if out != str(p) else None
    check(why == "" and got and got[0] == 2048 and 980 <= got[1] <= 987 and C.orig_of(out) == str(p)
          and Image.open(p).size == (2048, 1152), f"上下黑条去掉,原图不动 -> {got}")
    check(C.deborder(str(p))[0] == out, "做过的副本直接用(不重复做)")
    import numpy as _np
    rng = _np.random.default_rng(1)
    sky = _np.clip(rng.normal(12, 7, (600, 800, 3)), 0, 255).astype(_np.uint8)       # 夜空:很暗,但有噪点(真照片都有)
    night = Image.fromarray(sky)
    ImageDraw.Draw(night).ellipse([300, 200, 500, 400], fill=(230, 220, 120))
    pn = d / "夜景.png"
    night.save(pn)
    check(C.deborder(str(pn)) == (str(pn), ""), "暗的夜景:不当成黑边")
    tall = Image.new("RGB", (1152, 2048), (0, 0, 0))
    ImageDraw.Draw(tall).rectangle([0, 900, 1151, 1250], fill=(200, 100, 100))
    pt = d / "竖屏9x16.png"
    tall.save(pt)
    check(C.deborder(str(pt)) == (str(pt), "too_wide"), "大半是黑的(横图塞进竖画布):不去,交给人看")
    clean = d / "干净.png"
    Image.new("RGB", (640, 360), (200, 200, 200)).save(clean)
    check(C.deborder(str(clean)) == (str(clean), ""), "没黑边:用原图")
    # 16 个号传的封面:原样模式也去黑边
    up = apicover.prepare(str(p), "原样")
    check(up != str(p) and Image.open(up).size[1] < 1000, f"16 个号的封面:有黑边的去掉再传 -> {Image.open(up).size}")
    apicover.cleanup(up, str(p))
    check(apicover.prepare(str(clean), "原样") == str(clean), "16 个号的封面:没黑边的原图照传")


def test_quota_gate(tmp: Path):
    print("[每天规定条数的闸(2026-10-20 PUBG 多发了 21 条):不能多发、没发成的补上、多的挪到明天、开闸才不管;转竖屏不留原片]")
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi, plan, vertical
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    tmr = plan.tomorrow_folder()
    root = tmp / "api_q"
    accs = [("CSGO", "CSGO战神君")]
    apiacc.save([{"nick": n, "source_name": g, "client_id": g, "secret": apiacc.seal("s"), "category": 1604, "enabled": True} for g, n in accs])
    st = Settings(tmp / "st_q.json")
    check(st.get("quota_open") is False, "默认闸关着(不能超过每天规定条数)")
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"CSGO战神君": 3}},
                        {"cat": "教育", "label": "教育总发布量", "accounts": {"家长小天地": 2}}]
    st["api_roots"] = [str(root / "游戏视频")]
    st["api_archive"] = str(tmp / "发布q" / "{date}")
    st["api_interval"] = 0
    st["api_cover"] = "不传"
    sent = []

    class Fake:
        def __init__(self, *a, **k):
            pass

        def send(self, path, title, cat, before_publish=None, **k):
            if "坏" in title:
                raise omapi.ApiError("上传视频失败:invalid_file_upload(code 40023)", "40023")
            if before_publish:
                before_publish("V")
            sent.append(title)
            return {"vid": "V", "article_id": "A" + title, "transaction_id": "P", "status": "成功"}
    real = (omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning)
    boxes = []
    omapi.Client, T.ask = Fake, (lambda *a, **k: True)
    QMessageBox.information = QMessageBox.warning = staticmethod(lambda _p, t, x, *a, **k: boxes.append(x) or 0)
    try:
        def run(names, rec_name):
            d = root / "游戏视频" / "CSGO--CSGO战神君" / "10.20"
            if d.exists():
                shutil.rmtree(d)
            d.mkdir(parents=True)
            for nm in names:
                (d / f"{nm}.mp4").write_bytes(b"x")
                (d / f"{nm}.jpg").write_bytes(b"j")
            shutil.rmtree(root / "游戏视频" / "CSGO--CSGO战神君" / tmr, ignore_errors=True)
            rec = Records(tmp / rec_name)
            tab = T.ApiTab(st, rec, lambda m: None)
            tab.set_day("10.20")
            _wait(app, lambda: not tab.busy() and len(tab.videos) == len(names))
            for v in tab.videos:
                v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
            tab._update_scope()
            tab._update_sum()
            return tab, rec
        tab, rec = run(["q1", "q2", "q3", "q4", "q5"], "rq1.sqlite")
        check(tab.b_pub.text() == "现在发布 3 条" and "超过今天的规定条数" in tab.hero.sub.text(),
              f"按钮只写规定条数以内的;大字下面写多的发够以后挪到明天 -> {tab.b_pub.text()} / {tab.hero.sub.text()}")
        sent.clear()
        tab.run_publish()
        _wait(app, lambda: not tab.busy(), 30)
        app.processEvents()
        moved = sorted(os.listdir(root / "游戏视频" / "CSGO--CSGO战神君" / tmr)) if (root / "游戏视频" / "CSGO--CSGO战神君" / tmr).exists() else []
        check(sorted(sent) == ["q1", "q2", "q3"] and moved == ["q4.jpg", "q4.mp4", "q5.jpg", "q5.mp4"],
              f"规定 3 条:发 3 条就停,多的 2 条连封面挪到明天的文件夹({tmr})-> 发了 {sent},挪了 {moved}")
        tab, rec = run(["q1", "坏2", "q3", "q4", "q5"], "rq2.sqlite")
        sent.clear()
        tab.run_publish()
        _wait(app, lambda: not tab.busy(), 30)
        app.processEvents()
        moved = sorted(x for x in os.listdir(root / "游戏视频" / "CSGO--CSGO战神君" / tmr) if x.endswith(".mp4"))
        check(sorted(sent) == ["q1", "q3", "q4"] and moved == ["q5.mp4", "坏2.mp4"],
              f"中间一条没发成:拿下一条补上,还是发够 3 条(不少发);没发成的也挪到明天 -> 发了 {sent},挪了 {moved}")
        boxes.clear()
        tab, rec2 = run(["q6", "q7"], "rq2.sqlite")               # 同一个记录库:今天已经发够 3 条了
        sent.clear()
        tab.run_publish()
        app.processEvents()
        moved = sorted(x for x in os.listdir(root / "游戏视频" / "CSGO--CSGO战神君" / tmr) if x.endswith(".mp4"))
        check(not sent and (not tab.videos or (boxes and "已经发够" in boxes[-1])) and moved == ["q6.mp4", "q7.mp4"],
              f"今天已经发够的号:不发,还没发的挪到明天 -> {boxes[-1][:40] if boxes else ''} {moved}")
        # 打开 / 刷新页面时:今天已经发够的号,今天日期文件夹里没发的自动挪到明天(名字差个空格的封面也跟着)
        from qe.config import bj_now
        today_md = f"{bj_now().month}.{bj_now().day}"
        td = root / "游戏视频" / "CSGO--CSGO战神君" / today_md
        td.mkdir(parents=True, exist_ok=True)
        (td / "今天的.mp4").write_bytes(b"x")
        (td / "今天的 .jpg").write_bytes(b"j")
        st["api_picked"] = {"10.20": {"CSGO战神君": str(td)}}
        tab = T.ApiTab(st, Records(tmp / "rq2.sqlite"), lambda m: None)   # 这个记录库里 CSGO 今天已经发够 3 条
        tab.set_day("10.20")
        _wait(app, lambda: not tab.busy(), 30)
        app.processEvents()
        moved = sorted(os.listdir(root / "游戏视频" / "CSGO--CSGO战神君" / tmr))
        check("今天的.mp4" in moved and "今天的 .jpg" in moved and not os.listdir(td) and not tab.videos,
              f"打开页面:今天发够的号,今天文件夹里没发的自动挪到明天(差个空格的封面也挪)-> {[m for m in moved if '今天' in m]}")
        st["api_picked"] = {}
        st["quota_open"] = True
        tab, rec = run(["q1", "q2", "q3", "q4", "q5"], "rq3.sqlite")
        sent.clear()
        tab.run_publish()
        _wait(app, lambda: not tab.busy(), 30)
        check(len(sent) == 5 and tab.b_pub.text() != "现在发布 3 条", f"开了闸:能发的都发 -> {sent}")
        st["quota_open"] = False
    finally:
        omapi.Client, T.ask, QMessageBox.information, QMessageBox.warning = real
    # 教育号:最后一批只摆差的那几条;发够了还没发的挪到明天(文件夹结构不变)
    import gui.window as GW
    from qe import folders as F
    eroot = tmp / "素材q"
    d = eroot / "10.20" / "小天地" / "张"
    d.mkdir(parents=True)
    for t in ("一一一一视频", "二二二二视频", "三三三三视频"):
        (d / f"{t}.mp4").write_bytes(b"x")
        _img(d / f"{t}.png", 640, 360)
    st["material_root"] = str(eroot)
    rec = Records(tmp / "rq_edu.sqlite")
    real_dec = GW.C.decide

    def fake_decide(v):
        v.cover, v.cover_kind, v.cover_why, v.plan, v.warn = v.png, "png", "测试", {}, False
    GW.C.decide = fake_decide
    try:
        w = GW.MainWindow(st, rec)
        w.c_date.setCurrentIndex(w.c_date.findData("10.20"))
        from PySide6.QtCore import Qt
        top = next(w.tree.topLevelItem(k) for k in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(k).text(0) == "家长小天地")
        top.child(0).setCheckState(0, Qt.CheckState.Checked)
        _wait(app, lambda: not (w.cover_worker and w.cover_worker.isRunning()))
        acc, vids = w.next_batch()
        check(acc == "家长小天地" and len(vids) == 2, f"教育号规定 2 条:一批只摆 2 条 -> {len(vids)}")
        rec.manual_published(vids[0])
        acc, vids2 = w.next_batch()
        check(len(vids2) == 1, f"今天发了 1 条:下一批只摆差的 1 条 -> {len(vids2)}")
        rec.manual_published(vids2[0])
        check(w.next_batch()[1] == [] and "发够规定条数" in w._nothing_left(), "发够了:不再摆,写清楚发够了")
        n = w._edu_spill()
        left = [p.name for p in (eroot / tmr / "小天地" / "张").glob("*")] if (eroot / tmr).exists() else []
        rest = next(v.title for v in w.videos if v not in (vids[0], vids2[0]))
        check(n == 1 and sorted(left) == [f"{rest}.mp4", f"{rest}.png"],
              f"还没发的连封面挪到明天的文件夹,结构不变(小天地\\张)-> {left}")
        w.close()
    finally:
        GW.C.decide = real_dec
    # 转竖屏不留原片(16 个号页,用户 2026-10-20)
    ff = vertical.ffmpeg_path()
    if ff:
        import subprocess
        vd = tmp / "vq"
        vd.mkdir()
        subprocess.run([ff, "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25", "-t", "1",
                        "-c:v", "libx264", str(vd / "横.mp4")], check=True)
        out, orig = vertical.to_vertical(str(vd / "横.mp4"), ff, keep_orig=False)
        info = vertical.probe(out, ff) if hasattr(vertical, "probe") else None
        check(orig == "" and os.path.basename(out) == "横.mp4" and not (vd / "横屏原片").exists() and sorted(os.listdir(vd)) == ["横.mp4"]
              and os.path.getsize(out) > 1024, f"转竖屏不留原片:转好的放回原名,原片删了,不建「横屏原片」-> {os.listdir(vd)}")
        del info


def test_misfire_api(tmp: Path):
    print("[防误发 · 16 个号(改进清单 A1 A2):按钮叫「现在发布 N 条」、定时挪开;确认框默认「先不发」、红字写现在马上发、能改成定时]")
    from PySide6.QtWidgets import QApplication, QHBoxLayout, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    import gui.schedule as S
    root = tmp / "api_misfire"
    (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18").mkdir(parents=True)
    for f in ("m1.mp4", "m2.mp4"):
        (root / "游戏视频" / "CSGO--CSGO战神君" / "10.18" / f).write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    st = Settings(tmp / "st_misfire.json")
    st["api_roots"] = [str(root / "游戏视频")]
    st["api_archive"] = str(tmp / "发布misfire" / "{date}")
    tab = T.ApiTab(st, Records(tmp / "rec_misfire.sqlite"), lambda m: None)
    tab.c_date.setCurrentText("10.18")
    _wait(app, lambda: not tab.busy() and len(tab.videos) == 2)
    for v in tab.videos:
        v.checks = dict(v.checks or {}, bad_read=False, landscape=False, w=1080, h=1920)
    tab._update_scope()
    check(tab.b_pub.text() == "现在发布 2 条" and tab.b_sched.text() == "到点自动发…",
          f"按钮名字:「{tab.b_pub.text()}」「{tab.b_sched.text()}」")
    bar = next(lay for lay in tab.findChildren(QHBoxLayout) if lay.indexOf(tab.b_pub) >= 0)
    check(bar.indexOf(tab.b_sched) >= 0 and bar.indexOf(tab.b_pub) - bar.indexOf(tab.b_sched) >= 3,
          "定时按钮挪开了,不和「现在发布」挨着(中间隔着红字和空白)")
    seen, opened = [], []
    real = (T.ask, S.ask_schedule)

    def rec_ask(*a, **k):
        seen.append(k)
        if k.get("on_extra"):
            k["on_extra"]()                         # 人点了「改成定时…」
        return False
    T.ask = rec_ask
    S.ask_schedule = lambda _p, title, *a, **k: opened.append(title)
    try:
        tab.run_publish()
    finally:
        T.ask, S.ask_schedule = real
    k = seen[-1] if seen else {}
    check(k.get("default_no") is True and "现在马上发出去" in k.get("head", "") and "后台删" in k.get("head", ""),
          f"发布确认框:默认「先不发」,第一行写现在马上发、发了只能去后台删 -> {k.get('head')}")
    check(k.get("extra") == "改成定时…" and opened and "16 个号" in opened[0], "确认框里点「改成定时…」打开定时对话框")
    check(not tab.busy(), "点了「改成定时…」不会发")
    # 真的确认框:回车 = 默认按钮 =「先不发」
    clicks, hit = [], []
    real_exec = QMessageBox.exec

    def fake_exec(box):
        want = clicks.pop(0) if clicks else box.defaultButton().text()
        seen.append(("box", box.defaultButton().text(), box.text()))
        next(b for b in box.buttons() if b.text() == want).click()
        return 0
    QMessageBox.exec = fake_exec
    try:
        r1 = T._real_ask(tab, "现在发布", "发 2 条", "现在发出去", "先不发", default_no=True,
                         head="现在马上发出去(不是定时),发了只能去后台删", extra="改成定时…", on_extra=lambda: hit.append(1))
        b1 = seen[-1]
        clicks.append("改成定时…")
        r2 = T._real_ask(tab, "现在发布", "发 2 条", "现在发出去", "先不发", default_no=True, extra="改成定时…",
                         on_extra=lambda: hit.append(1))
        clicks.append("现在发出去")
        r3 = T._real_ask(tab, "现在发布", "发 2 条", "现在发出去", "先不发", default_no=True)
    finally:
        QMessageBox.exec = real_exec
    from gui.theme import cur as _theme_cur              # 红字 = 主题里的「危险」红(2026-10-18 换了新界面)
    check(r1 is False and b1[1] == "先不发" and _theme_cur()["danger"] in b1[2] and "现在马上发出去" in b1[2],
          f"确认框直接回车 = 先不发,第一行红字 -> {b1[1]}")
    check(r2 is False and hit == [1] and r3 is True, "点「改成定时…」不发、去定时;点「现在发出去」才发")
    tab.deleteLater()


def test_misfire_edu(tmp: Path):
    print("[防误发 + 定时可靠性 · 教育号(改进清单 A1 A3 A4 B5~B8)]")
    from PySide6.QtCore import QObject, QTimer, Signal
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    import datetime as dt
    import gui.schedule as S
    import gui.window as GW
    from qe.config import Settings
    from qe.records import Records
    root = tmp / "素材misfire"
    d = root / "10.20" / "小天地" / "王"
    d.mkdir(parents=True)
    for k in range(5):
        (d / f"误发测试第{k}条.mp4").write_bytes(b"x")
        _img(d / f"误发测试第{k}条.png", 640, 360)
    H = {"w": None, "mode": "ok"}
    calls = []

    class FakeBW(QObject):
        log = Signal(str)
        one = Signal(object, bool, str)
        finished_batch = Signal(object, str, str)

        def __init__(self, acc, vids, shots, parent=None, auto_publish=False, unattended=False, sched_at=""):
            super().__init__(parent)
            self.vids = vids
            calls.append([v.title for v in vids])

        def isRunning(self):
            return False

        def stop(self):
            pass

        def start(self):
            def go():
                w, mode = H["w"], H["mode"]
                if mode == "fail":                  # 一条都没摆上(罐头没开)
                    self.finished_batch.emit([], "罐头没开", "")
                    return
                if mode == "stopmid":               # 摆的时候人点了「停」,但这批照样定好了
                    w._chain_stop = True
                res = []
                for v in self.vids:
                    self.one.emit(v, True, "")
                    res.append((v, True, ""))
                self.finished_batch.emit(res, "", "ok")
                if mode == "tamper":                # 摆完这批,有人改了勾选
                    w._chain_key = ("别的日期", [])
            QTimer.singleShot(0, go)

    def fake_decide(v):
        v.cover, v.cover_kind, v.cover_why, v.plan = v.png, "png", "测试", {}
        v.warn = False
    real = (GW.BatchWorker, GW.C.decide, GW.C.note_wide, QMessageBox.information, S.ask_schedule)
    boxes = []
    GW.BatchWorker, GW.C.decide, GW.C.note_wide = FakeBW, fake_decide, (lambda v: None)
    QMessageBox.information = staticmethod(lambda _p, t, x, *a, **k: boxes.append(x))
    try:
        rec = Records(tmp / "r_misfire.sqlite")
        st = Settings(tmp / "s_misfire.json")
        st["material_root"] = str(root)
        st["api_roots"] = [str(tmp / "api_none")]
        st["ai_key"] = "test"
        st["batch_size"] = 1
        st["auto_publish"] = True                   # 老版本记住过「勾着」
        st.save()
        w = GW.MainWindow(st, rec)
        H["w"] = w
        check(not w.k_auto.isChecked(), "「摆好后自动点发布」每次打开都不勾(以前记住的也不算)")
        w.k_auto.setChecked(True)
        w.k_auto.setChecked(False)
        check(Settings(tmp / "s_misfire.json").get("auto_publish") is True, "勾 / 不勾都不再存进设置")
        w.k_auto.setChecked(True)
        check(w.b_sched.text() == "摆好并定时发出…", f"教育号定时按钮叫「摆好并定时发出…」-> {w.b_sched.text()}")
        # A3:勾着时「摆下一批」确认框第一行红字
        w.c_date.setCurrentIndex(w.c_date.findData("10.20"))
        w._load_tree()
        for k in range(w.tree.topLevelItemCount()):
            top = w.tree.topLevelItem(k)
            for j in range(top.childCount()):
                top.child(j).setCheckState(0, GW.Qt.CheckState.Checked)
        w._load_videos()
        _wait(app, lambda: not (w.cover_worker and w.cover_worker.isRunning()))
        asks = []

        def rec_ask(title, text, yes, no, default_no=False, head="", danger=False):
            asks.append({"title": title, "text": text, "yes": yes, "default_no": default_no, "head": head, "danger": danger})
            return False
        w._ask = rec_ask
        w.start_batch()
        check(asks and "这次会自动点发布" in asks[-1]["head"] and not calls, f"勾着自动发布:确认框第一行红字写「这次会自动点发布」-> {asks[-1]['head'] if asks else ''}")
        check(asks[-1]["default_no"] is True and asks[-1]["danger"] and asks[-1]["yes"] == "开始摆,并自动发布",
              f"勾着自动发布:默认「先不」(回车不会真发),开始按钮红、写「并自动发布」-> {asks[-1]['yes']}")
        w.k_auto.setChecked(False)
        w.start_batch()
        check(asks[-1]["head"] == "" and asks[-1]["default_no"] is False and asks[-1]["yes"] == "开始摆", "没勾:没有那行红字,照旧")
        # A4:「这批发好了」默认「还没」
        w.batch_id = "假的一批"
        w.batch_sent()
        check(asks[-1]["title"] == "这批发好了" and asks[-1]["default_no"] is True and "不会再摆" in asks[-1]["text"]
              and w.batch_id == "假的一批", "「这批发好了」默认按钮是「还没」,写清楚点了这批不会再摆")
        w.batch_id = ""
        w._refresh_buttons()
        # B7:估时间够不够
        check(w._sched_estimate("10.20") == (5, 5), f"估算:5 条、一批 1 条 = 5 批 -> {w._sched_estimate('10.20')}")
        now = dt.datetime.strptime(S.now_str(), "%Y-%m-%d %H:%M")

        def slot(mins):
            t = now + dt.timedelta(minutes=mins)
            t += dt.timedelta(minutes=(5 - t.minute % 5) % 5)
            return t.strftime("%Y-%m-%d %H:%M")
        at_near, at_far = slot(45), slot(180)
        answers = [(at_near, "10.20"), None]
        S.ask_schedule = lambda *a, **k: answers.pop(0) if answers else None
        asks.clear()
        w.schedule_clicked()
        check(asks and asks[0]["title"] == "时间可能不够" and "5 批" in asks[0]["text"] and asks[0]["default_no"] is True
              and not w._unattended, f"5 批只给 45 分钟:先提示时间可能不够(默认「改时间」)-> {[a['title'] for a in asks]}")
        answers[:] = [(at_far, "10.20")]
        asks.clear()
        w.schedule_clicked()
        check([a["title"] for a in asks] == ["定时发布"], f"时间够:不提示 -> {[a['title'] for a in asks]}")
        # B5:摆的时候点了「停」,这批照样定好了 → 停下,不再接着摆;总结不写「电脑可以关了」
        H["mode"] = "stopmid"
        calls.clear()
        boxes.clear()
        w._sched_fire("10.20", at_far)
        _wait(app, lambda: not w._unattended, 20)
        check(not w._unattended and len(calls) == 1, f"点了「停」就停,不再接着摆(以前会被成功那批吃掉)-> {len(calls)} 批")
        check(boxes and "停在这里" in boxes[-1] and "还有 4 条没定" in boxes[-1] and "电脑可以关" not in boxes[-1],
              f"停下的总结:停在哪、还剩几条,不说电脑可以关 -> {boxes[-1][:120] if boxes else ''}")
        # B8:连发中途有人改了勾选 → 不弹「自动发布停下了」卡住,写日志收尾
        H["mode"] = "tamper"
        calls.clear()
        boxes.clear()
        w._sched_fire("10.20", at_far)
        _wait(app, lambda: not w._unattended, 20)
        check(not w._unattended and len(calls) == 1 and boxes and "换了日期" in boxes[-1] and "停在这里" in boxes[-1],
              f"没人看着时改了勾选:收尾写总结,不卡住 -> {boxes[-1][:60] if boxes else ''}")
        # B6:之前点过「停」(_chain_stop 还是 True),再开定时、第一批整批出错 → 1 分钟后照样重试,发到完
        w._chain_stop = True
        H["mode"] = "fail"
        calls.clear()
        boxes.clear()
        w._sched_fire("10.20", at_far)
        _wait(app, lambda: "1 分钟后再试" in w.logbox.toPlainText(), 10)
        check(w._unattended and not w._chain_stop and len(calls) == 1, "再开定时:以前点过的「停」清掉了,整批出错等着重试")
        H["mode"] = "ok"
        w._sched_next()                             # 当它 1 分钟到了
        _wait(app, lambda: not w._unattended, 30)
        check(not w._unattended and len(calls) == 4 and boxes and "都定好了" in boxes[-1] and "电脑可以关" in boxes[-1],
              f"重试成功,剩下的都定好了,写「电脑可以关了」-> {len(calls)} 批 {boxes[-1][:40] if boxes else ''}")
        # 重试等着的时候点了「停」:到点收尾,不会一直「正在自己发」
        w._unattended, w._chain_stop = True, True
        boxes.clear()
        w._sched_next()
        check(not w._unattended and boxes and "停" in boxes[-1], "等重试的时候点了停:收尾")
        w.close()
    finally:
        GW.BatchWorker, GW.C.decide, GW.C.note_wide, QMessageBox.information, S.ask_schedule = real


def test_sched_notice(tmp: Path):
    print("[16 个号定时没开始 / 没发成(改进清单 B9):⏰ 那行红字常驻,下次打开弹一次]")
    from PySide6.QtWidgets import QApplication, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    st = Settings(tmp / "st_notice.json")
    st["api_roots"] = [str(tmp / "api_notice_none")]
    rec = Records(tmp / "rec_notice.sqlite")
    warns = []
    real_w = QMessageBox.warning
    QMessageBox.warning = staticmethod(lambda _p, t, x, *a, **k: warns.append(x))
    try:
        tab = T.ApiTab(st, rec, lambda m: None)
        tab._sched_fire("9.99")
        note = (st.get("sched_notice") or {}).get("api") or {}
        check("没发成" in tab.lab_sched.text() and "9.99" in tab.lab_sched.toolTip() and note.get("popped") is False,
              f"到点没开始:红字常驻,鼠标放上去看为什么 -> {tab.lab_sched.text()}")
        _wait(app, lambda: False, 2)
        check(not warns, "当时不弹框(已经有红字了),等下次打开再弹")
        tab2 = T.ApiTab(Settings(tmp / "st_notice.json"), rec, lambda m: None)
        _wait(app, lambda: bool(warns), 5)
        check(len(warns) == 1 and "9.99" in warns[0], "下次打开软件弹一次框")
        tab3 = T.ApiTab(Settings(tmp / "st_notice.json"), rec, lambda m: None)
        _wait(app, lambda: len(warns) > 1, 3)
        check(len(warns) == 1 and "没发成" in tab3.lab_sched.text(), "只弹一次;红字还在(重新定时才清掉)")
        tab3._clear_notice()
        tab3._sched_changed()
        check(tab3.lab_sched.text() == "", "重新定时以后红字清掉")
        st4 = Settings(tmp / "st_notice.json")
        st4["schedules"] = {"api": {"at": "2000-01-01 09:00", "date": "10.17"}}
        st4.save()
        tab4 = T.ApiTab(st4, rec, lambda m: None)
        check("软件没开着" in ((st4.get("sched_notice") or {}).get("api") or {}).get("msg", "") and "没发成" in tab4.lab_sched.text(),
              "到点时软件没开着:也红字常驻")
        _wait(app, lambda: len(warns) > 1, 5)
        check(len(warns) == 2 and "软件没开着" in warns[-1], "到点时软件没开着:这次打开就弹一次")
        # 用户 2026-10-24:10-21 那条红字挂了好几天 → 点一下看完就清;不是今天的不再挂
        infos = []
        real_i = QMessageBox.information
        QMessageBox.information = staticmethod(lambda _p, t, x, *a, **k: infos.append(x))
        try:
            tab4._ack_notice()
        finally:
            QMessageBox.information = real_i
        check(infos and "软件没开着" in infos[0] and tab4.lab_sched.text() == "", "点一下红字:说为什么,然后清掉")
        st5 = Settings(tmp / "st_notice.json")
        st5["sched_notice"] = {"api": {"at": "2026-10-21 09:12", "msg": "旧的", "popped": True}}
        st5.save()
        tab5 = T.ApiTab(st5, rec, lambda m: None)
        check(tab5.lab_sched.text() == "" and not (st5.get("sched_notice") or {}).get("api"), "不是今天的提醒:不再挂着")
        for t in (tab, tab2, tab3, tab4, tab5):
            t.deleteLater()
    finally:
        QMessageBox.warning = real_w


def _modal(app, fn, wait: float = 0.3):
    """等一会儿,对弹出来的模态框调 fn(框)(测试里看框里写了什么 / 改了再点确定)。"""
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication
    QTimer.singleShot(int(wait * 1000), lambda: fn(QApplication.activeModalWidget()))


def test_cd_items(tmp: Path):
    print("[改进清单 C / D(10~23):看得懂、知道下一步、小问题]")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QLabel, QLineEdit, QMessageBox
    app = QApplication.instance() or QApplication([])
    import gui.schedule as S
    import gui.window as GW
    from gui.widgets import newest_first
    from qe import apiacc, apisrc
    from qe.config import Settings, bj_now
    from qe.records import Records
    today = f"{bj_now().month}.{bj_now().day}"
    root = tmp / "素材cd"
    for emp, names in (("王", ("cd第一条视频", "cd第二条视频")), ("李", ("cd第一条视频",))):
        d = root / "10.21" / "小天地" / emp
        d.mkdir(parents=True)
        for n in names:
            (d / f"{n}.mp4").write_bytes(b"x")
            _img(d / f"{n}.png", 640, 360)
    (root / "10.21" / "小天地" / "31条（已发）").mkdir(parents=True)
    (root / "10.21" / "小天地" / "31条（已发）" / "发过的那条.mp4").write_bytes(b"x")
    (root / "10.21" / "不认识").mkdir(parents=True)
    (root / "9.30" / "小天地" / "张").mkdir(parents=True)
    (root / "9.30" / "小天地" / "张" / "旧的.mp4").write_bytes(b"x")
    if today not in ("10.21", "9.30"):
        (root / today).mkdir()                      # 今天的文件夹还空着:下拉里标「今天」、灰掉
    api_root = tmp / "api_cd" / "游戏视频"
    (api_root / "CSGO--CSGO战神君" / "10.22").mkdir(parents=True)
    (api_root / "CSGO--CSGO战神君" / "10.22" / "重试的那条.mp4").write_bytes(b"x")
    accs0 = apiacc.load()

    def fake_decide(v):
        v.cover, v.cover_kind, v.cover_why, v.plan = v.png, "png", "测试", {}
        v.warn = False
    real = (GW.C.decide, GW.C.note_wide, QMessageBox.information)
    boxes = []
    GW.C.decide, GW.C.note_wide = fake_decide, (lambda v: None)
    QMessageBox.information = staticmethod(lambda _p, t, x, *a, **k: boxes.append(x))
    try:
        apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"),
                      "category": 1604, "enabled": True}])
        st = Settings(tmp / "s_cd.json")
        st["material_root"] = str(root)
        st["api_roots"] = [str(api_root)]
        st["archive_root"] = str(tmp / "教育归档" / "10.15")          # 被人改坏了的(少 {month})
        st["ai_key"] = "test"
        rec = Records(tmp / "r_cd.sqlite")
        w = GW.MainWindow(st, rec)
        # 刚打开:今天的文件夹有东西就选今天(试用 2026-10-21);没有就是空白 → 13:告诉人先选日期
        from qe import folders as F
        has_today = any(F.day_of(d) == today and w._date_has_files(str(root), d) for d in F.list_dates(str(root)))
        check(w.c_date.currentData() == (today if has_today else ""), f"刚打开:今天有素材就选今天 -> {w.c_date.currentData()}")
        w.c_date.setCurrentIndex(0)
        check(w.lab_sum.text() == GW.PICK_DAY and w.lab_scope.text() == GW.PICK_DAY, f"13 没选日期提示先选日期 -> {w.lab_sum.text()}")
        # 15:最新的在上面、今天标出来、没视频的灰掉
        ds = [w.c_date.itemData(i) for i in range(1, w.c_date.count())]
        check(ds == newest_first(ds) and ds.index("10.21") < ds.index("9.30") and w.c_date.itemText(0) == "(选一天)",
              f"15 教育号日期:最新的在上面 -> {ds}")
        check(newest_first(["9.30", "10.14（101）", "10.2", "乱七八糟"]) == ["10.14（101）", "10.2", "9.30", "乱七八糟"], "15 改过名的日期也按开头排")
        if today not in ("10.21", "9.30"):
            i = w.c_date.findData(today)
            fg = w.c_date.itemData(i, Qt.ItemDataRole.ForegroundRole)
            check("今天" in w.c_date.itemText(i) and "没视频" in w.c_date.itemText(i) and fg is not None,
                  f"15 今天标出来、空的灰掉 -> {w.c_date.itemText(i)}")
        i = w.c_date.findData("10.21")
        check(w.c_date.itemText(i) in ("10.21", "10.21(今天)") and w.c_date.itemData(i, Qt.ItemDataRole.ForegroundRole) is None,
              "15 有视频的正常显示")
        # 10:没选日期 / 没勾员工点「摆下一批」
        w.start_batch()
        check(boxes and "左上角" in boxes[-1], f"10 没选日期:提示先选日期 -> {boxes[-1:]}")
        w.c_date.setCurrentIndex(i)
        check(w.lab_sum.text() == GW.PICK_EMP, f"13 选了日期没勾员工:提示勾员工 -> {w.lab_sum.text()}")
        w.start_batch()
        check("你还没勾员工" in boxes[-1] and "发完" not in boxes[-1], f"10 没勾员工:不再说「都发完了」 -> {boxes[-1]}")
        # 11:名字带「已发」的文件夹留灰行;不认识的文件夹也列出来
        top = next(w.tree.topLevelItem(k) for k in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(k).text(0) == "家长小天地")
        kids = [top.child(k) for k in range(top.childCount())]
        grey = [x for x in kids if "已发" in x.text(0)]
        check(len(grey) == 1 and grey[0].text(1) == "1 条,不发" and not grey[0].flags() & Qt.ItemFlag.ItemIsUserCheckable
              and "已发" in grey[0].toolTip(0), f"11 带「已发」的文件夹留一行灰的 -> {[x.text(0) for x in kids]}")
        unk = [w.tree.topLevelItem(k) for k in range(w.tree.topLevelItemCount()) if w.tree.topLevelItem(k).text(0) == "没读的文件夹"]
        check(unk and unk[0].child(0).text(0) == "不认识", "11 认不出号的文件夹也列出来")
        for x in kids:
            if x.flags() & Qt.ItemFlag.ItemIsUserCheckable:
                x.setCheckState(0, Qt.CheckState.Checked)
        grey[0].setCheckState(0, Qt.CheckState.Checked)                 # 灰行就算被勾上也不算
        check(sorted(e for _a, _d, e in w.checked_dirs()) == ["李", "王"], f"11 灰行勾不上 -> {w.checked_dirs()}")
        w._load_videos()
        _wait(app, lambda: not (w.cover_worker and w.cover_worker.isRunning()))
        w._update_scope()
        # 18:这次摆 N 条 / 跳过 N 条(为什么)
        sc = w.lab_scope.text()
        check("这次摆「家长小天地」2 条" in sc and "跳过 1 条" in sc and "同名" in sc, f"18 主按钮上面写这次摆几条、跳过几条 -> {sc}")
        # 17:用词统一
        stt = [w.table.item(r, 4).text() for r in range(w.table.rowCount())]
        check(stt.count("还没发") == 2 and not any(s == "没发" for s in stt), f"17 状态写「还没发」 -> {stt}")
        check(GW.READY == "摆好等发布" and w.t_sum.horizontalHeaderItem(4).text() == "摆好等发布", "17 「摆好没发」统一成「摆好等发布」")
        # 16:「封面都看过了」;摆好等确认时底栏只留要点的
        check(w.b_allok.text() == "封面都看过了", "16 「全部调好了」改名「封面都看过了」")
        check(not w.b_wm.isHidden() and not w.b_blur.isHidden(), "16 平时有「查水印 / 模糊水印」")
        w.batch_id = "测试批"
        w._refresh_buttons()
        check(w.b_wm.isHidden() and w.b_blur.isHidden() and w.b_allok.isHidden() and not w.b_sent.isHidden(),
              "16 摆好等确认时藏起「查水印 / 模糊水印 / 封面都看过了」")
        w.batch_id = ""
        w._refresh_buttons()
        # 12:定时框先写总览,标签「几点发出去」,两页意思写清楚
        ov = w._sched_overview("10.21")
        check("家长小天地 2 条" in ov, f"12 教育号总览:同名的只算一条 -> {ov}")
        seen = {}

        def look(dlg):
            seen["labels"] = [x.text() for x in dlg.findChildren(QLabel)] if dlg else []
            if dlg:
                dlg.reject()
        _modal(app, look)
        r = S.ask_schedule(w, "定时发布 · 教育号", ["10.21", "9.30"], "10.21", head="摆完电脑就能关", overview=w._sched_overview)
        labs = "\n".join(seen.get("labels", []))
        check(r is None and "家长小天地 2 条" in labs and "发出去" in labs and "几点发出去(北京时间):" in labs and "摆完电脑就能关" in labs,
              f"12 定时框最上面是总览、标签「几点发出去」 -> {labs[:200]}")
        api_ov = w.api_tab._sched_overview("10.22")
        check("CSGO战神君 1 条" in api_ov, f"12 16 个号总览 -> {api_ov}")
        # 15:16 个号页也是最新的在上、今天标出来
        tab = w.api_tab
        ti = tab.c_date.findData(today)
        check(ti > 0 and "今天" in tab.c_date.itemText(ti) and tab.c_date.itemData(1) == newest_first(
            [tab.c_date.itemData(k) for k in range(1, tab.c_date.count())])[0], "15 16 个号页:最新的在上、今天标出来")
        # 14:还没搬账号(nicks 空)时「今天达标」别写「都发够了」
        apiacc.save([])
        check(tab.set_day("10.22") and tab.videos, "16 个号页选到 10.22")
        sub = tab._quota_sub()
        check("差" in sub and "都发够了" not in sub, f"14 0/20 时不写「都发够了」 -> {sub}")
        apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "a", "secret": apiacc.seal("s"),
                      "category": 1604, "enabled": True}])
        # 20:数据记录页空表有提示
        w.refresh_records()
        check(not w.rec_empty.isHidden() and "还没有记录" in w.rec_empty.text(), f"20 空表有提示 -> {w.rec_empty.text()}")
        check(w.t_rec.horizontalHeader().sectionResizeMode(7) == GW.QHeaderView.ResizeMode.Stretch, "20 备注那列占剩下的宽度")
        # 23:没发成的直接「重试 / 标成已发」
        tab.reload_dates()
        tab.set_day("10.22")
        v = tab.videos[0]
        rec.api_result(v, "没发成", "网断了")
        w.refresh_records()
        check(w.rec_empty.isHidden(), "20 有记录了提示就藏起来")
        w.t_rec.selectRow(0)
        check(not w.b_rec_retry.isHidden() and not w.b_rec_sent.isHidden() and w.b_rec_unsent.isHidden(),
              "23 选中没发成的:有「重试」「标成已发」")
        w.stack.setCurrentIndex(3)
        w.rec_retry()
        check(w.stack.currentWidget() is tab and tab.day() == "10.22" and tab.selected_rows() == [0],
              f"23 重试 = 跳到 16 个号页、选中这条(不自己发) -> {tab.day()} {tab.selected_rows()}")
        check(rec.status_of(v.path) == "", "23 重试不会自己发")
        w._ask = lambda *a, **k: True
        w.t_rec.selectRow(0)
        w.rec_mark(True)
        check(rec.status_of(v.path) == "已发", "23 标成已发")
        p2 =api_root / "CSGO--CSGO战神君" / "10.22" / "不确定的那条.mp4"
        p2.write_bytes(b"x")
        tab.reload_dates()
        v2 = next(x for x in tab.videos if x.path == str(p2))
        rec.api_begin(v2)
        w.refresh_records()
        row = next(k for k in range(w.t_rec.rowCount()) if w.t_rec.item(k, 3).text() == "不确定的那条")
        w.t_rec.selectRow(row)
        check(w.t_rec.item(row, 4).text() == "不确定发没发" and not w.b_rec_unsent.isHidden() and w.b_rec_retry.isHidden(),
              "23 不确定发没发的:有「标成没发出去」,先不能重试")
        w.rec_mark(False)
        check(rec.status_of(str(p2)) == "", "23 标成没发出去:之后能再发")
        # 21:日期统一写 10.19
        check(w.daily.d.displayFormat() == "M.d", "21 每日汇总的日期写 10.19")
        w.refresh_records()
        pubs = [w.t_rec.item(k, 6).text() for k in range(w.t_rec.rowCount()) if w.t_rec.item(k, 6).text()]
        check(pubs and all(re.match(r"^\d{1,2}\.\d{1,2} \d\d:\d\d", x) for x in pubs), f"21 数据记录的时间写 10.19 10:30 -> {pubs}")
        # 19:设置里「发完挪到」只能改前面那段;改坏了的换回按月份
        got = {}

        def settings(dlg):
            if not dlg:
                return
            got["warn"] = [x.text() for x in dlg.findChildren(QLabel) if "没有 {month}" in x.text()]
            got["tpl"] = [x.text().split("后面自动接上:")[-1] for x in dlg.findChildren(QLabel) if x.text().startswith("后面自动接上:")]
            ed = [x for x in dlg.findChildren(QLineEdit) if x.text() == r"X:\知乎视频\内容查看"]
            got["n"] = len(ed)
            if ed:
                ed[0].setText(str(tmp / "教育归档"))
            dlg.findChild(QDialogButtonBox).button(QDialogButtonBox.StandardButton.Ok).click()
        _modal(app, settings)
        w.edit_settings()
        check(got.get("warn") and r"\{month}月发布视频\企鹅号\教育" in got.get("tpl", []) and got.get("n") == 2,
              f"19 模板那半截灰字不能改;改坏了的红字说换回来了 -> {got}")
        check(st["archive_root"] == os.path.join(str(tmp / "教育归档"), r"{month}月发布视频\企鹅号\教育"),
              f"19 存的是 前半截 + 模板 -> {st['archive_root']}")
        check(st["api_archive"] == r"X:\知乎视频\内容查看\{month}月发布视频\企鹅号\企鹅号软件发布", f"19 没动的不变 -> {st['api_archive']}")
        w.close()
        w.deleteLater()
    finally:
        GW.C.decide, GW.C.note_wide, QMessageBox.information = real
        apiacc.save(accs0)
    # 22:relink 只在原来的文件已经不在了才改(同号同名的另一份不算成已发)
    from qe import folders as F
    a = tmp / "relink" / "甲" / "同名.mp4"
    b = tmp / "relink" / "乙" / "同名.mp4"
    for p in (a, b):
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"x")
    r2 = Records(tmp / "r_relink.sqlite")
    v = F.Video(path=str(a), title="同名", account="家长小天地", employee="甲", date="10.21")
    r2.manual_published(v)
    rid = r2.rows()[0]["id"]
    check(not r2.relink(rid, str(b)) and r2.status_of(str(b)) == "", "22 原来那份还在:另一份同名的不算已发")
    a.unlink()
    check(r2.relink(rid, str(b)) and r2.status_of(str(b)) == "已发", "22 原来那份没了(被挪了 / 改了名):认新位置")


def _no_dialogs():
    """测试里不弹框:中文确认框(gui.api_tab.ask)一律点「是」,提示框一律关掉(不然 offscreen 下会一直等人点)。"""
    from PySide6.QtWidgets import QApplication, QMessageBox
    QApplication.instance() or QApplication([])
    import gui.api_tab as T
    T._real_ask = T.ask                            # test_misfire 要测真的确认框
    T.ask = lambda *a, **k: True
    QMessageBox.information = staticmethod(lambda *a, **k: None)
    QMessageBox.warning = staticmethod(lambda *a, **k: None)


def test_find_redo(tmp: Path):
    print("[找视频重发(用户 2026-10-25):输入标题找发过的 → 视频连封面挪回今天 / 明天的文件夹(结构照原样)、记录改没发成;16 个号能在企鹅号上删]")
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi, redo
    from qe import folders as F
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    from gui.find_dialog import FindDialog
    base = tmp / "fr"
    root = base / "小学生教育类"
    day = root / "10.24" / "家长研习社"
    day.mkdir(parents=True)
    t = "幼小衔接必备 第3集"
    (day / f"{t}.mp4").write_bytes(b"x")
    (day / f"{t}.jpg").write_bytes(b"j")
    (day / t).mkdir()
    (day / t / "图片修改需求 (1).png").write_bytes(b"p")
    (day / "别的.mp4").write_bytes(b"x")
    st = Settings(tmp / "st_fr.json")
    st["material_root"] = str(root)
    st["archive_root"] = str(base / "内容查看" / r"{month}月发布视频\企鹅号\教育")
    rec = Records(tmp / "r_fr.sqlite")
    v = F.Video(path=str(day / f"{t}.mp4"), title=t, account="家长研习社", employee="小王", date="10.24")
    rec.manual_published(v)
    rec.manual_published(F.Video(path=str(day / "别的.mp4"), title="别的", account="家长研习社", employee="小王", date="10.24"))
    F.move_date(st, "10.24")                                  # 这天发完整天挪去归档
    arc = Path(F.archive_target(st, "10.24")) / "家长研习社"
    check((arc / f"{t}.mp4").is_file() and not (root / "10.24").exists(), "准备:教育号这天整个挪去了归档")
    rows = redo.search(rec, "衔接必备")
    check(len(rows) == 1 and rows[0]["title"] == t and redo.search(rec, "  ") == [], "标题里的几个字就找得到;空的不找")
    check(redo.locate(st, rows[0]) == str(arc / f"{t}.mp4"), f"教育号:从归档里认出视频现在在哪 -> {redo.locate(st, rows[0])}")
    check(sorted(os.path.basename(p) for p, _k, _r in redo.extras(str(arc / f"{t}.mp4"))) == sorted([f"{t}.jpg", t]),
          "跟着一起挪的:同名的图 + 封面文件夹,别的视频不带")
    (root / "10.25(星期六)").mkdir()
    check(redo.target_dir(st, rows[0], "10.25") == str(root / "10.25(星期六)" / "家长研习社"),
          "挪回去:那天的文件夹已经有了(带星期的也认)就放进去,下面照原样建「家长研习社」")
    check(redo.target_dir(st, rows[0], "10.26") == str(root / "10.26" / "家长研习社"), "没有那天的文件夹:建「10.26」")

    real_ask = T.ask
    asked = []
    T.ask = lambda *a, **k: asked.append(a) or True
    try:
        dlg = FindDialog(rec, st, after=lambda: asked.append("after"))
        dlg.e.setText("衔接")
        dlg.find()
        check(dlg.lst.count() == 1 and dlg.b_back.isEnabled() and dlg.b_copy.isEnabled() and dlg.b_del.isHidden(),
              "窗口里找得到;教育号能挪回去、能复制标题,「在企鹅号上删掉」不显示(罐头发的删不了)")
        dlg.copy()
        check(QApplication.clipboard().text() == t, "复制标题(去后台搜)")
        dlg.c_day.setCurrentIndex(1)
        md = dlg.c_day.currentData()
        dlg.pull_back()
        dst = Path(redo.target_dir(st, rows[0], md))
        check((dst / f"{t}.mp4").is_file() and (dst / f"{t}.jpg").is_file() and (dst / t / "图片修改需求 (1).png").is_file(),
              f"点「挪回去重发」:视频、封面图、封面文件夹都挪到明天的文件夹 -> {dst}")
        check(not (arc / f"{t}.mp4").exists() and (arc / "别的.mp4").is_file(), "归档里只拿走这一条,别的不动")
        check(rec.status_of(str(day / f"{t}.mp4")) == "" and rec.status_of(str(dst / f"{t}.mp4")) == "" and "after" in asked,
              "记录改成没发,新位置能再发;刷新了页面")
        check(dlg.lst.count() == 1 and "没发" in dlg.lst.item(0).text(), "窗口里那条显示「没发成」")
    finally:
        T.ask = real_ask

    # 16 个号:挪走的位置记录里有;在企鹅号上删
    aroot = base / "游戏视频"
    acc = aroot / "CSGO--CSGO战神君"
    (acc / "10.24").mkdir(parents=True)
    (acc / "10.26").mkdir()
    aarc = base / "接口归档"
    aarc.mkdir()
    st["api_roots"] = [str(aroot)]
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    w = F.Video(path=str(acc / "10.24" / "横屏的Ace.mp4"), title="横屏的Ace", account="CSGO战神君", employee="", date="10.24")
    rid = rec.api_begin(w)
    rec.api_finish(rid, "已发", "article_id=A77 vid=V77")
    (aarc / "横屏的Ace.mp4").write_bytes(b"x")
    (aarc / "横屏的Ace.jpg").write_bytes(b"j")
    rec.mark_moved(rid, str(aarc / "横屏的Ace.mp4"))
    r = redo.search(rec, "Ace")[0]
    check(redo.is_api(r) and os.path.normcase(redo.locate(st, r)) == os.path.normcase(str(aarc / "横屏的Ace.mp4")),
          "16 个号:按记下的挪去位置找到")
    check(redo.target_dir(st, r, "10.26") == str(acc / "10.26"), "挪回:号的文件夹下面那天的日期文件夹")
    deleted = []

    class C:
        def __init__(self, cid, sec, **k):
            self.sec = sec

        def delete_article(self, aid):
            deleted.append((aid, self.sec))
            return {}
    real = omapi.Client
    omapi.Client = C
    T.ask = lambda *a, **k: True
    try:
        dlg = FindDialog(rec, st)
        dlg.e.setText("Ace")
        dlg.find()
        check(not dlg.b_del.isHidden() and dlg.b_del.isEnabled(), "16 个号已发的:「在企鹅号上删掉」能点")
        dlg.delete()
        check(deleted == [("A77", "s")], f"在企鹅号上删掉:用记下的文章编号、这个号的密钥 -> {deleted}")
        check("在企鹅号上删掉了" in rec.conn.execute("SELECT note FROM posts WHERE id=?", (rid,)).fetchone()[0], "记录里写上删掉了")
        new, why = redo.pull_back(st, rec, r, "10.26")
        check(why == "" and (acc / "10.26" / "横屏的Ace.mp4").is_file() and (acc / "10.26" / "横屏的Ace.jpg").is_file()
              and rec.status_of(new) == "", f"16 个号挪回 10.26:视频连封面回去,能再发 -> {why}")
        check(redo.delete_on_platform(dict(r, note="")).startswith("软件没记下"), "没记文章编号的:删不了,说清楚去后台删")
    finally:
        omapi.Client = real
        T.ask = real_ask
    _ = app


def test_edu_verify(tmp: Path):
    print("[教育号也能核对(用户 2026-10-26):钥匙单独存、只查不发;按标题对;打回的打勾、对不上的不打勾;拎回来挪回今天的文件夹]")
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi, verify
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    from gui.verify_dialog import VerifyDialog
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    apiacc.set_qkey("家长研习空间", "ce", "se")
    check("家长研习空间" not in apiacc.nicks() and all(a["nick"] != "家长研习空间" for a in apiacc.load()),
          "教育号的钥匙不进 16 个号的名单:发布那边看不到,不可能被接口发")
    ca = {a["nick"]: a for a in apiacc.check_accounts()}
    check(ca["家长研习空间"].get("query_only") and not ca["CSGO战神君"].get("query_only") and apiacc.unseal(ca["家长研习空间"]["secret"]) == "se",
          "核对的名单:16 个号 + 教育号(标着只查不发,钥匙加密存)")
    n = bj_now()
    day, md = n.strftime("%Y-%m-%d"), f"{n.month}.{n.day}"
    root = tmp / "eduv" / "小学生教育类"
    d = root / md / "家长研习空间"
    d.mkdir(parents=True)
    st = Settings(tmp / "st_eduv.json")
    st["material_root"] = str(root)
    st["archive_root"] = str(tmp / "eduv" / "归档" / "{month}月")
    rec = Records(tmp / "r_eduv.sqlite")
    titles = {"A": "孩子写作业磨蹭怎么办 #家庭教育", "B": "幼小衔接必备:上学前一定要会的十件事,家长赶紧收藏", "C": "被打回的那条", "D": "标题对不上的那条"}
    for k, t in titles.items():
        (d / f"{t}.mp4").write_bytes(b"x")
        rec.manual_published(F.Video(path=str(d / f"{t}.mp4"), title=t, account="家长研习空间", employee="小王", date=md))
    t0 = f"{day} 10:00:00"
    arts = [{"article_title": "孩子写作业磨蹭怎么办#家庭教育", "article_pub_time": t0, "article_pub_flag": "发布成功"},
            {"article_title": "幼小衔接必备:上学前一定要会的十件事", "article_pub_time": t0, "article_pub_flag": "发布成功"},  # 平台截短了
            {"article_title": "被打回的那条", "article_pub_time": t0, "article_pub_flag": "审核不通过"},
            {"article_title": "人在后台手动发的", "article_pub_time": t0, "article_pub_flag": "发布成功"}]

    class C:
        def __init__(self, cid, sec, **k):
            self.cid = cid

        def list_articles(self, page=1, limit=10):
            return {"articles": arts if page == 1 and self.cid == "ce" else [], "total": len(arts)}

        def article_info(self, aid):
            return {}
    r = verify.check_edu(C("ce", "se"), rec, "家长研习空间", day)
    why = {m["title"]: (m["why"], m["sure"]) for m in r["missing"]}
    check(r["ours"] == 4 and r["live"] == 3 and len(r["extra"]) == 1 and set(why) == {titles["C"], titles["D"]},
          f"按标题对:去空格标点一样的、被平台截短的都算对上;后台多出来的列为软件没记的 -> {r['ours']} {r['live']} {why}")
    check("被审核打回" in why[titles["C"]][0] and why[titles["C"]][1] is False and why[titles["D"]][1] is False
          and "改了标题" in why[titles["D"]][0],
          "被打回的:列出来不打勾(再发多半还会被打回);标题找不到、这天后台又有多出来的:不打勾,提示可能改了标题")
    extra_one = arts.pop()                                     # 这天后台没有多出来的 → 找不到的就是没发上,打勾
    r2 = verify.check_edu(C("ce", "se"), rec, "家长研习空间", day)
    why2 = {m["title"]: (m["why"], m["sure"]) for m in r2["missing"]}
    check(why2[titles["D"]][1] is True and "没发上" in why2[titles["D"]][0], f"标题找不到、这天也没多出来的:打勾 -> {why2}")
    arts.append(extra_one)
    from datetime import timedelta
    from qe.config import bj_str
    yday = (n - timedelta(days=1)).strftime("%Y-%m-%d")
    rec.conn.execute("UPDATE posts SET published_at=? WHERE title=?", (f"{yday} 18:00:00", titles["D"]))   # D 改成昨天发的、昨天没发上
    rec.conn.commit()
    rec.manual_published(F.Video(path=str(d / "补发的D.mp4"), title=titles["D"], account="家长研习空间", employee="小王", date=md))  # 今天补发过了
    r3 = verify.check_edu(C("ce", "se"), rec, "家长研习空间", yday)
    m3 = {m["title"]: m for m in r3["missing"]}
    check(titles["D"] in m3 and m3[titles["D"]]["sure"] is False and "已经补发过了" in m3[titles["D"]]["why"],
          f"那天没发上、但后来已经用同样标题补发过的:不打勾,说清楚(不然拎回来会重复发) -> {m3.get(titles['D'], {}).get('why')}")
    rec.conn.execute("DELETE FROM posts WHERE path=?", (str(d / "补发的D.mp4"),))
    rec.conn.execute("UPDATE posts SET published_at=? WHERE title=?", (bj_str(), titles["D"]))
    rec.conn.commit()
    real = omapi.Client
    omapi.Client = C
    try:
        dlg = VerifyDialog(rec, day, st=st)
        _wait(app, lambda: dlg.w is None or dlg.w.isFinished())
        app.processEvents()
        check("家长研习空间" in dlg.results and dlg.results["家长研习空间"].get("edu") and "CSGO战神君" in dlg.results,
              "「去企鹅号核对」16 个号和教育号一起查")
        check(dlg.miss.count() == 2, f"假发成功列了 2 条 -> {dlg.miss.count()}")
        from PySide6.QtCore import Qt
        its = [dlg.miss.item(k) for k in range(dlg.miss.count())]
        check(all(it.checkState() == Qt.CheckState.Unchecked for it in its), "打回的、标题对不上的:默认都不打勾")
        for it in its:
            if "被打回的那条" in it.text():
                it.setCheckState(Qt.CheckState.Checked)       # 人确定要重发:自己勾上
        dlg.pull_back()
        check(rec.status_of(str(d / f"{titles['C']}.mp4")) == "" and (d / f"{titles['C']}.mp4").is_file()
              and not list(d.glob("*(2).mp4")),
              "拎回来:被打回的改成没发,视频就在今天的文件夹里的不挪(不会变成「(2)」)")
        check(rec.status_of(str(d / f"{titles['D']}.mp4")) == "已发", "没打勾的(标题对不上)不动")
        rows = verify.table_rows(st, rec, day, dlg.results)
        check(any(nick == "家长研习空间" and a == 3 for _c, nick, _q, a, _t in rows), "导出表格:教育号用企鹅号上的实发")
        dlg.close()
    finally:
        omapi.Client = real


def test_redo_room(tmp: Path):
    print("[补发不占规定条数(用户 2026-10-26):拎回来重发的是以前哪天发的 → 那天给这个号多加 1 条;当天发的不加;没规定条数的不管]")
    from datetime import timedelta
    from qe import plan
    from qe.config import Settings, bj_now
    st = Settings(tmp / "st_room.json")
    st["daily_plan"] = [{"cat": "教育", "label": "教育总发布量", "accounts": {"家长研习空间": 102}}]
    st["quota_open"] = False
    n = bj_now()
    today, yday = n.strftime("%Y-%m-%d"), (n - timedelta(days=1)).strftime("%Y-%m-%d")
    row = {"account": "家长研习空间", "published_at": f"{yday} 18:00:00"}
    check(plan.room_for_redo(st, row) == 1 and plan.quota_on(st, "家长研习空间", today) == 103,
          "昨天发的被打回、今天补发:今天这个号多加 1 条(今天找好的 102 条照样全发)")
    check(plan.room_for_redo(st, dict(row, published_at=f"{today} 09:00:00")) == 0 and plan.quota_on(st, "家长研习空间", today) == 103,
          "今天发的今天拎回来:不加(改回没发以后今天自然空出一条)")
    check(plan.room_for_redo(st, dict(row, account="没规定条数的号")) == 0, "没规定条数的号不管")
    tmr = n + timedelta(days=1)
    check(plan.md_to_day(f"{tmr.month}.{tmr.day}") == tmr.strftime("%Y-%m-%d") and plan.md_to_day("乱写") == "",
          "「10.27」认成那天的日期")
    check(plan.room_for_redo(st, row, tmr.strftime("%Y-%m-%d")) == 1 and plan.quota_on(st, "家长研习空间", tmr.strftime("%Y-%m-%d")) == 103,
          "挪回明天重发:加在明天")
    check(plan.room_for_redo(st, dict(row, status="没发成")) == 0 and plan.quota_on(st, "家长研习空间", today) == 103,
          "已经拎回来过(没发成)再挪一次:不再加(上次加过了)")
    from qe.records import Records
    s = plan.summary(Records(tmp / "r_room.sqlite"), st, today)
    tsv = plan.as_tsv(s)
    check("教育总发布量\t102\t0" in tsv and "总发布量\t102\t0" in tsv.splitlines()[-1] and s["rows"][0][1] == 103 and s["extra"] == {"家长研习空间": 1},
          f"复制到企业微信的「规定」写原本的 102(补发加的那 1 条不算;用户 2026-10-26「规定的 80」),还差几条照样按 103 算 -> {tsv!r}")
    # 16 个号拎回来:那天给这个号手选过文件夹的,就放进手选的那个(用户 2026-10-26:挪回老日期文件夹,今天的页面看不到)
    from qe import redo
    md = f"{n.month}.{n.day}"
    st["api_picked"] = {md: {"日漫漫谈社": str(tmp / "rm" / "日漫漫谈社" / "10.25（周日）")}}
    arow = {"account": "日漫漫谈社", "employee": "接口", "path": str(tmp / "rm" / "日漫漫谈社" / "10.21" / "痛.mp4")}
    check(redo.target_dir(st, arow, md) == str(tmp / "rm" / "日漫漫谈社" / "10.25（周日）"), "手选过文件夹的号:挪进今天手选的那个文件夹")
    (tmp / "rm" / "日漫漫谈社" / "10.21").mkdir(parents=True)
    check(redo.target_dir(st, dict(arow, account="日漫解说馆"), md) == str(tmp / "rm" / "日漫漫谈社" / md),
          "没手选的:号的文件夹下今天的日期文件夹")


def test_verify_range(tmp: Path):
    print("[批量查询(用户 2026-10-26):从哪天查到哪天,一个号只翻一遍;导出成企业微信那张表(每天一列)+ 没发上的视频]")
    from datetime import timedelta
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi, verify
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    from gui.verify_dialog import VerifyDialog
    import openpyxl
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    apiacc.set_qkey("家长研习空间", "ce", "se")
    n = bj_now()
    days = [(n - timedelta(days=k)).strftime("%Y-%m-%d") for k in (2, 1, 0)]
    st = Settings(tmp / "st_vr.json")
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"CSGO战神君": 2}},
                        {"cat": "教育", "label": "教育总发布量", "accounts": {"家长研习空间": 1}}]
    rec = Records(tmp / "r_vr.sqlite")
    arts = {"c": [], "ce": []}
    for k, d in enumerate(days):
        for j in range(2):                                        # CSGO 每天记了 2 条;第一天的第 2 条企鹅号上没有
            v = F.Video(path=str(tmp / "vr" / f"{d}-{j}.mp4"), title=f"c{k}{j}", account="CSGO战神君", employee="", date=d[5:])
            rid = rec.api_begin(v)
            rec.api_finish(rid, "已发", f"article_id=A{k}{j} vid=V{k}{j}")
            rec.conn.execute("UPDATE posts SET published_at=? WHERE id=?", (f"{d} 10:0{j}:00", rid))
            if not (k == 0 and j == 1):
                arts["c"].append({"article_id": f"A{k}{j}", "article_title": f"c{k}{j}", "article_pub_time": f"{d} 10:0{j}:00",
                                  "article_pub_flag": "发布成功", "article_video_info": {"vid": f"V{k}{j}"}})
        rec.manual_published(F.Video(path=str(tmp / "vr" / f"e{d}.mp4"), title=f"教育{k}", account="家长研习空间", employee="小王", date=d[5:]))
        rec.conn.execute("UPDATE posts SET published_at=? WHERE title=?", (f"{d} 18:00:00", f"教育{k}"))
        arts["ce"].append({"article_title": f"教育{k}", "article_pub_time": f"{d} 18:00:00", "article_pub_flag": "发布成功"})
    rec.conn.commit()
    arts["c"].append({"article_id": "OLD", "article_title": "老的", "article_pub_time": "2020-01-01 00:00:00", "article_pub_flag": "发布成功"})
    pages = []

    class C:
        def __init__(self, cid, sec, **k):
            self.cid = cid

        def list_articles(self, page=1, limit=10):
            pages.append((self.cid, page))
            a = sorted(arts[self.cid], key=lambda x: x["article_pub_time"], reverse=True)
            return {"articles": a[(page - 1) * limit: page * limit], "total": len(a)}

        def article_info(self, aid):
            return {"article_info": {"article_pub_flag": "未发布", "article_del_flag": "未删除"}}
    # 用户 2026-10-26 当天核对:定时还没到点的,企鹅号状态是代号「10」—— 不能当成没发上(原来默认打勾,差点拎回来重发)
    v = F.Video(path=str(tmp / "vr" / "定时的.mp4"), title="定时的", account="CSGO战神君", employee="", date=days[-1][5:])
    rid = rec.api_begin(v)
    rec.api_finish(rid, "已发", "article_id=T1 vid=VT1")
    rec.conn.execute("UPDATE posts SET published_at=? WHERE id=?", (f"{days[-1]} 23:00:00", rid))
    rec.conn.commit()
    arts["c"].append({"article_id": "T1", "article_title": "定时的", "article_pub_time": f"{days[-1]} 23:00:00",
                      "article_pub_flag": "10", "article_video_info": {"vid": "VT1"}})
    one = verify.check_account(C("c", "s"), rec, "CSGO战神君", days[-1])
    check(not one["missing"] and one["live"] == 3 and one["review"] == 1,
          f"定时没到点(状态「10」):算在企鹅号上、等着发,不算没发上 -> missing={[m['why'] for m in one['missing']]} live={one['live']}")
    arts["c"].pop()
    rec.conn.execute("DELETE FROM posts WHERE id=?", (rid,))
    rec.conn.commit()
    pages.clear()
    per = verify.check_range(C("c", "s"), rec, "CSGO战神君", days)
    check([per[d]["live"] for d in days] == [1, 2, 2] and len(per[days[0]]["missing"]) == 1 and pages == [("c", 1)],
          f"一个号一段日子:后台只翻一遍,一天一天对出来 -> {[per[d]['live'] for d in days]} 翻了 {pages}")
    real = omapi.Client
    omapi.Client = C
    try:
        dlg = VerifyDialog(rec, days[-1], st=st)
        _wait(app, lambda: dlg.w is None or dlg.w.isFinished())
        dlg.start(days[0], days[-1])
        _wait(app, lambda: dlg.w is None or dlg.w.isFinished())
        app.processEvents()
        r = dlg.results.get("CSGO战神君") or {}
        check(r.get("live") == 5 and set(dlg.by_day) == {"CSGO战神君", "家长研习空间"} and dlg.miss.count() == 1
              and "·" in dlg.miss.item(0).text() and f"{int(days[0][5:7])}.{int(days[0][8:10])}" in dlg.miss.item(0).text(),
              f"窗口里查一段:合起来一行一个号,没发上的写清楚哪天 -> {r.get('live')} {dlg.miss.count()}")
        p = tmp / "vr.xlsx"
        verify.export_range_xlsx(str(p), st, rec, dlg.days, dlg.by_day)
        wb = openpyxl.load_workbook(p)
        ws, w2 = wb.worksheets[0], wb.worksheets[1]
        row = next(rr for rr in range(2, ws.max_row + 1) if ws.cell(rr, 2).value == "CSGO战神君")
        check([ws.cell(row, 4 + k).value for k in range(3)] == [1, 2, 2] and ws.cell(row, 4).fill.fgColor.rgb.endswith("FFC7CE")
              and ws.cell(1, 4).value == f"{int(days[0][5:7])}月{int(days[0][8:10])}日",
              "导出第一页:每个号一行、每天一列(企鹅号上核实的),比规定少的标红")
        check(w2.cell(2, 2).value == "CSGO战神君" and w2.cell(2, 3).value == "c01" and w2.cell(2, 5).value == "要补",
              "导出第二页:没发上的视频(哪天、哪个号、标题、原因、要不要补)")
        dlg.close()
    finally:
        omapi.Client = real


def test_mark_sent_keep(tmp: Path):
    print("[标成已发保留原来的发布时间;教育号同一天同标题记了两条算一条(用户 2026-10-26:核对说发上的视频没发上)]")
    from qe import folders as F
    from qe import verify
    from qe.records import Records
    rec = Records(tmp / "r_msk.sqlite")
    v = F.Video(path=str(tmp / "msk" / "a.mp4"), title="耗尽自身五行之力", account="国漫趣谈", employee="", date="10.24")
    rid = rec.api_begin(v)
    rec.api_finish(rid, "已发", "article_id=X1 vid=Y1")
    rec.conn.execute("UPDATE posts SET published_at='2026-10-24 15:32:49' WHERE id=?", (rid,))
    rec.conn.commit()
    rec.mark_unsent(rid, "(拎回来)")
    rec.mark_sent(rid)
    got = rec.conn.execute("SELECT status, published_at FROM posts WHERE id=?", (rid,)).fetchone()
    check(tuple(got) == ("已发", "2026-10-24 15:32:49"), f"拎回来以后又标成已发:还是原来那天发的(不变成今天)-> {tuple(got)}")
    w = F.Video(path=str(tmp / "msk" / "b.mp4"), title="没发过的", account="国漫趣谈", employee="", date="10.24")
    rec.failed(w, "b1", "没摆成")
    rid2 = rec.conn.execute("SELECT id FROM posts WHERE title='没发过的'").fetchone()[0]
    rec.mark_sent(rid2)
    check(rec.conn.execute("SELECT published_at FROM posts WHERE id=?", (rid2,)).fetchone()[0], "原来没有发布时间的:记成现在(补记)")
    day = "2026-10-26"
    for k, folder in enumerate(("10.24", "10.26")):              # 同一条视频记了两条(两个文件夹)
        x = F.Video(path=str(tmp / "msk" / folder / "幼小衔接.mp4"), title="幼小衔接必备 AI 工具汇总", account="家长研习空间",
                    employee="小王", date=folder)
        rec.manual_published(x)
    rec.conn.execute("UPDATE posts SET published_at=? WHERE account='家长研习空间'", (f"{day} 14:31:16",))
    rec.conn.commit()

    class C:
        def list_articles(self, page=1, limit=10):
            a = [{"article_title": "幼小衔接必备 AI 工具汇总", "article_pub_time": f"{day} 10:37:16", "article_pub_flag": "发布成功"}]
            return {"articles": a if page == 1 else [], "total": 1}
    r = verify.check_edu(C(), rec, "家长研习空间", day)
    check(r["ours"] == 1 and not r["missing"], f"教育号同一天同标题记了两条:算一条,不报没发上 -> ours={r['ours']} {[m['why'] for m in r['missing']]}")


def test_key_guards(tmp: Path):
    print("[钥匙串号的硬拦(审查 2026-10-26):扫错号不换、教育号进不了发布名单、搬账号不改回旧钥匙、选择框按位置取]")
    import sqlite3
    from PySide6.QtWidgets import QApplication, QInputDialog, QMessageBox
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omlogin
    from qe.config import Settings
    from qe.records import Records
    import gui.api_tab as T
    apiacc.save([{"nick": "PUBG突围君", "client_id": "CID_PUBG", "secret": apiacc.seal("SP"), "category": 1604, "enabled": True},
                 {"nick": "家长研习空间", "client_id": "CID_X", "secret": apiacc.seal("SX"), "category": 1, "enabled": True}])
    apiacc.set_qkey("家长研习空间", "CID_EDU", "SE")
    check([a["nick"] for a in apiacc.load()] == ["PUBG突围君"] and "家长研习空间" not in apiacc.nicks(),
          "只查不发的号就算混进了接口账号.json,发布那边也看不到")
    for f, why in ((lambda: apiacc.set_key("PUBG突围君", "CID_EDU", "SE"), "教育号的钥匙换给 16 个号"),
                   (lambda: apiacc.set_qkey("家长小天地", "CID_PUBG", "SP"), "16 个号的钥匙存给教育号")):
        try:
            f()
            ok = False
        except ValueError:
            ok = True
        check(ok, f"{why}:拒绝")
    keys = [{"name": "国漫解说馆", "client_id": "A"}, {"name": "国漫解说馆z", "client_id": "B"}]
    check(omlogin.match(keys, "国漫解说馆z")["client_id"] == "B" and omlogin.name_ok({"name": "家长研习空间"}, "PUBG突围君") is False
          and omlogin.name_ok({"name": ""}, "PUBG突围君"), "名字完全一样的先认;名字对不上的算扫错号")
    # 搬账号:教育号不搬;已经有钥匙的号不改回旧库里的老钥匙
    db = tmp / "kg_old.db"
    con = sqlite3.connect(db)
    con.execute("create table accounts(nick, source_name, source_id, client_id, client_secret, category, enabled, sort_order, id)")
    con.executemany("insert into accounts values(?,?,?,?,?,?,?,?,?)",
                    [("PUBG突围君", "PUBG", "1", "OLD", "OLDSEC", 1604, 0, 1, 1), ("家长小天地", "家", "2", "C2", "S2", 1, 1, 2, 2),
                     ("CSGO战神君", "CSGO", "3", "C3", "S3", 1604, 1, 3, 3)])
    con.commit()
    con.close()
    apiacc.migrate(str(db), skip={"家长小天地"})
    p = apiacc.find("PUBG突围君")
    check(p["client_id"] == "CID_PUBG" and p["enabled"] is True and apiacc.find("CSGO战神君")["client_id"] == "C3"
          and apiacc.find("家长小天地") is None, "搬账号:已有的号钥匙 / 开关不动,新号照搬,教育号不搬")
    # 登录取钥匙:选 PUBG、扫的是教育号 → 不换;选教育号、扫的是 PUBG → 不存
    st = Settings(tmp / "st_kg.json")
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏", "accounts": {"PUBG突围君": 20}},
                        {"cat": "教育", "label": "教育", "accounts": {"家长研习空间": 102, "家长小天地": 54}}]
    tab = T.ApiTab(st, Records(tmp / "r_kg.sqlite"), lambda m: None)
    scanned = {}

    class P:
        def terminate(self):
            pass
    real = (omlogin.launch, omlogin.wait_keys, QInputDialog.getItem, QMessageBox.warning, QMessageBox.information, T.ask)
    warned = []
    omlogin.launch = lambda nick: (P(), 1)
    omlogin.wait_keys = lambda port, stopped=None: ("读到了", [scanned["k"]], "")
    QMessageBox.warning = staticmethod(lambda *a, **k: warned.append(a[2]))
    QMessageBox.information = staticmethod(lambda *a, **k: None)
    T.ask = lambda *a, **k: True
    try:
        scanned["k"] = {"name": "家长研习空间", "client_id": "CID_EDU2", "secret": "SE2", "media_id": "9"}
        QInputDialog.getItem = staticmethod(lambda *a, **k: ("PUBG突围君", True))
        tab.relogin_key()
        check(apiacc.find("PUBG突围君")["client_id"] == "CID_PUBG" and warned and "不是「PUBG突围君」" in warned[-1],
              f"选 PUBG、手机扫的是教育号:不换,说清楚 -> {warned[-1:] }")
        scanned["k"] = {"name": "", "client_id": "CID_EDU", "secret": "SE", "media_id": ""}
        tab.relogin_key()
        check(apiacc.find("PUBG突围君")["client_id"] == "CID_PUBG" and "教育号" in warned[-1], "回话里没名字、但钥匙是教育号的:也不换")
        scanned["k"] = {"name": "PUBG突围君", "client_id": "CID_PUBG", "secret": "SP", "media_id": "1"}
        QInputDialog.getItem = staticmethod(lambda *a, **k: ("家长小天地" + T.EDU_TAG, True))
        tab.relogin_key()
        check(apiacc.qfind("家长小天地") is None and "不是「家长小天地」" in warned[-1], "选教育号、扫成 PUBG:不存")
        scanned["k"] = {"name": "家长小天地", "client_id": "CID_XTD", "secret": "SX2", "media_id": "8"}
        tab.relogin_key()
        check(apiacc.qfind("家长小天地") and apiacc.qfind("家长小天地")["client_id"] == "CID_XTD", "扫对了:照常存")
        # 一个账号下两个没名字的号:选第二个就是第二个
        two = [{"name": "", "client_id": "CID_A1", "secret": "a", "media_id": ""}, {"name": "", "client_id": "CID_B2", "secret": "b", "media_id": ""}]
        omlogin.wait_keys = lambda port, stopped=None: ("读到了", two, "")
        answers = iter(["PUBG突围君", None])

        def pick(parent, title, text, items, *a, **k):
            v = next(answers)
            return (v, True) if v else (items[1], True)
        QInputDialog.getItem = staticmethod(pick)
        tab.relogin_key()
        check(apiacc.find("PUBG突围君")["client_id"] == "CID_B2", f"几个没名字的号:选第二个换上的就是第二个 -> {apiacc.find('PUBG突围君')['client_id']}")
    finally:
        omlogin.launch, omlogin.wait_keys, QInputDialog.getItem, QMessageBox.warning, QMessageBox.information, T.ask = real
    _ = app


def test_audit_redo(tmp: Path):
    print("[审查 2026-10-26 拎回来重发:不丢文件 / 不加(2) / 用原名 / 不进脏话文件夹 / 不确定的不挪 / 记录跟着走 / 封面跟着走 / 额度能退]")
    from qe import folders as F
    from qe import plan, redo
    from qe.config import Settings, bj_now
    from qe.records import Records
    n = bj_now()
    md, today = f"{n.month}.{n.day}", n.strftime("%Y-%m-%d")
    rec = Records(tmp / "r_ar.sqlite")
    st = Settings(tmp / "st_ar.json")
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏", "accounts": {"CSGO战神君": 20}},
                        {"cat": "教育", "label": "教育", "accounts": {"家长研习空间": 102}}]
    root = tmp / "ar" / "游戏视频"
    acc = root / "CSGO--CSGO战神君"
    (acc / "10.20").mkdir(parents=True)
    (acc / f"{md} 脏话").mkdir()
    (acc / f"{md}（星期一）").mkdir()
    st["api_roots"] = [str(root)]
    arc = tmp / "ar" / "归档"
    arc.mkdir(parents=True)

    def api_row(title, arc_name):
        v = F.Video(path=str(acc / "10.20" / f"{title}.mp4"), title=title, account="CSGO战神君", employee="", date="10.20")
        rid = rec.api_begin(v)
        rec.api_finish(rid, "已发", f"article_id=A{rid} vid=V{rid}")
        rec.conn.execute("UPDATE posts SET published_at='2026-10-20 10:00:00' WHERE id=?", (rid,))
        rec.conn.commit()
        (arc / f"{arc_name}.mp4").write_bytes(b"x")
        (arc / f"{arc_name}.jpg").write_bytes(b"j")
        rec.mark_moved(rid, str(arc / f"{arc_name}.mp4"))
        return dict(rec.conn.execute("SELECT * FROM posts WHERE id=?", (rid,)).fetchone())
    r1 = api_row("好标题", "好标题 (2)")
    check(redo.target_dir(st, r1, md) == str(acc / f"{md}（星期一）"), "认今天的文件夹和发布页一样:「脏话」那个不算")
    out = {}
    new, why = redo.pull_back(st, rec, r1, md, out)
    check(why == "" and os.path.basename(new) == "好标题.mp4" and (acc / f"{md}（星期一）" / "好标题.jpg").is_file()
          and not list((acc / f"{md}（星期一）").glob("*(2)*")), f"归档里叫「好标题 (2)」的:挪回来用原名,封面跟着改名 -> {new} {why}")
    got = rec.conn.execute("SELECT status, path, moved_pkey FROM posts WHERE id=?", (r1["id"],)).fetchone()
    check(got[0] == "没发成" and os.path.normcase(got[1]) == os.path.normcase(new) and got[2] is None and out.get("room") == 1,
          "记录改成没发成、位置改成新的;补发以前的给今天加 1 条")
    from qe.records import pkey
    check(rec.status_of(new) == "" and rec.conn.execute("SELECT count(*) FROM posts WHERE pkey=? AND status='没发成'", (pkey(new),)).fetchone()[0] == 1
          and plan.quota_on(st, "CSGO战神君", today) == 21, "新位置认得出这条、能再发(之后重发成了就不再挂着没发成)")
    check(plan.undo_room(st, r1["id"]) == 1 and plan.quota_on(st, "CSGO战神君", today) == 20 and plan.undo_room(st, r1["id"]) == 0,
          "拎回来以后又标成已发:多加的那 1 条退回去(只退一次)")
    # 目标已经有同名的:停下来,原处不动,不加 (2)
    r2 = api_row("撞名的", "撞名的")
    (acc / f"{md}（星期一）" / "撞名的.mp4").write_bytes(b"y")
    new2, why2 = redo.pull_back(st, rec, r2, md)
    check(new2 == "" and "已经有" in why2 and (arc / "撞名的.mp4").is_file()
          and rec.conn.execute("SELECT status FROM posts WHERE id=?", (r2["id"],)).fetchone()[0] == "已发",
          f"今天的文件夹里已经有同名的:不挪、不加 (2)、记录不动 -> {why2}")
    # 不确定发没发 / 摆好的:不让挪
    r3 = api_row("不确定的", "不确定的")
    rec.conn.execute("UPDATE posts SET status='结果不明' WHERE id=?", (r3["id"],))
    rec.conn.commit()
    new3, why3 = redo.pull_back(st, rec, dict(r3, status="结果不明"), md)
    check(new3 == "" and "不确定发没发" in why3 and (arc / "不确定的.mp4").is_file(), f"「不确定发没发」的不让挪 -> {why3}")
    # 同一个视频两条「已发」记录:一起改,额度只加一次
    r4 = api_row("两条记录", "两条记录")
    rid_b = rec.api_begin(F.Video(path=r4["path"], title="两条记录", account="CSGO战神君", employee="", date="10.20"))
    rec.api_finish(rid_b, "已发", "")
    rec.conn.execute("UPDATE posts SET published_at='2026-10-20 11:00:00', moved_pkey=? WHERE id=?", (r4["moved_pkey"], rid_b))
    rec.conn.commit()
    o4 = {}
    redo.pull_back(st, rec, r4, md, o4)
    o4b = {}
    redo.pull_back(st, rec, dict(rec.conn.execute("SELECT * FROM posts WHERE id=?", (rid_b,)).fetchone()), md, o4b)
    check([x[0] for x in rec.conn.execute("SELECT status FROM posts WHERE id IN (?,?)", (r4["id"], rid_b))] == ["没发成", "没发成"]
          and o4.get("room") == 1 and not o4b.get("room"), "同一个视频的两条记录一起改成没发;额度只加一次")
    # 号的视频直接放在号文件夹里(没日期文件夹、名单里也找不到):挪回原文件夹
    flat = tmp / "ar" / "别处" / "电竞赛事"
    flat.mkdir(parents=True)
    v = F.Video(path=str(flat / "赛事.mp4"), title="赛事", account="不在名单的号", employee="", date="10.20")
    rid5 = rec.api_begin(v)
    rec.api_finish(rid5, "已发", "")
    (arc / "赛事.mp4").write_bytes(b"x")
    rec.mark_moved(rid5, str(arc / "赛事.mp4"))
    check(redo.target_dir(st, dict(rec.conn.execute("SELECT * FROM posts WHERE id=?", (rid5,)).fetchone()), md) == str(flat),
          "没日期文件夹的号:挪回原来的文件夹(原来直接报找不到)")
    # 教育号:封面在「梁\\封面\\标题\\」;原来套了一层同名日期
    eroot = tmp / "ar" / "小学生教育类"
    st["material_root"] = str(eroot)
    st["archive_root"] = str(tmp / "ar" / "教育归档" / "{month}月")
    day = eroot / "10.20"
    (day / "10.20" / "家长研习空间" / "梁" / "视频").mkdir(parents=True)
    vp = day / "10.20" / "家长研习空间" / "梁" / "视频" / "孩子磨蹭.mp4"
    vp.write_bytes(b"x")
    (day / "10.20" / "家长研习空间" / "梁" / "封面" / "孩子磨蹭").mkdir(parents=True)
    (day / "10.20" / "家长研习空间" / "梁" / "封面" / "孩子磨蹭" / "横屏16x9.png").write_bytes(b"p")
    rec.manual_published(F.Video(path=str(vp), title="孩子磨蹭", account="家长研习空间", employee="梁", date="10.20"))
    F.move_date(st, "10.20")
    er = dict(rec.conn.execute("SELECT * FROM posts WHERE title='孩子磨蹭'").fetchone())
    (eroot / f"{md}不能用").mkdir(parents=True)
    tdir = redo.target_dir(st, er, md)
    check(tdir == str(eroot / md / "家长研习空间" / "梁" / "视频"), f"教育号:不进「不能用」,去掉套着的同名日期那层 -> {tdir}")
    o5 = {}
    new5, why5 = redo.pull_back(st, rec, er, md, o5)
    check(why5 == "" and (eroot / md / "家长研习空间" / "梁" / "封面" / "孩子磨蹭" / "横屏16x9.png").is_file(),
          f"教育号放在「梁\\封面\\标题\\」的封面跟着挪 -> {why5} {o5.get('warn')}")


def test_audit_verify(tmp: Path):
    print("[审查 2026-10-26 核对:只有确定没发上的才打勾(跨半夜 / 翻页重复 / 单查出错 / 删掉的重复 / 系列标题 / 空状态)]")
    from qe import folders as F
    from qe import verify
    from qe.records import Records
    rec = Records(tmp / "r_av.sqlite")
    d0 = "2026-10-20"

    def put(title, edu=False, aid="", at=f"{d0} 10:00:00", acc="CSGO战神君"):
        v = F.Video(path=str(tmp / "av" / f"{title}.mp4"), title=title, account=acc, employee="小王" if edu else "", date="10.20")
        if edu:
            rec.manual_published(v)
            rid = rec.conn.execute("SELECT max(id) FROM posts").fetchone()[0]
        else:
            rid = rec.api_begin(v)
            rec.api_finish(rid, "已发", f"article_id={aid} vid=V{aid}" if aid else "")
        rec.conn.execute("UPDATE posts SET published_at=? WHERE id=?", (at, rid))
        rec.conn.commit()

    class C:
        def __init__(self, arts, info=None, dup_first=False):
            self.arts, self.info, self.dup_first = sorted(arts, key=lambda a: a["article_pub_time"], reverse=True), info, dup_first

        def list_articles(self, page=1, limit=10):
            a = self.arts
            if self.dup_first and page == 2:                 # 翻页中间有新文章进来:整张表往下挪,第 1 页最后一条又出现
                a = [a[0]] + a
            return {"articles": a[(page - 1) * limit: page * limit], "total": len(self.arts) + 1}

        def article_info(self, aid):
            if self.info == "err":
                raise RuntimeError("查文章:网络连不上")
            return {}

    def art(t, at, flag="发布成功", aid="", dele=""):
        return {"article_id": aid, "article_title": t, "article_pub_time": at, "article_pub_flag": flag, "article_del_flag": dele}
    # 跨半夜(教育号):记录记在 20 号,企鹅号上是 21 号 00:05
    put("跨半夜的", edu=True, at=f"{d0} 23:55:00", acc="家长研习空间")
    r = verify.check_edu(C([art("跨半夜的", "2026-10-21 00:05:00")]), rec, "家长研习空间", d0)
    check(not r["missing"], f"跨半夜发的:前后一天有就算发上了 -> {[m['why'] for m in r['missing']]}")
    # 翻页重复:同一篇不算两次
    many = [art(f"t{k}", f"{d0} 0{k % 9}:00:00", aid=f"P{k}") for k in range(15)]
    got = verify.platform_between(C(many, dup_first=True), d0, d0)
    check(len(got) == len({p["aid"] for p in got}) == 15, f"翻页中间表往下挪:同一篇只算一次 -> {len(got)}")
    # 翻到上限还没翻完:报查不了,不当成读完
    try:
        verify.platform_between(C([art("x", f"{d0} 10:00:00", aid=f"Q{k}") for k in range(30)]), "2020-01-01", d0, max_pages=2)
        ok = False
    except verify.ListTooLong:
        ok = True
    check(ok, "翻满页数上限还没翻到那天以前:报「查不了」")
    # 单查出错:不打勾
    put("单查出错的", aid="E1")
    r = verify.check_account(C([], info="err"), rec, "CSGO战神君", d0)
    m = next(x for x in r["missing"] if x["title"] == "单查出错的")
    check(m["sure"] is False and "看一眼" in m["why"], f"列表里没有、单查又出错:不打勾 -> {m['why']}")
    # 删掉的是重复的那条:同标题的在线 → 不打勾
    put("发重了", aid="D1", at=f"{d0} 11:00:00")
    put("发重了", aid="D2", at=f"{d0} 11:05:00")
    r = verify.check_account(C([art("发重了", f"{d0} 11:00:00", aid="D1"), art("发重了", f"{d0} 11:05:00", aid="D2", dele="已删除")]),
                             rec, "CSGO战神君", d0)
    m = next((x for x in r["missing"] if x["title"] == "发重了"), None)
    check(m is not None and m["sure"] is False and "同标题" in m["why"], f"删掉的是发重了的那条:不打勾 -> {m and m['why']}")
    # 被打回的:列出来,不打勾
    put("被打回的", aid="R1")
    r = verify.check_account(C([art("被打回的", f"{d0} 10:00:00", flag="审核不通过", aid="R1")]), rec, "CSGO战神君", d0)
    m = next(x for x in r["missing"] if x["title"] == "被打回的")
    check(m["sure"] is False and "打回" in m["why"], "被打回的:不打勾(再发多半还会被打回)")
    # 系列标题:第12集被打回,第13集后来发了 → 不算「已经补发过了」
    put("三角洲行动高光时刻第12集", aid="S12")
    put("三角洲行动高光时刻第13集", aid="S13", at="2026-10-21 10:00:00")
    r = verify.check_account(C([art("三角洲行动高光时刻第12集", f"{d0} 10:00:00", flag="未发布", aid="S12")]), rec, "CSGO战神君", d0)
    m = next(x for x in r["missing"] if x["title"] == "三角洲行动高光时刻第12集")
    check(m["sure"] is True and not m.get("redone"), f"第12集没发上、第13集后来发了:不认成补发过,照样打勾 -> {m['why']}")
    # 状态是空的:认不出,不打勾,也不算在线
    put("状态空的", aid="N1")
    r = verify.check_account(C([art("状态空的", f"{d0} 10:00:00", flag="", aid="N1")]), rec, "CSGO战神君", d0)
    m = next((x for x in r["missing"] if x["title"] == "状态空的"), None)
    check(m is not None and m["sure"] is False and "认不出" in m["why"], f"状态是空的:列出来看一眼 -> {m and m['why']}")


def test_autocover(tmp: Path):
    print("[没放封面的自动挑一帧(用户 2026-10-26:平台拿第一帧当封面,很多是黑屏;选了全自动)]")
    import subprocess
    import cv2
    import numpy as np
    from qe import autocover
    from qe.vertical import ffmpeg_path
    ff = ffmpeg_path()
    if not ff:
        print("  (没有 ffmpeg,跳过)")
        return
    d = tmp / "ac"
    d.mkdir()
    v1, v2 = d / "开头黑屏.mp4", d / "全黑.mp4"
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "color=c=black:s=320x568:d=2",
                    "-f", "lavfi", "-i", "testsrc=s=320x568:d=6", "-filter_complex", "[0][1]concat=n=2:v=1[v]", "-map", "[v]",
                    "-pix_fmt", "yuv420p", str(v1)], check=True)
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "color=c=black:s=320x568:d=5",
                    "-pix_fmt", "yuv420p", str(v2)], check=True)
    p = autocover.pick(str(v1))
    img = cv2.imdecode(np.fromfile(p, np.uint8), cv2.IMREAD_GRAYSCALE) if p else None
    check(p and img is not None and img.mean() > 40 and img.std() > 20 and os.path.dirname(p) == autocover.cache_dir(),
          f"开头黑屏的:挑的是后面清楚的一帧,放在软件自己的文件夹里(不碰素材) -> {p}")
    check(not [x for x in os.listdir(d) if x.lower().endswith(".jpg")], "素材文件夹里没多出东西")
    n = autocover.next_frame(str(v1))
    check(n and n != p and autocover.pick(str(v1)) == n, "「换一帧」换成另一张,之后发的时候就用它")
    check(autocover.pick(str(v2)) == "", "整条都是黑屏的:挑不出来,让平台自己截")
    v3 = d / "黑边竖屏.mp4"                                  # 转过竖屏的:上下黑边,中间是画面 —— 不能当成黑屏
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "testsrc=s=540x304:d=5",
                    "-vf", "pad=540:960:0:328:black", "-pix_fmt", "yuv420p", str(v3)], check=True)
    check(autocover.pick(str(v3)) != "", "上下黑边、中间有画面的(转过竖屏的):照样挑得出封面")


def test_unsure_check(tmp: Path):
    print("[结果不明的软件自己去企鹅号查(用户 2026-10-26,选了 B):查到了记成已发;查不到只提示,不自动改]")
    from datetime import timedelta
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, omapi, unsure
    from qe import folders as F
    from qe.config import Settings, bj_now, bj_str
    from qe.records import Records, pkey
    import gui.api_tab as T
    n = bj_now()
    md = f"{n.month}.{n.day}"
    root = tmp / "uc" / "游戏视频"
    d = root / "CSGO--CSGO战神君" / md
    d.mkdir(parents=True)
    for t in ("发上了的", "没发上的"):
        (d / f"{t}.mp4").write_bytes(b"x")
    apiacc.save([{"nick": "CSGO战神君", "source_name": "CSGO", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    rec = Records(tmp / "r_uc.sqlite")
    ago = bj_str(n - timedelta(minutes=20))
    ids = {}
    for t, vid in (("发上了的", "VID1"), ("没发上的", "VID2")):
        v = F.Video(path=str(d / f"{t}.mp4"), title=t, account="CSGO战神君", employee="", date=md)
        rid = rec.api_begin(v)
        rec.api_finish(rid, "结果不明", f"vid={vid};发布:网络连不上企鹅号")
        rec.conn.execute("UPDATE posts SET prepared_at=? WHERE id=?", (ago, rid))
        ids[t] = rid
    rec.conn.commit()
    pub = bj_str(n - timedelta(minutes=19))
    arts = [{"article_id": "A1", "article_title": "别的", "article_pub_time": pub, "article_video_info": {"vid": "VID1"}},
            {"article_id": "A9", "article_title": "老的", "article_pub_time": bj_str(n - timedelta(days=3)), "article_video_info": {"vid": "OLD"}}]

    class C:
        def __init__(self, *a, **k):
            pass

        def list_articles(self, page=1, limit=10):
            return {"articles": arts if page == 1 else [], "total": len(arts)}
    r1 = dict(rec.conn.execute("SELECT * FROM posts WHERE id=?", (ids["发上了的"],)).fetchone())
    check(unsure.find(C(), r1)["article_id"] == "A1", "按视频编号认出来(标题被平台改了也认得)")
    r2 = dict(rec.conn.execute("SELECT * FROM posts WHERE id=?", (ids["没发上的"],)).fetchone())
    check(unsure.find(C(), r2) is None, "后台没有的:找不到")
    st = Settings(tmp / "st_uc.json")
    st["api_roots"] = [str(root)]
    real = omapi.Client
    omapi.Client = C
    try:
        tab = T.ApiTab(st, rec, lambda m: None)
        tab.set_day(md)
        _wait(app, lambda: not tab.busy() and len(tab.videos) == 2)
        tab.check_unsure(force=True)
        _wait(app, lambda: tab._unsure_thread is not None and not tab._unsure_thread.is_alive(), 20)
        for _ in range(20):
            app.processEvents()
        got = rec.conn.execute("SELECT status, published_at, note FROM posts WHERE id=?", (ids["发上了的"],)).fetchone()
        check(got[0] == "已发" and got[1] == pub and got[2].startswith("article_id=A1"),
              f"查到了:记成已发,发布时间 / 文章编号用企鹅号上的 -> {tuple(got)}")
        i2 = next(k for k, v in enumerate(tab.videos) if v.title == "没发上的")
        i1 = next(k for k, v in enumerate(tab.videos) if v.title == "发上了的")
        check(rec.conn.execute("SELECT status FROM posts WHERE id=?", (ids["没发上的"],)).fetchone()[0] == "结果不明"
              and "多半没发上" in tab._status_text(i2)[0] and tab.state.get(i1) == "已发",
              f"查不到、发出去 20 分钟了:不自动改,只提示「多半没发上」(用户选的 B)-> {tab._status_text(i2)[0]}")
    finally:
        omapi.Client = real


def test_spill_picked_root(tmp: Path):
    print("[发够了:手选的是号的总文件夹(没日期文件夹)的,剩下的也挪进它里面明天的日期文件夹(用户 2026-10-27)]")
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    from qe import apiacc, plan
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    import gui.api_tab as T
    n = bj_now()
    md, tmr = f"{n.month}.{n.day}", plan.tomorrow_folder()
    root = tmp / "spr" / "游戏视频"
    acc = root / "地下城与勇士--DNF连招档案馆"
    acc.mkdir(parents=True)
    for k in range(3):
        (acc / f"剩下的{k}.mp4").write_bytes(b"x")
        (acc / f"剩下的{k}.jpg").write_bytes(b"j")
    apiacc.save([{"nick": "DNF连招档案馆", "source_name": "DNF", "client_id": "c", "secret": apiacc.seal("s"), "category": 1604, "enabled": True}])
    rec = Records(tmp / "r_spr.sqlite")
    rec.api_result(F.Video(path=str(tmp / "spr" / "发过的.mp4"), title="发过的", account="DNF连招档案馆", employee="", date=md), "已发")
    st = Settings(tmp / "st_spr.json")
    st["api_roots"] = [str(root)]
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"DNF连招档案馆": 1}}]
    st["api_picked"] = {md: {"DNF连招档案馆": str(acc)}}
    tab = T.ApiTab(st, rec, lambda m: None)
    tab.set_day(md)
    _wait(app, lambda: not tab.busy() and not list(acc.glob("*.mp4")), 20)
    got = sorted(x.name for x in (acc / tmr).iterdir()) if (acc / tmr).is_dir() else []
    check(not list(acc.glob("*.mp4")) and got == sorted([f"剩下的{k}{e}" for k in range(3) for e in (".jpg", ".mp4")]),
          f"今天发够了:总文件夹里剩下的连封面挪进「{tmr}」(明天页面自己认得出) -> {got}")


def test_quota_free(tmp: Path):
    print("[有多少发多少的号(用户 2026-10-27:电竞情报馆规定 80,但强时效、不一定有比赛):不卡条数、少了不算没发够,规定照写]")
    from qe import plan
    from qe import folders as F
    from qe.config import Settings, bj_now
    from qe.records import Records
    st = Settings(tmp / "st_qf.json")
    st["daily_plan"] = [{"cat": "游戏", "label": "游戏总发布量", "accounts": {"电竞情报馆": 80, "PUBG突围君": 20}}]
    rec = Records(tmp / "r_qf.sqlite")
    day = bj_now().strftime("%Y-%m-%d")
    for k in range(30):
        rec.api_result(F.Video(path=str(tmp / "qf" / f"e{k}.mp4"), title=f"e{k}", account="电竞情报馆", employee="", date="x"), "已发")
    check(plan.is_free(st, "电竞情报馆") and not plan.is_free(st, "PUBG突围君"), "默认:电竞情报馆有多少发多少,别的号照常")
    check(plan.left_today(rec, st, "电竞情报馆", day) is None and plan.left_today(rec, st, "PUBG突围君", day) == 20,
          "电竞不卡条数(不限);别的号照样按规定算还能发几条")
    s = plan.summary(rec, st, day)
    short = [n for n, _q, _a in plan.short_list(s)]
    check("电竞情报馆" not in short and "PUBG突围君" in short and s["rows"][0][1] == 30 + 20,
          "电竞发了 30 条不算没发够;还差几条只算别的号")
    tsv = plan.as_tsv(s)
    check("游戏总发布量\t100\t30" in tsv, f"复制到企业微信:规定照写 100(含电竞的 80),实际 30 -> {tsv!r}")
    plan.set_free(st, "电竞情报馆", False)
    check(plan.left_today(rec, st, "电竞情报馆", day) == 50, "关掉以后:电竞也按 80 条卡")


if __name__ == "__main__":
    tmp = Path(tempfile.mkdtemp(prefix="qe_test_"))
    _no_dialogs()
    test_folders(tmp)
    test_cover_dir(tmp)
    test_scan_any_layout(tmp)
    test_scan_tricky(tmp)
    test_scan_cover_match(tmp)
    test_window_dups(tmp)
    test_cover(tmp)
    test_records(tmp)
    test_failed_steps(tmp)
    test_bj_time()
    test_guard()
    test_window(tmp)
    test_window_scan(tmp)
    test_preview(tmp)
    test_title_zone(tmp)
    test_group(tmp)
    test_api(tmp)
    test_api_tab(tmp)
    test_audit_r1(tmp)
    test_audit_r2(tmp)
    test_ui_r1(tmp)
    test_squarish(tmp)
    test_audit_r3(tmp)
    test_bughunt(tmp)
    test_daily_plan(tmp)
    test_login_key(tmp)
    test_window_renamed(tmp)
    test_at_titles(tmp)
    test_schedule(tmp)
    test_prepare_retry(tmp)
    test_sched_edu(tmp)
    test_sched_api(tmp)
    test_ui_g(tmp)
    test_pub_progress(tmp)
    test_awake(tmp)
    test_yellow_manual(tmp)
    test_api_per_acc(tmp)
    test_quick_review(tmp)
    test_quota_gate(tmp)
    test_game_guard(tmp)
    test_ux_fixes(tmp)
    test_packs(tmp)
    test_clean_dirs(tmp)
    test_black_border(tmp)
    test_edit_quota(tmp)
    test_extra_vert_order(tmp)
    test_verify(tmp)
    test_unsend(tmp)
    test_api_dups(tmp)
    test_view_tomorrow(tmp)
    test_cover_move_keep(tmp)
    test_export_and_edu_dims(tmp)
    test_misfire_api(tmp)
    test_misfire_edu(tmp)
    test_sched_notice(tmp)
    test_cd_items(tmp)
    test_api_wm(tmp)
    test_edu_wm(tmp)
    test_blur(tmp)
    test_ads(tmp)
    test_cover_cache_unique(tmp)
    test_cover_dialog_src(tmp)
    test_api_direct(tmp)
    test_api_cover_pic(tmp)
    test_api_cover(tmp)
    test_wm_detect_unit(tmp)
    test_find_redo(tmp)
    test_edu_verify(tmp)
    test_redo_room(tmp)
    test_verify_range(tmp)
    test_mark_sent_keep(tmp)
    test_key_guards(tmp)
    test_audit_redo(tmp)
    test_audit_verify(tmp)
    test_autocover(tmp)
    test_unsure_check(tmp)
    test_spill_picked_root(tmp)
    test_quota_free(tmp)
    shutil.rmtree(tmp, ignore_errors=True)
    print("\n" + ("全部通过" if not fails else f"失败 {len(fails)} 项:\n  - " + "\n  - ".join(fails)))
    sys.exit(1 if fails else 0)
